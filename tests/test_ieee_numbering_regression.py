from __future__ import annotations

import io
from pathlib import Path
from zipfile import ZipFile
import unittest

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from formatter.equation_detector import detect_equation_blocks, detect_unnumbered_equation_blocks
from formatter.equation_utils import is_unnumbered_equation_block, match_equation_label_paragraph
from formatter.formatter_engine import FormattingOptions, ProcessingMode, process_document
from formatter.streamlit_pipeline import extract_docx_text_lines, process_uploaded_docx


FIXTURE = Path(__file__).resolve().parent / "fixtures" / "ieee_numbering_regression.docx"


class IeeeNumberingRegressionTests(unittest.TestCase):
    def test_explanatory_prose_with_variables_is_not_an_equation_candidate(self):
        doc = Document()
        prose = doc.add_paragraph("The equivalent core-loss resistance Rc is calculated as:\t(17)")
        prose.alignment = WD_ALIGN_PARAGRAPH.CENTER
        vector_text = doc.add_paragraph(
            "The no-load primary current Ip(oc) is vectorially decomposed into the active core-loss current"
        )
        vector_text.alignment = WD_ALIGN_PARAGRAPH.CENTER

        self.assertIsNone(match_equation_label_paragraph(prose))
        self.assertFalse(is_unnumbered_equation_block(vector_text))
        self.assertFalse(detect_equation_blocks(doc))
        self.assertFalse(detect_unnumbered_equation_blocks(doc))

    def test_display_formula_text_is_still_detected(self):
        doc = Document()
        numbered = doc.add_paragraph("Rc = vp(oc)^2 / Poc \u2248 6297.6 \u03a9\t(17)")
        numbered.alignment = WD_ALIGN_PARAGRAPH.CENTER
        unnumbered = doc.add_paragraph("V = IR")
        unnumbered.alignment = WD_ALIGN_PARAGRAPH.CENTER

        self.assertIsNotNone(match_equation_label_paragraph(numbered))
        self.assertTrue(is_unnumbered_equation_block(unnumbered))

    def test_template_matching_preserves_table_colon_separator(self):
        doc = Document()
        caption = doc.add_paragraph("Table 1: Measured Secondary Voltage and Current Data")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Voltage"
        table.cell(0, 1).text = "Current"
        before_xml = caption._p.xml

        result = process_document(
            doc,
            mode=ProcessingMode.FIX_NUMBERING_AND_FORMATTING,
            options=FormattingOptions(format_profile_name="template_matching"),
        )

        self.assertEqual(doc.paragraphs[0].text, "Table 1: Measured Secondary Voltage and Current Data")
        self.assertEqual(before_xml, caption._p.xml)
        self.assertFalse([action for action in result.actions if action.kind == "table"])

    def test_template_matching_noop_does_not_restyle_existing_figure_caption(self):
        doc = Document()
        caption = doc.add_paragraph("Fig. 1. Phasor diagram for the short-circuit test")
        caption.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        before_xml = caption._p.xml

        result = process_document(
            doc,
            mode=ProcessingMode.FIX_NUMBERING_AND_FORMATTING,
            options=FormattingOptions(format_profile_name="template_matching"),
        )

        self.assertEqual(before_xml, caption._p.xml)
        self.assertFalse([action for action in result.actions if action.kind == "figure"])

    def test_template_matching_noop_does_not_restyle_existing_equation(self):
        doc = Document()
        equation = doc.add_paragraph("V = IR\t(1)")
        equation.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        before_xml = equation._p.xml

        result = process_document(
            doc,
            mode=ProcessingMode.FIX_NUMBERING_AND_FORMATTING,
            options=FormattingOptions(format_profile_name="template_matching"),
        )

        self.assertEqual(before_xml, equation._p.xml)
        self.assertFalse([action for action in result.actions if action.kind == "equation"])
        self.assertFalse([action for action in result.actions if action.kind == "table"])

    def test_ieee_table_and_equation_numbering_is_idempotent(self):
        
        if not FIXTURE.exists():
            self.skipTest("Optional document fixture is not distributed")
        doc = Document(FIXTURE)

        first = process_document(
            doc,
            mode=ProcessingMode.FIX_NUMBERING_AND_FORMATTING,
            options=FormattingOptions(format_profile_name="ieee"),
        )
        first_text = _document_text(doc)
        first_table_values = _table_values(doc)

        second = process_document(
            doc,
            mode=ProcessingMode.FIX_NUMBERING_AND_FORMATTING,
            options=FormattingOptions(format_profile_name="ieee"),
        )
        second_text = _document_text(doc)

        self.assertEqual(first_text, second_text)
        self.assertFalse(second.actions)
        self.assertIn("Table I systematically delineates the divergence between the results.", first_text)
        self.assertIn("As shown in Table II, the measured value increases.", first_text)
        self.assertIn("The final values are summarized in Table III.", first_text)
        self.assertNotIn("Table I. systematically", first_text)
        self.assertNotIn("Table I: systematically", first_text)
        self.assertIn("TABLE I", first_text)
        self.assertIn("TABLE II", first_text)
        self.assertIn("TABLE III", first_text)
        self.assertIn("E = mc^2     (4) (4)", first_text)
        self.assertIn("P = VI", first_text)
        self.assertTrue(any("automatic equation renumbering is suspended" in item for item in first.manual_review))
        self.assertEqual(first_table_values, _table_values(doc))
        self.assertIn(("Case 1", "R = 10 ohm"), first_table_values)
        self.assertIn(("Case 2", "V = 12 V"), first_table_values)
        self.assertIn(("Case 3", "P = 24 W"), first_table_values)
        self.assertTrue(first.actions)
        self.assertFalse([action for action in first.actions if action.kind == "equation"])

    def test_streamlit_pipeline_export_matches_regression_contract(self):
        if not FIXTURE.exists():
            self.skipTest("Optional document fixture is not distributed")
        run = process_uploaded_docx(
            report_bytes=FIXTURE.read_bytes(),
            input_file_name=FIXTURE.name,
            mode=ProcessingMode.FIX_NUMBERING_AND_FORMATTING,
            options=FormattingOptions(format_profile_name="ieee"),
        )

        exported_text = "\n".join(extract_docx_text_lines(run.docx_bytes))
        document_xml = _document_xml(run.docx_bytes)

        self.assertIn("Table I systematically delineates the divergence between the results.", exported_text)
        self.assertIn("TABLE I", exported_text)
        self.assertIn("TABLE II", exported_text)
        self.assertIn("TABLE III", exported_text)
        self.assertIn("P = VI", exported_text)
        self.assertNotIn("Table I. systematically", exported_text)
        self.assertIn("E = mc^2     (4) (4)", exported_text)
        self.assertTrue(all(not _has_duplicate_equation_label(action.after) for action in run.result.actions))


def _document_text(doc: Document) -> str:
    return "\n".join(paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip())


def _table_values(doc: Document) -> set[tuple[str, str]]:
    values: set[tuple[str, str]] = set()
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text for cell in row.cells]
            if len(cells) >= 2:
                values.add((cells[0], cells[1]))
    return values


def _document_xml(docx_bytes: bytes) -> str:
    with ZipFile(io.BytesIO(docx_bytes)) as package:
        return package.read("word/document.xml").decode("utf-8")


def _has_duplicate_equation_label(text: str) -> bool:
    import re

    return bool(re.search(r"\((\d+)\)\s+\(\1\)", text or ""))


if __name__ == "__main__":
    unittest.main()
