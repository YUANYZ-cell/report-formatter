import re
import io
from html.parser import HTMLParser
from pathlib import Path
import unittest

from docx import Document
import fitz
from streamlit.testing.v1 import AppTest

from formatter.formatter_engine import FormattingOptions, ProcessingMode
from formatter.localization import STAT_CARD_I18N_KEYS, UI_TEXT, stat_card_text
from formatter.streamlit_pipeline import extract_docx_text_lines, process_uploaded_docx


APP_PATH = Path(__file__).resolve().parents[1] / "app.py"
FORBIDDEN_ZH_STATS_RE = re.compile(r"Card|title|caption|native|image|equation|reference|skipped|toc")
ZH_RESULT_CARD_TITLES = [
    "检测到的图题",
    "检测到的表题",
    "原生表格数量",
    "图片表格数量",
    "公式块数量",
    "公式编号数量",
    "未编号公式数量",
    "参考文献引用数量",
    "已跳过目录项",
    "发现的问题",
]


class _VisibleTextParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []

    def handle_data(self, data):
        cleaned = data.strip()
        if cleaned:
            self.parts.append(cleaned)


def _summary_card_visible_text(app: AppTest) -> str:
    parts = []
    for markdown in app.markdown:
        value = getattr(markdown, "value", "")
        if '<div class="summary-grid">' not in value:
            continue
        parser = _VisibleTextParser()
        parser.feed(value)
        parts.extend(parser.parts)
    return "\n".join(parts)


class LocalizationTest(unittest.TestCase):
    def test_chinese_stat_card_text_contains_no_english_fallback_tokens(self):
        for card_id in STAT_CARD_I18N_KEYS:
            for field in ("title", "caption"):
                text = stat_card_text("zh", card_id, field)
                self.assertNotRegex(text, FORBIDDEN_ZH_STATS_RE, f"{card_id}.{field}: {text}")

    def test_missing_chinese_stat_card_key_uses_chinese_fallback(self):
        key = STAT_CARD_I18N_KEYS["native_tables"]["title"]
        original = UI_TEXT["zh"].pop(key)
        try:
            self.assertEqual(stat_card_text("zh", "native_tables", "title"), "原生表格数量")
        finally:
            UI_TEXT["zh"][key] = original


class AppStateTest(unittest.TestCase):
    def test_issue_preview_button_enters_preview_without_duplicate_uploaders(self):
        """Opening issue preview must render the route instead of recreating the main uploader."""
        app = AppTest.from_file(str(APP_PATH), default_timeout=40)
        sample_path = APP_PATH.parent / "samples" / "sample_bad_numbering_report.docx"

        app.run()
        app.file_uploader(key="report_file").upload(
            sample_path.name,
            sample_path.read_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        app.run()
        app.button(key="run_formatter").click()
        app.run()

        self.assertEqual(len(app.exception), 0)
        self.assertFalse(app.button(key="open_issue_preview").disabled)

        app.button(key="open_issue_preview").click()
        app.run()

        self.assertEqual(len(app.exception), 0)
        self.assertEqual(app.session_state["active_preview_mode"], "issue")
        self.assertEqual(len(app.file_uploader), 2)
        self.assertTrue(
            any("preview-workspace-title'>Issue preview" in item.value for item in app.markdown)
        )

    def test_report_uploader_accepts_pdf_as_an_analysis_input(self):
        self.assertIn('type=["docx", "pdf"]', APP_PATH.read_text(encoding="utf-8"))

    def test_app_does_not_force_flush_debug_prints_at_import_time(self):
        """A hidden Windows launcher has no console handle for forced stdout flushes."""
        self.assertNotIn("flush=True", APP_PATH.read_text(encoding="utf-8"))

    def test_stale_preview_state_does_not_render_duplicate_uploaders(self):
        """A preview route without its cached result must fall back to the upload page cleanly."""
        app = AppTest.from_file(str(APP_PATH), default_timeout=20)

        app.run()
        app.session_state["active_preview_mode"] = "single"
        app.run()

        self.assertEqual(len(app.exception), 0)
        self.assertIsNone(app.session_state["active_preview_mode"])
        self.assertEqual(app.file_uploader(key="report_file").label, "Upload report (.docx or .pdf)")

    def test_hidden_launcher_redirects_streamlit_output_to_logs(self):
        launcher_path = APP_PATH.with_name("启动报告格式化工具.ps1")
        launcher_source = launcher_path.read_text(encoding="utf-8")
        self.assertIn("-RedirectStandardOutput", launcher_source)
        self.assertIn("-RedirectStandardError", launcher_source)

    def test_hidden_launcher_quotes_the_app_path(self):
        launcher_path = APP_PATH.with_name("启动报告格式化工具.ps1")
        self.assertIn("('\"{0}\"' -f $appPath)", launcher_path.read_text(encoding="utf-8"))

    def test_streamlit_pipeline_preserves_duplicate_equation_text_while_auto_edit_is_suspended(self):
        doc = Document()
        doc.add_paragraph("Abstract")
        doc.add_paragraph("V = IR     (1) (1)")
        buffer = io.BytesIO()
        doc.save(buffer)

        run = process_uploaded_docx(
            report_bytes=buffer.getvalue(),
            input_file_name="duplicate_equation.docx",
            mode=ProcessingMode.FIX_NUMBERING_AND_FORMATTING,
            options=FormattingOptions(format_profile_name="ieee"),
        )

        exported_text = "\n".join(extract_docx_text_lines(run.docx_bytes))
        self.assertIn("V = IR     (1) (1)", exported_text)
        self.assertFalse([action for action in run.result.actions if action.kind == "equation"])
        self.assertTrue(any("automatic equation renumbering is suspended" in item for item in run.result.manual_review))

    def test_chinese_ui_language_survives_rerun(self):
        app = AppTest.from_file(str(APP_PATH), default_timeout=20)

        app.run()
        self.assertEqual(len(app.exception), 0)

        app.selectbox(key="ui_language").set_value("zh")
        app.run()
        self.assertEqual(len(app.exception), 0)
        self.assertEqual(app.session_state["ui_language"], "zh")
        self.assertEqual(app.selectbox(key="format_profile_name_zh_select").value, "template_matching")
        self.assertEqual(app.selectbox(key="processing_mode_zh_select").value, "analyze_only")
        self.assertEqual(app.selectbox(key="style_source_zh_select").value, "auto_detect")

        app.run()
        self.assertEqual(len(app.exception), 0)
        self.assertEqual(app.session_state["ui_language"], "zh")
        self.assertIn("设置", [item.value for item in app.subheader])
        self.assertNotIn("输入文件", [item.value for item in app.subheader])
        self.assertEqual(app.file_uploader(key="report_file").label, "上传实验报告（.docx 或 .pdf）")
        self.assertEqual(app.file_uploader(key="template_file").label, "上传可选样式模板（.docx）")

    def test_chinese_upload_and_run_keeps_language(self):
        app = AppTest.from_file(str(APP_PATH), default_timeout=30)
        sample_path = APP_PATH.parent / "samples" / "sample_bad_numbering_report.docx"

        app.run()
        app.selectbox(key="ui_language").set_value("zh")
        app.run()
        app.file_uploader(key="report_file").upload(
            sample_path.name,
            sample_path.read_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        app.run()
        self.assertEqual(len(app.exception), 0)
        self.assertEqual(app.session_state["ui_language"], "zh")
        language_card_text = _summary_card_visible_text(app)
        self.assertIn("自动识别的报告语言", language_card_text)
        self.assertNotRegex(language_card_text, FORBIDDEN_ZH_STATS_RE)

        app.button(key="run_formatter").click()
        app.run()
        self.assertEqual(len(app.exception), 0)
        self.assertEqual(app.session_state["ui_language"], "zh")
        self.assertIn("下载结果", [tab.label for tab in app.tabs])
        zh_card_text = _summary_card_visible_text(app)
        for expected in ZH_RESULT_CARD_TITLES:
            self.assertIn(expected, zh_card_text)
        self.assertNotRegex(zh_card_text, FORBIDDEN_ZH_STATS_RE)

    def test_pdf_upload_runs_analysis_without_creating_a_formatted_docx(self):
        pdf = fitz.open()
        page = pdf.new_page()
        page.insert_text((72, 72), "Fig. 1. PDF analysis sample\nSee Fig. 1 for the response.")
        pdf_bytes = pdf.tobytes()
        pdf.close()

        app = AppTest.from_file(str(APP_PATH), default_timeout=30)
        app.run()
        app.file_uploader(key="report_file").upload("sample.pdf", pdf_bytes, "application/pdf")
        app.run()
        self.assertEqual(len(app.exception), 0)

        app.button(key="run_formatter").click()
        app.run()
        self.assertEqual(len(app.exception), 0)
        self.assertIsNone(app.session_state["latest_docx_bytes"])
        self.assertTrue(app.session_state["latest_report_bytes"])
        self.assertIn("PDF analysis only", app.session_state["latest_result"].report_text)

        app.selectbox(key="ui_language").set_value("en")
        app.run()
        self.assertEqual(len(app.exception), 0)
        en_card_text = _summary_card_visible_text(app)
        self.assertIn("Native Table Count", en_card_text)
        self.assertIn("Reference Citation Count", en_card_text)

        app.selectbox(key="ui_language").set_value("zh")
        app.run()
        self.assertEqual(len(app.exception), 0)
        zh_card_text = _summary_card_visible_text(app)
        self.assertIn("原生表格数量", zh_card_text)
        self.assertIn("参考文献引用数量", zh_card_text)
        self.assertNotRegex(zh_card_text, FORBIDDEN_ZH_STATS_RE)

    def test_streamlit_download_bytes_match_ieee_broken_fixture(self):
        fixture = Path(__file__).resolve().parent / "fixtures" / "optional_external_report.docx"
        if not fixture.exists():
            self.skipTest(f"{fixture} is not available on this machine")

        app = AppTest.from_file(str(APP_PATH), default_timeout=40)
        app.run()
        self.assertEqual(len(app.exception), 0)

        app.selectbox(key="format_profile_name_en_select").set_value("ieee")
        app.selectbox(key="processing_mode_en_select").set_value("fix_numbering_and_formatting")
        app.file_uploader(key="report_file").upload(
            fixture.name,
            fixture.read_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        app.run()
        self.assertEqual(len(app.exception), 0)

        app.button(key="run_formatter").click()
        app.run()
        self.assertEqual(len(app.exception), 0)

        self.assertIn("latest_docx_bytes", app.session_state)
        self.assertIn("latest_run_id", app.session_state)
        self.assertIn(app.session_state["latest_run_id"], app.session_state["latest_docx_file_name"])
        self.assertNotEqual(app.session_state["latest_docx_file_name"], "formatted_report.docx")

        exported_text = "\n".join(extract_docx_text_lines(app.session_state["latest_docx_bytes"]))
        debug_info = app.session_state["latest_debug_info"]
        debug_preview = "\n".join(debug_info["exported_docx_preview"])

        for expected in [
            "V = IR                                                        (1)",
            "Equation (1) defines the voltage-current relationship.",
            "P = VI                                                        (3)",
            "Eq. (3) is referenced here.",
            "Fig. 2. Motor speed response after step input",
            "TABLE II",
            "MEASURED MOTOR SPEED RESULTS",
        ]:
            self.assertIn(expected, exported_text)

        for stale in [
            "P = VI                                                        (2)",
            "Eq. (2) is referenced here.",
            "Figure 3 Motor speed response after step input",
            "Table 3",
        ]:
            self.assertNotIn(stale, exported_text)

        self.assertIn("P = VI                                                        (3)", debug_preview)
        self.assertIn("TABLE II", debug_preview)
        self.assertTrue(debug_info["formatter_engine.__file__"].endswith(r"formatter\formatter_engine.py"))
