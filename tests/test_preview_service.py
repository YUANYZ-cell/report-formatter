from dataclasses import dataclass
import io
import os
import tempfile
import unittest

from docx import Document
import formatter.preview_service as preview_service_module

from formatter.preview_service import (
    PreviewEnvironment,
    PreviewMode,
    RenderedPage,
    build_change_items,
    build_document_version,
    build_pixel_diff_items,
    build_issue_items,
    build_preview_version,
    evaluate_preview_gate,
    estimate_page_number,
    items_for_page,
    render_docx_to_pages,
    render_page_html,
    safe_file_name,
    parse_pdftotext_bbox_layout,
    locate_action_in_pdf,
    transform_pdf_bbox,
)


@dataclass
class DummyIssue:
    severity: str
    category: str
    message: str
    paragraph_index: int | None = None


@dataclass
class DummyAction:
    kind: str
    paragraph_index: int | None
    before: str
    after: str
    message: str


class PreviewGateTest(unittest.TestCase):
    def setUp(self):
        self.available = PreviewEnvironment(
            converter_available=True,
            renderer_available=True,
            converter_label="LibreOffice headless",
            renderer_label="pdftoppm",
            libreoffice_path="soffice",
            pdftoppm_path="pdftoppm",
        )

    def test_issue_preview_requires_latest_result(self):
        gate = evaluate_preview_gate(
            mode=PreviewMode.ISSUE,
            report_bytes=b"docx",
            result_present=False,
            modified_docx_bytes=None,
            current_signature="v1",
            latest_signature="v1",
            environment=self.available,
        )

        self.assertFalse(gate.enabled)
        self.assertEqual(gate.reason_key, "preview_run_required")
        self.assertEqual(gate.button_key, "open_issue_preview")

    def test_comparison_preview_requires_modified_docx(self):
        gate = evaluate_preview_gate(
            mode=PreviewMode.COMPARISON,
            report_bytes=b"docx",
            result_present=True,
            modified_docx_bytes=None,
            current_signature="v1",
            latest_signature="v1",
            environment=self.available,
        )

        self.assertFalse(gate.enabled)
        self.assertEqual(gate.reason_key, "preview_missing_modified")

    def test_stale_settings_disable_preview(self):
        gate = evaluate_preview_gate(
            mode=PreviewMode.ISSUE,
            report_bytes=b"docx",
            result_present=True,
            modified_docx_bytes=None,
            current_signature="v2",
            latest_signature="v1",
            environment=self.available,
        )

        self.assertFalse(gate.enabled)
        self.assertEqual(gate.reason_key, "preview_stale_settings")

    def test_missing_high_fidelity_converter_uses_low_fidelity_fallback(self):
        low_fidelity = PreviewEnvironment(
            converter_available=False,
            renderer_available=True,
            converter_label="DOCX browser preview",
            renderer_label="Pillow text renderer",
            pdftoppm_path="pdftoppm",
            low_fidelity_available=True,
        )
        gate = evaluate_preview_gate(
            mode=PreviewMode.ISSUE,
            report_bytes=b"docx",
            result_present=True,
            modified_docx_bytes=None,
            current_signature="v1",
            latest_signature="v1",
            environment=low_fidelity,
        )

        self.assertTrue(gate.enabled)
        self.assertEqual(gate.reason_key, "preview_low_fidelity_ready")

    def test_preview_route_can_still_be_unavailable(self):
        unavailable = PreviewEnvironment(
            converter_available=False,
            renderer_available=False,
            converter_label="Unavailable",
            renderer_label="Unavailable",
            low_fidelity_available=False,
        )
        gate = evaluate_preview_gate(
            mode=PreviewMode.ISSUE,
            report_bytes=b"docx",
            result_present=True,
            modified_docx_bytes=None,
            current_signature="v1",
            latest_signature="v1",
            environment=unavailable,
        )

        self.assertFalse(gate.enabled)
        self.assertEqual(gate.reason_key, "preview_route_unavailable")

    def test_matching_versions_enable_preview(self):
        gate = evaluate_preview_gate(
            mode=PreviewMode.COMPARISON,
            report_bytes=b"docx",
            result_present=True,
            modified_docx_bytes=b"changed",
            current_signature="v1",
            latest_signature="v1",
            environment=self.available,
        )

        self.assertTrue(gate.enabled)
        self.assertEqual(gate.reason_key, "preview_ready")


class PreviewEnvironmentDetectionTest(unittest.TestCase):
    def test_poppler_lookup_prefers_a_native_executable_over_a_cmd_shim(self):
        original_which = preview_service_module.shutil.which

        def fake_which(command):
            return {
                "pdftoppm": r"C:\\runtime\\pdftoppm.cmd",
                "pdftoppm.exe": r"C:\\TeXLive\\pdftoppm.exe",
            }.get(command)

        preview_service_module.shutil.which = fake_which
        try:
            executable = preview_service_module._find_poppler_executable("pdftoppm")
        finally:
            preview_service_module.shutil.which = original_which

        self.assertEqual(executable, r"C:\\TeXLive\\pdftoppm.exe")

    def test_detect_preview_environment_caches_expensive_probes(self):
        preview_service_module.detect_preview_environment.cache_clear()
        calls = {"word": 0, "libreoffice": 0, "low_fidelity": 0, "which": 0}

        def fake_word_capability():
            calls["word"] += 1
            return preview_service_module._WordCapability(True, "usable", path="WINWORD.EXE")

        def fake_find_libreoffice():
            calls["libreoffice"] += 1
            return None

        def fake_low_fidelity_available():
            calls["low_fidelity"] += 1
            return True

        def fake_which(command):
            calls["which"] += 1
            return "pdftoppm" if command == "pdftoppm" else None

        original_word = preview_service_module._detect_word_capability
        original_libreoffice = preview_service_module._find_libreoffice
        original_low_fidelity = preview_service_module._low_fidelity_available
        original_which = preview_service_module.shutil.which
        preview_service_module._detect_word_capability = fake_word_capability
        preview_service_module._find_libreoffice = fake_find_libreoffice
        preview_service_module._low_fidelity_available = fake_low_fidelity_available
        preview_service_module.shutil.which = fake_which
        try:
            first = preview_service_module.detect_preview_environment()
            second = preview_service_module.detect_preview_environment()
        finally:
            preview_service_module._detect_word_capability = original_word
            preview_service_module._find_libreoffice = original_libreoffice
            preview_service_module._low_fidelity_available = original_low_fidelity
            preview_service_module.shutil.which = original_which
            preview_service_module.detect_preview_environment.cache_clear()

        self.assertIs(first, second)
        self.assertEqual(first.converter_label, "Microsoft Word")
        self.assertEqual(first.renderer_label, "pdftoppm")
        self.assertEqual(calls, {"word": 1, "libreoffice": 1, "low_fidelity": 1, "which": 4})


class PreviewVersionTest(unittest.TestCase):
    def test_document_version_changes_when_modified_docx_changes(self):
        first = build_document_version(
            input_signature="same",
            original_docx_bytes=b"original",
            modified_docx_bytes=b"modified-a",
        )
        second = build_document_version(
            input_signature="same",
            original_docx_bytes=b"original",
            modified_docx_bytes=b"modified-b",
        )

        self.assertNotEqual(first, second)
        self.assertEqual(build_preview_version(first, PreviewMode.COMPARISON), build_preview_version(first, "comparison"))

    def test_selected_page_numbers_preload_only_current_and_adjacent_pages(self):
        self.assertEqual(preview_service_module._selected_page_numbers(30, 15, 1), [14, 15, 16])
        self.assertEqual(preview_service_module._selected_page_numbers(30, 1, 1), [1, 2])
        self.assertEqual(preview_service_module._selected_page_numbers(30, 30, 1), [29, 30])
        self.assertEqual(preview_service_module._selected_page_numbers(30, 15, 0), [15])
        self.assertEqual(preview_service_module._selected_page_numbers(3, None, 1), [1, 2, 3])


class PreviewMappingTest(unittest.TestCase):
    def test_paragraph_indices_map_to_bounded_pages(self):
        self.assertEqual(estimate_page_number(None, max_index=100, page_count=8), 1)
        self.assertEqual(estimate_page_number(0, max_index=100, page_count=8), 1)
        self.assertEqual(estimate_page_number(100, max_index=100, page_count=8), 8)

    def test_issue_items_keep_page_filterable_metadata(self):
        issues = [
            DummyIssue("error", "equation", "Missing equation number.", 0),
            DummyIssue("warning", "table", "Possible table gap.", 90),
        ]

        items = build_issue_items(issues, page_count=3)

        self.assertEqual(items[0].item_id, "issue-1")
        self.assertEqual(items[0].page_number, 1)
        self.assertEqual(items[1].page_number, 1)
        self.assertIsNone(items[1].rect)
        self.assertEqual(items[1].location_status, "unlocated")

    def test_change_items_keep_before_after_for_review_list(self):
        actions = [DummyAction("table", 10, "Table 2", "TABLE II", "Renumbered table.")]

        items = build_change_items(actions, page_count=2)

        self.assertEqual(items[0].item_id, "change-1")
        self.assertEqual(items[0].before, "Table 2")
        self.assertEqual(items[0].after, "TABLE II")
        self.assertEqual(items[0].severity, "change")
        self.assertIsNone(items[0].rect)

    def test_bbox_layout_locates_one_column_line_and_rejects_cross_column_guess(self):
        xml = b'''<html><body><doc><page width="600" height="800">
        <flow><block><line xMin="320" yMin="200" xMax="560" yMax="214">
        <word xMin="320" yMin="200" xMax="380" yMax="214">TABLE</word>
        <word xMin="390" yMin="200" xMax="405" yMax="214">II</word>
        </line><line xMin="40" yMin="250" xMax="560" yMax="264">
        <word xMin="40" yMin="250" xMax="560" yMax="264">WIDE HEADING</word>
        </line></block></flow></page></doc></body></html>'''
        lines = parse_pdftotext_bbox_layout(xml)
        table = DummyAction("table", 10, "Table 2", "TABLE II", "Renumbered")
        located = locate_action_in_pdf(table, lines)
        self.assertEqual(located[0], 1)
        self.assertLess(located[1][2], 0.48)
        wide = DummyAction("heading", 11, "Old", "WIDE HEADING", "Heading")
        self.assertIsNone(locate_action_in_pdf(wide, lines))
        wide.is_cross_column = True
        self.assertIsNotNone(locate_action_in_pdf(wide, lines))

    def test_cropbox_rotation_transform_uses_top_left_normalized_space(self):
        normal = transform_pdf_bbox((20, 30, 60, 50), (10, 20, 110, 220), 0)
        self.assertEqual(normal, (0.1, 0.85, 0.4, 0.1))
        rotated = transform_pdf_bbox((20, 30, 60, 50), (10, 20, 110, 220), 90)
        self.assertTrue(all(0 <= value <= 1 for value in rotated))

    def test_unlocated_item_never_renders_marker(self):
        page = RenderedPage(page_number=1, image_bytes=b"png")
        item = build_change_items([DummyAction("table", 1, "a", "b", "c")], 1)[0]
        html = render_page_html(page, [item])
        self.assertNotIn("preview-marker preview-marker-change", html)

    def test_pixel_diff_items_use_rendered_page_pixels(self):
        try:
            from PIL import Image, ImageDraw
        except Exception as exc:  # pragma: no cover - depends on local optional dependency.
            self.skipTest(f"Pillow unavailable: {exc}")

        original = Image.new("RGB", (200, 100), "white")
        modified = Image.new("RGB", (200, 100), "white")
        draw = ImageDraw.Draw(modified)
        draw.rectangle((95, 45, 105, 52), fill="black")
        original_bytes = io.BytesIO()
        modified_bytes = io.BytesIO()
        original.save(original_bytes, format="PNG")
        modified.save(modified_bytes, format="PNG")

        items = build_pixel_diff_items(
            [RenderedPage(page_number=8, image_bytes=original_bytes.getvalue(), width=200, height=100)],
            [RenderedPage(page_number=8, image_bytes=modified_bytes.getvalue(), width=200, height=100)],
        )

        self.assertEqual(len(items), 1)
        self.assertEqual(items[0].page_number, 8)
        self.assertEqual(items[0].location_status, "located")
        self.assertEqual(items[0].coordinate_space, "image_normalized_top_left")
        x, y, width, height = items[0].rect
        self.assertGreater(x, 0.35)
        self.assertLess(x + width, 0.7)
        self.assertGreater(y, 0.25)
        self.assertLess(y + height, 0.75)

    def test_pixel_diff_suppresses_unreliable_large_or_many_markers(self):
        try:
            from PIL import Image, ImageDraw
        except Exception as exc:  # pragma: no cover - depends on local optional dependency.
            self.skipTest(f"Pillow unavailable: {exc}")

        original = Image.new("RGB", (200, 100), "white")
        modified = Image.new("RGB", (200, 100), "white")
        draw = ImageDraw.Draw(modified)
        draw.rectangle((10, 10, 190, 90), fill="black")
        original_bytes = io.BytesIO()
        modified_bytes = io.BytesIO()
        original.save(original_bytes, format="PNG")
        modified.save(modified_bytes, format="PNG")
        diagnostics = {}

        items = build_pixel_diff_items(
            [RenderedPage(page_number=9, image_bytes=original_bytes.getvalue(), width=200, height=100)],
            [RenderedPage(page_number=9, image_bytes=modified_bytes.getvalue(), width=200, height=100)],
            diagnostics=diagnostics,
        )

        self.assertEqual(items, [])
        self.assertIn(diagnostics[9]["status"], {"suppressed_total_area_too_large", "located"})
        if diagnostics[9]["status"] == "located":
            self.assertEqual(diagnostics[9]["final_count"], 0)


class PreviewRenderingFallbackTest(unittest.TestCase):
    def test_low_fidelity_docx_preview_renders_png_page(self):
        try:
            import PIL  # noqa: F401
        except Exception as exc:  # pragma: no cover - depends on local optional dependency.
            self.skipTest(f"Pillow unavailable: {exc}")

        document = Document()
        document.add_paragraph("This is a browser fallback preview paragraph.")
        output = io.BytesIO()
        document.save(output)
        environment = PreviewEnvironment(
            converter_available=False,
            renderer_available=False,
            converter_label="DOCX browser preview",
            renderer_label="Pillow text renderer",
            low_fidelity_available=True,
        )

        rendered = render_docx_to_pages(
            output.getvalue(),
            "fallback.docx",
            label="original",
            dpi=120,
            max_pages=2,
            environment=environment,
        )

        self.assertEqual(rendered.converter_label, "DOCX browser preview")
        self.assertEqual(rendered.page_count, 1)
        self.assertTrue(rendered.rendered_pages[0].image_bytes.startswith(b"\x89PNG"))

    def test_safe_file_name_preserves_pdf_for_direct_preview(self):
        self.assertEqual(safe_file_name("final report.pdf"), "final_report.pdf")

    def test_docx_to_pdf_cache_survives_page_change(self):
        document = Document()
        document.add_paragraph("Preview cache smoke test.")
        output = io.BytesIO()
        document.save(output)
        docx_bytes = output.getvalue()
        environment = PreviewEnvironment(
            converter_available=True,
            renderer_available=True,
            converter_label="Microsoft Word",
            renderer_label="pdftoppm",
            pdftoppm_path="pdftoppm",
            word_available=True,
            word_status="usable",
        )
        fake_pdf = b"%PDF-1.7\n" + (b"0" * 700)
        calls = {"convert": 0, "pages": []}

        def fake_convert(input_path, working_dir, environment):
            calls["convert"] += 1
            return fake_pdf

        def fake_render(pdf_bytes, working_dir, *, dpi, max_pages, page_number, preload_radius, environment, timings=None, label="document"):
            calls["pages"].append(page_number)
            return ([RenderedPage(page_number=page_number or 1, image_bytes=b"png", width=1, height=1)], 3)

        original_convert = preview_service_module.convert_docx_to_pdf
        original_render = preview_service_module.render_pdf_bytes_to_png_pages
        original_cache_dir = os.environ.get("REPORT_FORMATTER_PREVIEW_CACHE_DIR")
        with tempfile.TemporaryDirectory() as cache_dir:
            os.environ["REPORT_FORMATTER_PREVIEW_CACHE_DIR"] = cache_dir
            preview_service_module.convert_docx_to_pdf = fake_convert
            preview_service_module.render_pdf_bytes_to_png_pages = fake_render
            try:
                first = render_docx_to_pages(
                    docx_bytes,
                    "cache.docx",
                    label="original",
                    dpi=120,
                    max_pages=3,
                    page_number=1,
                    preload_radius=0,
                    environment=environment,
                )
                second = render_docx_to_pages(
                    docx_bytes,
                    "cache.docx",
                    label="original",
                    dpi=120,
                    max_pages=3,
                    page_number=2,
                    preload_radius=0,
                    environment=environment,
                )
            finally:
                preview_service_module.convert_docx_to_pdf = original_convert
                preview_service_module.render_pdf_bytes_to_png_pages = original_render
                if original_cache_dir is None:
                    os.environ.pop("REPORT_FORMATTER_PREVIEW_CACHE_DIR", None)
                else:
                    os.environ["REPORT_FORMATTER_PREVIEW_CACHE_DIR"] = original_cache_dir

        self.assertEqual(calls["convert"], 1)
        self.assertEqual(calls["pages"], [1, 2])
        self.assertEqual(first.page_count, 3)
        self.assertEqual(second.rendered_pages[0].page_number, 2)


if __name__ == "__main__":
    unittest.main()
