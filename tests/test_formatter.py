from __future__ import annotations

import base64
import io
from pathlib import Path
import unittest

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from formatter import image_table_detector
from formatter.equation_detector import match_equation_number
from formatter.formatter_engine import FormattingOptions, ProcessingMode, process_document
from formatter.equation_utils import (
    get_visible_text_from_xml,
    protect_ieee_citation_tokens,
    restore_ieee_citation_tokens,
)
from formatter.language_detector import detect_report_language


def _tiny_png_stream() -> io.BytesIO:
    png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="
    )
    return io.BytesIO(png)


def _add_omml_equation(paragraph, text: str = "V=IR") -> None:
    paragraph._p.append(
        parse_xml(
            f'<m:oMathPara {nsdecls("m")}><m:oMath><m:r><m:t>{text}</m:t></m:r></m:oMath></m:oMathPara>'
        )
    )


class FormatterEngineTests(unittest.TestCase):
    def test_ieee_profile_renumbers_figures_tables_but_suspends_equation_writes(self):
        doc = Document()
        doc.add_paragraph("Abstract")
        doc.add_paragraph(
            "This lab report uses Fig. 1, Table I, and Eq. (1) in the main discussion section."
        )
        doc.add_paragraph("Figure 3 Sensor output before calibration")
        doc.add_paragraph("Table 7")
        doc.add_paragraph("SENSOR SUMMARY")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "Sensor"
        doc.add_paragraph("y = mx + b     (4)")

        result = process_document(
            doc,
            mode=ProcessingMode.FIX_NUMBERING_AND_FORMATTING,
            options=FormattingOptions(format_profile_name="ieee"),
        )
        texts = [paragraph.text for paragraph in doc.paragraphs]

        self.assertIn("Fig. 1. Sensor output before calibration", texts)
        self.assertIn("TABLE I", texts)
        self.assertIn("y = mx + b     (4)", texts)
        self.assertTrue(any(action.kind == "figure" for action in result.actions))
        self.assertTrue(any(action.kind == "table" for action in result.actions))
        self.assertFalse([action for action in result.actions if action.kind == "equation"])
        self.assertTrue(any("automatic equation renumbering is suspended" in item for item in result.manual_review))

    def test_equation_renumbering_distinguishes_labels_from_references(self):
        doc = Document()
        doc.add_paragraph("Abstract")
        first_equation = doc.add_paragraph()
        first_equation.add_run("V = IR")
        first_equation.add_run("                                                        ")
        first_equation.add_run("(1)")
        doc.add_paragraph(
            "Equation (1) defines the voltage-current relationship. "
            "The next equation deliberately skips numbering."
        )
        second_equation = doc.add_paragraph()
        second_equation.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        second_equation.add_run("P = VI")
        second_equation.add_run("\t................................ ")
        second_equation.add_run("(")
        second_equation.add_run("3")
        second_equation.add_run(")")
        doc.add_paragraph("Eq. (3) is referenced here. Equation (4) is referenced here but does not exist.")

        result = process_document(
            doc,
            mode=ProcessingMode.FIX_NUMBERING_AND_FORMATTING,
            options=FormattingOptions(format_profile_name="ieee"),
        )
        texts = [paragraph.text for paragraph in doc.paragraphs]
        messages = "\n".join(issue.message for issue in result.issues)

        self.assertIn("V = IR                                                        (1)", texts)
        self.assertIn("P = VI\t................................ (3)", texts)
        self.assertIn(
            "Equation (1) defines the voltage-current relationship. "
            "The next equation deliberately skips numbering.",
            texts,
        )
        self.assertIn("Eq. (3) is referenced here. Equation (4) is referenced here but does not exist.", texts)
        self.assertNotIn("(2) defines the voltage-current relationship.", "\n".join(texts))
        self.assertEqual(second_equation.alignment, WD_ALIGN_PARAGRAPH.RIGHT)
        self.assertEqual(second_equation.runs[0].text, "P = VI")
        self.assertEqual(second_equation.runs[1].text, "\t................................ ")
        self.assertEqual("".join(run.text for run in second_equation.runs[2:]), "(3)")
        self.assertEqual([equation.number for equation in result.equations], [1, 3])
        self.assertNotIn("Eq. (3)", messages)
        self.assertIn("Missing equation reference: Equation (4)", messages)
        self.assertFalse([action for action in result.actions if action.kind == "equation"])

    def test_existing_equation_number_is_idempotent(self):
        doc = Document()
        doc.add_paragraph("Abstract")
        doc.add_paragraph("V = IR     (1)")

        process_document(
            doc,
            mode=ProcessingMode.FIX_NUMBERING_AND_FORMATTING,
            options=FormattingOptions(format_profile_name="ieee"),
        )
        process_document(
            doc,
            mode=ProcessingMode.FIX_NUMBERING_AND_FORMATTING,
            options=FormattingOptions(format_profile_name="ieee"),
        )

        self.assertEqual(doc.paragraphs[-1].text, "V = IR     (1)")
        self.assertNotIn("(1) (1)", doc.paragraphs[-1].text)

    def test_unnumbered_equation_is_reported_but_not_numbered(self):
        doc = Document()
        doc.add_paragraph("Abstract")
        equation = doc.add_paragraph("V = IR")
        equation.alignment = WD_ALIGN_PARAGRAPH.CENTER

        process_document(
            doc,
            mode=ProcessingMode.FIX_NUMBERING_AND_FORMATTING,
            options=FormattingOptions(format_profile_name="ieee"),
        )
        process_document(
            doc,
            mode=ProcessingMode.FIX_NUMBERING_AND_FORMATTING,
            options=FormattingOptions(format_profile_name="ieee"),
        )

        self.assertEqual(doc.paragraphs[-1].text, "V = IR")
        self.assertEqual(get_visible_text_from_xml(doc.paragraphs[-1]), "V = IR")
        self.assertNotIn("(1) (1)", doc.paragraphs[-1].text)

    def test_wrong_equation_number_is_preserved_while_auto_edit_is_suspended(self):
        doc = Document()
        doc.add_paragraph("Abstract")
        doc.add_paragraph("a = b     (1)")
        doc.add_paragraph("c = d     (2)")
        doc.add_paragraph("V = IR     (7)")

        process_document(
            doc,
            mode=ProcessingMode.FIX_NUMBERING_AND_FORMATTING,
            options=FormattingOptions(format_profile_name="ieee"),
        )

        self.assertEqual(doc.paragraphs[-1].text, "V = IR     (7)")
        self.assertNotIn("(7) (3)", doc.paragraphs[-1].text)

    def test_duplicate_equation_numbers_are_not_rewritten_while_auto_edit_is_suspended(self):
        doc = Document()
        doc.add_paragraph("Abstract")
        doc.add_paragraph("V = IR     (1) (1)")

        result = process_document(
            doc,
            mode=ProcessingMode.FIX_NUMBERING_AND_FORMATTING,
            options=FormattingOptions(format_profile_name="ieee"),
        )

        self.assertEqual(doc.paragraphs[-1].text, "V = IR     (1) (1)")
        self.assertFalse([action for action in result.actions if action.kind == "equation"])
        self.assertTrue(any("automatic equation renumbering is suspended" in item for item in result.manual_review))

    def test_omml_equation_with_existing_number_is_not_renumbered_twice(self):
        doc = Document()
        doc.add_paragraph("Abstract")
        equation = doc.add_paragraph()
        _add_omml_equation(equation)
        equation.add_run("\t")
        equation.add_run("(1)")

        process_document(
            doc,
            mode=ProcessingMode.FIX_NUMBERING_AND_FORMATTING,
            options=FormattingOptions(format_profile_name="ieee"),
        )
        process_document(
            doc,
            mode=ProcessingMode.FIX_NUMBERING_AND_FORMATTING,
            options=FormattingOptions(format_profile_name="ieee"),
        )

        self.assertEqual(equation.text, "\t(1)")
        self.assertNotIn("(1) (1)", equation.text)

    def test_table_row_equation_number_is_preserved_while_auto_edit_is_suspended(self):
        doc = Document()
        doc.add_paragraph("Abstract")
        table = doc.add_table(rows=1, cols=2)
        _add_omml_equation(table.cell(0, 0).paragraphs[0])
        table.cell(0, 1).paragraphs[0].add_run("(7)")

        process_document(
            doc,
            mode=ProcessingMode.FIX_NUMBERING_AND_FORMATTING,
            options=FormattingOptions(format_profile_name="ieee"),
        )
        second = process_document(
            doc,
            mode=ProcessingMode.FIX_NUMBERING_AND_FORMATTING,
            options=FormattingOptions(format_profile_name="ieee"),
        )

        self.assertEqual(table.cell(0, 1).text, "(7)")
        self.assertEqual([equation.number for equation in second.equations], [7])
        self.assertFalse(second.actions)

    def test_ieee_citations_in_captions_are_not_equation_numbers(self):
        doc = Document()
        doc.add_paragraph("Abstract")
        doc.add_paragraph("Fig. 4. Equivalent Circuit Modeling [1]")
        centered_caption = doc.add_paragraph("Fig. 5. Phasor diagram for the open-circuit test [1]")
        centered_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph("The equivalent model is as shown in [1, Fig. 2].")
        doc.add_paragraph("V = IR     (1)")

        result = process_document(
            doc,
            mode=ProcessingMode.FIX_NUMBERING_AND_FORMATTING,
            options=FormattingOptions(
                format_profile_name="ieee",
                renumber_figures=False,
                renumber_tables=False,
            ),
        )
        texts = [paragraph.text for paragraph in doc.paragraphs]

        self.assertIn("Fig. 4. Equivalent Circuit Modeling [1]", texts)
        self.assertIn("Fig. 5. Phasor diagram for the open-circuit test [1]", texts)
        self.assertIn("The equivalent model is as shown in [1, Fig. 2].", texts)
        self.assertIn("V = IR     (1)", texts)
        self.assertNotIn("Fig. 4. Equivalent Circuit Modeling (1)", texts)
        self.assertNotIn("Fig. 5. Phasor diagram for the open-circuit test (2)", texts)
        self.assertEqual([equation.number for equation in result.equations], [1])

    def test_requested_ieee_classification_contract(self):
        doc = Document()
        doc.add_paragraph("Abstract")
        doc.add_paragraph("Fig. 4. Equivalent Circuit Modeling [1]")
        doc.add_paragraph("Fig. 5. Phasor diagram for the open-circuit test [1]")
        doc.add_paragraph("Table 1: Measured Secondary Voltage and Current Data")
        equation = doc.add_paragraph("V = IR")
        equation.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph("The method is shown in [1]. Prior work [1]-[3] should remain citations.")

        result = process_document(
            doc,
            mode=ProcessingMode.FIX_NUMBERING_AND_FORMATTING,
            options=FormattingOptions(
                format_profile_name="ieee",
                renumber_figures=False,
                renumber_tables=False,
            ),
        )
        texts = [paragraph.text for paragraph in doc.paragraphs]

        self.assertIn("Fig. 4. Equivalent Circuit Modeling [1]", texts)
        self.assertIn("Fig. 5. Phasor diagram for the open-circuit test [1]", texts)
        self.assertIn("Table 1: Measured Secondary Voltage and Current Data", texts)
        self.assertIn("The method is shown in [1]. Prior work [1]-[3] should remain citations.", texts)
        self.assertIn("V = IR", texts)
        self.assertEqual([figure.number for figure in result.figures], [4, 5])
        self.assertTrue(any(table.number == 1 and table.kind == "table" for table in result.tables))
        self.assertEqual([equation.number for equation in result.equations], [])
        self.assertEqual(len(result.unnumbered_equations), 1)
        self.assertEqual(
            [citation["text"] for citation in result.reference_citations],
            ["[1]", "[1]", "[1]", "[1]-[3]"],
        )
        self.assertIsNone(match_equation_number("The method is shown in [1]."))
        self.assertIsNone(match_equation_number("Fig. 4. Equivalent Circuit Modeling [1]"))
        self.assertIsNotNone(match_equation_number("x = y + z     (A1)"))
        self.assertIsNotNone(match_equation_number("x = y + z     (1a)"))

    def test_image_table_caption_resolves_table_reference(self):
        doc = Document()
        doc.add_paragraph("Abstract")
        doc.add_paragraph("Table 2")
        doc.add_paragraph("Measured motor speed results")
        doc.add_picture(_tiny_png_stream())
        doc.add_paragraph("The measured data are summarized in Table 2.")

        result = process_document(
            doc,
            mode=ProcessingMode.ANALYZE_ONLY,
            options=FormattingOptions(format_profile_name="ieee"),
        )
        messages = "\n".join(issue.message for issue in result.issues)

        self.assertTrue(any(table.number == 2 and table.source == "image_table_candidate" for table in result.tables))
        self.assertNotIn("Table 2 points to a nonexistent table", messages)

    def test_embedded_table_caption_in_image_is_reported_with_ocr_title(self):
        doc = Document()
        doc.add_paragraph("Abstract")
        doc.add_picture(_tiny_png_stream())

        original_ocr = image_table_detector._ocr_image_blob
        original_structure = image_table_detector._detect_table_structure

        def fake_ocr(_blob):
            return image_table_detector.OcrResult(
                True,
                "stub",
                "Table 1: Measured Secondary Voltage and Current Data\nLoad Current Voltage Power",
            )

        def fake_structure(_blob):
            return image_table_detector.TableStructureResult(True, 0.91, 4, 4, "stub")

        image_table_detector._ocr_image_blob = fake_ocr
        image_table_detector._detect_table_structure = fake_structure
        try:
            result = process_document(
                doc,
                mode=ProcessingMode.ANALYZE_ONLY,
                options=FormattingOptions(format_profile_name="ieee"),
            )
        finally:
            image_table_detector._ocr_image_blob = original_ocr
            image_table_detector._detect_table_structure = original_structure

        self.assertEqual(len(result.images), 1)
        image = result.images[0]
        self.assertTrue(image.is_suspected_table)
        self.assertEqual(image.caption_type, "table_caption_embedded_in_image")
        self.assertEqual(image.ocr_title_text, "Table 1: Measured Secondary Voltage and Current Data")
        self.assertTrue(any(table.number == 1 and table.source == "image_table_confirmed" for table in result.tables))
        self.assertIn("Suspected image table: image 1", result.report_text)
        self.assertIn("Table 1: Measured Secondary Voltage and Current Data", result.report_text)
        self.assertIn("Caption is embedded in the image", result.report_text)

    def test_nearby_image_table_candidate_is_warning_not_confirmed_missing(self):
        doc = Document()
        doc.add_paragraph("Abstract")
        doc.add_paragraph("The measured data are summarized in Table 2.")
        doc.add_picture(_tiny_png_stream())

        result = process_document(
            doc,
            mode=ProcessingMode.ANALYZE_ONLY,
            options=FormattingOptions(format_profile_name="ieee"),
        )
        table_issues = [issue for issue in result.issues if issue.category == "table"]
        messages = "\n".join(issue.message for issue in table_issues)

        self.assertTrue(any(table.number == 2 and table.source == "image_table_candidate" for table in result.tables))
        self.assertTrue(any(issue.severity == "warning" and "possible image table" in issue.message for issue in table_issues))
        self.assertNotIn("Table 2 points to a nonexistent table", messages)

    def test_reference_citations_are_counted_separately_from_equations(self):
        doc = Document()
        doc.add_paragraph("Abstract")
        doc.add_paragraph("Fig. 4. Equivalent Circuit Modeling [1]")
        doc.add_paragraph("The circuit is adapted from [1, Fig. 2].")
        doc.add_paragraph("V = IR     (1)")

        result = process_document(
            doc,
            mode=ProcessingMode.ANALYZE_ONLY,
            options=FormattingOptions(format_profile_name="ieee"),
        )

        self.assertEqual([equation.number for equation in result.equations], [1])
        self.assertEqual([citation["text"] for citation in result.reference_citations], ["[1]", "[1, Fig. 2]"])

    def test_ieee_citation_ranges_are_tokenized_and_restored(self):
        text = "Prior work [1], [1, eq. (8)], and [3]\u2013[5] should remain citations."
        protected, tokens = protect_ieee_citation_tokens(text)

        self.assertNotIn("[1]", protected)
        self.assertNotIn("[3]\u2013[5]", protected)
        self.assertEqual(restore_ieee_citation_tokens(protected, tokens), text)

    def test_current_report_fig_captions_and_image_tables_do_not_cross_into_equations(self):
        doc = Document()
        doc.add_paragraph("Abstract")
        doc.add_paragraph("Table 1 systematically delineates the divergence between these theoretical extrapolations and the empirical measurements.")
        doc.add_paragraph("Fig. 4. Equivalent Circuit Modeling [1]")
        doc.add_paragraph("Fig. 5. Phasor diagram for the open-circuit test [1]")
        doc.add_paragraph("Table 2.")
        doc.add_picture(_tiny_png_stream())
        doc.add_paragraph("Table 3.")
        doc.add_picture(_tiny_png_stream())
        doc.add_paragraph("The results are summarized in Table 2 and Table 3.")
        equation = doc.add_paragraph("V = IR")
        equation.alignment = WD_ALIGN_PARAGRAPH.CENTER

        result = process_document(
            doc,
            mode=ProcessingMode.FIX_NUMBERING_AND_FORMATTING,
            options=FormattingOptions(
                format_profile_name="ieee",
                renumber_figures=False,
                renumber_tables=False,
            ),
        )
        texts = [paragraph.text for paragraph in doc.paragraphs]
        messages = "\n".join(issue.message for issue in result.issues)

        self.assertIn("Fig. 4. Equivalent Circuit Modeling [1]", texts)
        self.assertIn("Fig. 5. Phasor diagram for the open-circuit test [1]", texts)
        self.assertIn(
            "Table 1 systematically delineates the divergence between these theoretical extrapolations and the empirical measurements.",
            texts,
        )
        self.assertIn("V = IR", texts)
        self.assertEqual([equation.number for equation in result.equations], [])
        self.assertEqual(len(result.unnumbered_equations), 1)
        self.assertFalse(
            any(
                table.number == 1
                and table.source in {"native_table", "caption_only_table"}
                and "systematically delineates" in table.raw_text
                for table in result.tables
            )
        )
        self.assertTrue(any(table.number == 2 and table.source == "image_table_candidate" for table in result.tables))
        self.assertTrue(any(table.number == 3 and table.source == "image_table_candidate" for table in result.tables))
        self.assertNotIn("Fig. 4. Equivalent Circuit Modeling (1)", "\n".join(texts))
        self.assertNotIn("Fig. 5. Phasor diagram for the open-circuit test (2)", "\n".join(texts))
        self.assertNotIn("Table 2 points to a nonexistent table", messages)
        self.assertNotIn("Table 3 points to a nonexistent table", messages)

    def test_toc_equation_numbers_do_not_create_duplicates(self):
        doc = Document()
        doc.add_paragraph("Contents")
        doc.add_paragraph("Equivalent Circuit Modeling\t........\t(1)")
        doc.add_paragraph("Phasor diagram\t........\t(1)")
        doc.add_paragraph("Abstract")
        doc.add_paragraph("V = IR     (1)")

        result = process_document(
            doc,
            mode=ProcessingMode.ANALYZE_ONLY,
            options=FormattingOptions(format_profile_name="ieee"),
        )
        messages = "\n".join(issue.message for issue in result.issues)

        self.assertEqual([equation.number for equation in result.equations], [1])
        self.assertGreaterEqual(result.skipped_toc_entries_count, 3)
        self.assertNotIn("Duplicate equation number 1", messages)

    def test_toc_table_reference_is_not_body_missing_table_error(self):
        doc = Document()
        doc.add_paragraph("Table of Contents")
        doc.add_paragraph("Table 2\t........\t5")
        doc.add_paragraph("Abstract")
        doc.add_paragraph("The report discusses measurements without citing that table.")

        result = process_document(
            doc,
            mode=ProcessingMode.ANALYZE_ONLY,
            options=FormattingOptions(format_profile_name="ieee"),
        )
        messages = "\n".join(issue.message for issue in result.issues)

        self.assertNotIn("Table 2 points to a nonexistent table", messages)

    def test_missing_body_table_reference_still_reports_error(self):
        doc = Document()
        doc.add_paragraph("Abstract")
        doc.add_paragraph("The missing measurements are shown in Table 9.")

        result = process_document(
            doc,
            mode=ProcessingMode.ANALYZE_ONLY,
            options=FormattingOptions(format_profile_name="ieee"),
        )
        messages = "\n".join(issue.message for issue in result.issues)

        self.assertIn("Table 9 points to a nonexistent table", messages)

    def test_fix_mode_reports_unnumbered_display_equation_without_inserting_number(self):
        doc = Document()
        doc.add_paragraph("Abstract")
        equation = doc.add_paragraph("V = IR")
        equation.alignment = WD_ALIGN_PARAGRAPH.CENTER

        result = process_document(
            doc,
            mode=ProcessingMode.FIX_NUMBERING_AND_FORMATTING,
            options=FormattingOptions(format_profile_name="ieee"),
        )

        self.assertIn("V = IR", [paragraph.text for paragraph in doc.paragraphs])
        self.assertEqual([equation.number for equation in result.equations], [])
        self.assertEqual(len(result.unnumbered_equations), 1)
        self.assertFalse([action for action in result.actions if action.kind == "equation"])

    def test_downloaded_ieee_broken_fixture_equation_regression_when_available(self):
        fixture = Path(__file__).resolve().parent / "fixtures" / "optional_external_report.docx"
        if not fixture.exists():
            self.skipTest(f"{fixture} is not available on this machine")

        doc = Document(fixture)
        result = process_document(
            doc,
            mode=ProcessingMode.FIX_NUMBERING_AND_FORMATTING,
            options=FormattingOptions(format_profile_name="ieee"),
        )
        texts = [paragraph.text for paragraph in doc.paragraphs]
        messages = "\n".join(issue.message for issue in result.issues)

        self.assertIn("Fig. 1. PWM input waveform at 1.5 ms pulse width.", texts)
        self.assertIn("Fig. 2. Motor speed response after step input", texts)
        self.assertIn("Fig. 3. Duplicate number for controller block diagram.", texts)
        self.assertNotIn("Figure 3 Motor speed response after step input", texts)
        self.assertIn("TABLE I", texts)
        self.assertIn("MOTOR TEST PARAMETERS", texts)
        self.assertIn("Table I is correctly referenced here.", texts)
        self.assertIn("TABLE II", texts)
        self.assertIn("MEASURED MOTOR SPEED RESULTS", texts)
        self.assertNotIn("Table 3", texts)
        self.assertIn("V = IR                                                        (1)", texts)
        self.assertIn("P = VI                                                        (3)", texts)
        self.assertIn(
            "Equation (1) defines the voltage-current relationship. "
            "The next equation deliberately skips numbering.",
            texts,
        )
        self.assertIn("Eq. (3) is referenced here. Equation (4) is referenced here but does not exist.", texts)
        self.assertNotIn("(2) defines the voltage-current relationship.", "\n".join(texts))
        self.assertNotIn("P = VI                                                        (2)", texts)
        self.assertIn("Missing equation reference: Equation (4)", messages)
        self.assertIn("Table 5 points to a nonexistent table", messages)

    def test_template_matching_preserves_detected_figure_prefix(self):
        doc = Document()
        doc.add_paragraph("Abstract")
        doc.add_paragraph("This report uses Figure 1 and Table 1 in the discussion.")
        doc.add_paragraph("Figure 7 Sensor output before calibration")
        doc.add_paragraph("Table 4")
        doc.add_paragraph("Measured values")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "Measured"

        process_document(
            doc,
            mode=ProcessingMode.FIX_NUMBERING_AND_FORMATTING,
            options=FormattingOptions(format_profile_name="template_matching"),
        )
        texts = [paragraph.text for paragraph in doc.paragraphs]

        self.assertIn("Figure 1. Sensor output before calibration", texts)
        self.assertIn("Table 1", texts)
        self.assertIn("Measured values", texts)

    def test_reference_checker_reports_nonexistent_and_inconsistent_refs(self):
        doc = Document()
        doc.add_paragraph("Abstract")
        doc.add_paragraph(
            "The setup is shown in Figure 1. The failed result is shown in Fig. 5. "
            "Table 2 is cited, and Eq. (3) is also cited."
        )
        doc.add_paragraph("Fig. 1. Experimental setup")
        doc.add_paragraph("TABLE I")
        doc.add_paragraph("MEASURED VALUES")
        doc.add_paragraph("a = b + c     (1)")

        result = process_document(doc, mode=ProcessingMode.ANALYZE_ONLY)
        messages = "\n".join(issue.message for issue in result.issues)

        self.assertIn("Fig. 5 points to a nonexistent figure", messages)
        self.assertIn("Table 2 points to a nonexistent table", messages)
        self.assertIn("Missing equation reference: Eq. (3)", messages)
        self.assertIn("Figure 1 is inconsistent", messages)

    def test_chinese_language_detection_and_chinese_report_messages(self):
        doc = Document()
        doc.add_paragraph("摘要")
        doc.add_paragraph("本文展示图 1 和表 1 的实验结果，并讨论公式 (1)。")
        doc.add_paragraph("图 3. 原始波形")
        doc.add_paragraph("表 2. 测量结果")
        doc.add_paragraph("y = mx + b     (2)")

        detection = detect_report_language(doc)
        result = process_document(
            doc,
            mode=ProcessingMode.FIX_NUMBERING_AND_FORMATTING,
            options=FormattingOptions(ui_language="zh", report_language_override="zh"),
        )

        self.assertEqual(detection["language"], "zh")
        self.assertEqual(result.report_language, "zh")
        self.assertIn("模板感知实验报告格式化工具检查报告", result.report_text)
        self.assertIn("已选择界面语言：中文", result.report_text)

    def test_preserves_cover_page_when_start_detected(self):
        doc = Document()
        doc.add_paragraph("Fig. 9. Cover page illustration")
        doc.add_paragraph("Student name and course information")
        doc.add_page_break()
        doc.add_paragraph("Abstract")
        doc.add_paragraph("Fig. 4. Body figure")

        process_document(
            doc,
            mode=ProcessingMode.FIX_NUMBERING_ONLY,
            options=FormattingOptions(preserve_cover_page=True),
        )
        texts = [paragraph.text for paragraph in doc.paragraphs]

        self.assertEqual(texts[0], "Fig. 9. Cover page illustration")
        self.assertIn("Fig. 1. Body figure", texts)


if __name__ == "__main__":
    unittest.main()
