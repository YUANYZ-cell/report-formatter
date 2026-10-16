from __future__ import annotations

from pathlib import Path
import unittest

from docx import Document

from formatter.equation_utils import get_visible_text_from_xml
from formatter.formatter_engine import FormattingOptions, ProcessingMode, process_document
from formatter.utils import int_to_roman


TARGET = Path(__file__).resolve().parent / "fixtures" / "optional_transformer_report.docx"


class RootCauseRegressionTest(unittest.TestCase):
    def test_roman_boundaries(self):
        self.assertEqual([int_to_roman(value) for value in (1, 2, 4, 9, 40)], ["I", "II", "IV", "IX", "XL"])

    @unittest.skipUnless(TARGET.exists(), "real Transformer report fixture is unavailable")
    def test_real_report_strict_ieee_conversion_is_semantic_and_idempotent(self):
        document = Document(TARGET)
        result = process_document(
            document,
            mode=ProcessingMode.FIX_NUMBERING_AND_FORMATTING,
            options=FormattingOptions(format_profile_name="ieee", ui_language="en"),
        )

        self.assertEqual(result.start_index, 67)
        self.assertFalse([action for action in result.actions if action.kind == "equation"])
        self.assertFalse(result.unnumbered_equations)

        texts = [" ".join(paragraph.text.split()) for paragraph in document.paragraphs]
        for number in (1, 2, 3):
            marker = f"TABLE {int_to_roman(number)}"
            index = texts.index(marker)
            self.assertTrue(texts[index + 1])
            self.assertEqual(texts[index + 1], texts[index + 1].upper())
        self.assertNotIn("Table 2: Equivalent Circuit Parameters Derived from Empirical Testing", texts)

        expected_headings = {
            "I. INTRODUCTION",
            "A. The Ideal vs. Real Transformer",
            "II. METHODOLOGY",
            "III. RESULTS AND DISCUSSION",
            "2) Open-Circuit Test: Excitation Branch Analysis:",
            "C. Loss Distribution Analysis at Full Load",
            "D. Full-Load Efficiency Evaluation",
            "IV. CONCLUSION & FUTURE WORKS",
        }
        self.assertTrue(expected_headings.issubset(set(texts)))
        full_load_body = next(text for text in texts if text.startswith("Building upon the holistic loss distribution"))
        self.assertIn("efficiency", full_load_body.lower())

        visible = [get_visible_text_from_xml(paragraph) for paragraph in document.paragraphs]
        for equation_number in range(14, 24):
            self.assertTrue(any(text.rstrip().endswith(f"({equation_number})") for text in visible))

        second = process_document(
            document,
            mode=ProcessingMode.FIX_NUMBERING_AND_FORMATTING,
            options=FormattingOptions(format_profile_name="ieee", ui_language="en"),
        )
        numbering_actions = [
            action for action in second.actions if action.kind in {"table", "equation", "heading"}
        ]
        self.assertFalse(numbering_actions)


if __name__ == "__main__":
    unittest.main()
