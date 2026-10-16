from __future__ import annotations

from datetime import datetime
import html
import importlib
import io
import os
from pathlib import Path
import re
import subprocess

import streamlit as st
from docx import Document

from formatter import caption_rules, equation_detector, equation_rules, formatter_engine
import formatter.localization as localization_module
import formatter.preview_service as preview_service_module
from formatter.formatter_engine import FormattingOptions, ProcessingMode
from formatter.language_detector import detect_report_language, resolve_report_language
from formatter.localization import (
    ALIGNMENT_LABELS,
    CATEGORY_LABELS,
    EQUATION_WRAPPER_LABELS,
    LANGUAGE_NAMES,
    MODE_LABELS,
    NUMBER_STYLE_LABELS,
    PROFILE_LABELS,
    ROLE_LABELS,
    SEVERITY_LABELS,
    STYLE_SOURCE_LABELS,
    TABLE_LAYOUT_LABELS,
    label,
    reverse_label,
    stat_card_text,
    t,
)
from formatter.preview_service import (
    PreviewEnvironment,
    PreviewMode,
    PreviewUnavailable,
    build_document_version,
    build_preview_version,
    detect_preview_environment,
    evaluate_preview_gate,
    items_for_page,
    render_page_html,
    render_preview_bundle,
)


_EXPECTED_PREVIEW_CACHE_VERSION = "preview-cache-2026-06-25-safe-pixel-diff-v2"


def _bind_preview_service_symbols() -> None:
    """Keep Streamlit hot-reload from mixing old preview modules with new app code."""
    global preview_service_module
    global PreviewEnvironment, PreviewMode, PreviewUnavailable
    global build_document_version, build_preview_version, detect_preview_environment
    global evaluate_preview_gate, items_for_page, render_page_html, render_preview_bundle

    env_fields = getattr(getattr(preview_service_module, "PreviewEnvironment", None), "__dataclass_fields__", {})
    version = getattr(preview_service_module, "PREVIEW_CACHE_VERSION", "")
    needs_reload = version != _EXPECTED_PREVIEW_CACHE_VERSION or "pdftotext_path" not in env_fields
    if needs_reload:
        preview_service_module = importlib.reload(preview_service_module)

    PreviewEnvironment = preview_service_module.PreviewEnvironment
    PreviewMode = preview_service_module.PreviewMode
    PreviewUnavailable = preview_service_module.PreviewUnavailable
    build_document_version = preview_service_module.build_document_version
    build_preview_version = preview_service_module.build_preview_version
    detect_preview_environment = preview_service_module.detect_preview_environment
    evaluate_preview_gate = preview_service_module.evaluate_preview_gate
    items_for_page = preview_service_module.items_for_page
    render_page_html = preview_service_module.render_page_html
    render_preview_bundle = preview_service_module.render_preview_bundle


_bind_preview_service_symbols()


def _preview_environment_from_kwargs(**kwargs) -> PreviewEnvironment:
    """Construct PreviewEnvironment across old cached module shapes during hot reload."""
    env_fields = getattr(PreviewEnvironment, "__dataclass_fields__", None)
    if env_fields:
        kwargs = {key: value for key, value in kwargs.items() if key in env_fields}
    return PreviewEnvironment(**kwargs)


def _preview_environment_attr(environment: PreviewEnvironment, name: str, default=None):
    return getattr(environment, name, default)


from formatter.pdf_pipeline import process_uploaded_pdf
from formatter.streamlit_pipeline import build_processing_signature, process_uploaded_docx


def _style_rows(profile, ui_language: str):
    rows = []
    for key in ["body_text", "figure_caption", "table_number", "table_caption", "equation"]:
        item = getattr(profile, key)
        rows.append(
            {
                t(ui_language, "role"): label(ROLE_LABELS, ui_language, key),
                t(ui_language, "font"): _style_value(item.font_name, ui_language),
                t(ui_language, "size"): _style_value(f"{item.font_size:g}" if item.font_size else None, ui_language),
                t(ui_language, "bold"): _format_boolean_value(item.bold, ui_language),
                t(ui_language, "italic"): _format_boolean_value(item.italic, ui_language),
                t(ui_language, "alignment"): label(ALIGNMENT_LABELS, ui_language, item.alignment)
                if item.alignment
                else t(ui_language, "not_detected"),
                t(ui_language, "source"): _recognition_source_label(item.source, item.role, ui_language),
                t(ui_language, "pattern"): _recognition_mode_label(item.source, item.pattern, item.role, ui_language),
            }
        )
    return rows


def _style_debug_rows(profile, ui_language: str):
    rows = []
    for key in ["body_text", "figure_caption", "table_number", "table_caption", "equation"]:
        item = getattr(profile, key)
        rows.append(
            {
                t(ui_language, "role"): label(ROLE_LABELS, ui_language, key),
                t(ui_language, "source"): item.source or t(ui_language, "not_detected"),
                t(ui_language, "pattern"): item.pattern or t(ui_language, "not_detected"),
            }
        )
    return rows


def _caption_rows(items, ui_language: str):
    if not items:
        return [
            {
                t(ui_language, "number"): t(ui_language, "not_detected"),
                t(ui_language, "paragraph"): t(ui_language, "not_applicable"),
                t(ui_language, "text"): t(ui_language, "none_detected"),
            }
        ]
    has_pdf_page = any(getattr(item, "pdf_page_number", None) is not None for item in items)
    page_header = "页码" if ui_language == "zh" else "Page"
    rows = []
    for item in items:
        row = {
            t(ui_language, "number"): item.number,
            t(ui_language, "paragraph"): (item.paragraph_index + 1)
            if item.paragraph_index is not None
            else t(ui_language, "not_applicable"),
            t(ui_language, "source"): getattr(item, "source", "paragraph"),
            t(ui_language, "text"): item.text or item.raw_text,
        }
        if has_pdf_page:
            row = {page_header: getattr(item, "pdf_page_number", None) or t(ui_language, "not_detected"), **row}
        rows.append(row)
    return rows


def _image_rows(items, ui_language: str):
    if ui_language == "zh":
        headers = {
            "image": "图片",
            "paragraph": "段落",
            "file": "文件",
            "classification": "分类",
            "ocr_title": "OCR 标题",
            "confidence": "置信度",
        }
    else:
        headers = {
            "image": "Image",
            "paragraph": "Paragraph",
            "file": "File",
            "classification": "Classification",
            "ocr_title": "OCR title",
            "confidence": "Confidence",
        }
    if not items:
        return [
            {
                headers["image"]: t(ui_language, "not_detected"),
                headers["paragraph"]: t(ui_language, "not_applicable"),
                headers["file"]: t(ui_language, "none_detected"),
                headers["classification"]: t(ui_language, "none_detected"),
                headers["ocr_title"]: t(ui_language, "none_detected"),
                headers["confidence"]: t(ui_language, "not_applicable"),
            }
        ]
    rows = []
    for item in items:
        if getattr(item, "is_suspected_table", False):
            classification = "image_table"
        elif getattr(item, "is_suspected_figure", False):
            classification = "image_figure"
        else:
            classification = "unparsed_image"
        rows.append(
            {
                headers["image"]: item.image_index,
                headers["paragraph"]: (item.paragraph_index + 1)
                if item.paragraph_index is not None
                else t(ui_language, "not_applicable"),
                headers["file"]: item.file_name,
                headers["classification"]: classification,
                headers["ocr_title"]: item.ocr_title_text or t(ui_language, "not_detected"),
                headers["confidence"]: f"{item.confidence:.2f}",
            }
        )
    return rows


def _yes_no(value, ui_language: str):
    return _format_boolean_value(value, ui_language)


def _format_boolean_value(value, ui_language: str) -> str:
    if value is True:
        return t(ui_language, "yes")
    if value is False:
        return t(ui_language, "no")
    return t(ui_language, "not_detected")


def _style_value(value, ui_language: str) -> str:
    if value is None or value == "":
        return t(ui_language, "not_detected")
    return str(value)


def _humanize_internal_text(value: str | None) -> str:
    if not value:
        return ""
    return value.replace("_", " ").replace("-", " ").capitalize()


def _recognition_source_label(source: str | None, role: str, ui_language: str) -> str:
    source_value = (source or "").lower()
    if "fallback" in source_value:
        return t(ui_language, "recognition_fallback")
    if "template" in source_value:
        return t(ui_language, "recognition_template")
    if "uploaded report" in source_value or "inferred" in source_value:
        return t(ui_language, "recognition_uploaded_report")
    return _recognition_mode_label(source, None, role, ui_language)


def _recognition_mode_label(source: str | None, pattern: str | None, role: str, ui_language: str) -> str:
    source_value = (source or "").lower()
    pattern_value = pattern or ""
    if role == "body_text" or "body" in source_value:
        return t(ui_language, "recognition_inferred_body")
    if role == "figure_caption" or "figure" in source_value:
        return t(ui_language, "recognition_figure_regex")
    if role == "table_number" or "table number" in source_value:
        return t(ui_language, "recognition_table_number_regex")
    if role == "table_caption" or "table caption" in source_value or pattern_value == "next paragraph after TABLE N":
        return t(ui_language, "recognition_table_caption_after_number")
    if role == "equation" or "equation" in source_value:
        return t(ui_language, "recognition_equation_regex")
    if "heading" in source_value:
        return t(ui_language, "recognition_inferred_heading")
    if "fallback" in source_value:
        return t(ui_language, "recognition_fallback")
    return _humanize_internal_text(source or pattern) or t(ui_language, "not_detected")


def _option_labels(mapping: dict[str, dict[str, str]], ui_language: str, keys: list[str]) -> list[str]:
    return [label(mapping, ui_language, key) for key in keys]


UI_LANGUAGE_OPTIONS = ["en", "zh"]
UI_LANGUAGE_LABELS = {"en": "English", "zh": "中文"}
UI_LANGUAGE_ALIASES = {
    "English": "en",
    "英文": "en",
    "Chinese": "zh",
    "中文": "zh",
}
SESSION_DEFAULTS = {
    "ui_language": "en",
    "report_language_override": "auto",
    "format_profile_name": "template_matching",
    "processing_mode": "analyze_only",
    "style_source": "auto_detect",
    "figure_prefix": "Fig.",
    "figure_number_style": "arabic",
    "figure_separator": ".",
    "table_prefix": "Table",
    "table_number_style": "roman_upper",
    "table_layout": "single_line",
    "equation_wrapper": "parentheses",
    "preserve_cover_page": True,
    "preserve_two_column_layout": True,
    "insert_placeholder_captions": False,
    "renumber_figures": True,
    "renumber_tables": True,
    "renumber_equations": False,
    "check_references": True,
    "generate_report": True,
    "active_preview_mode": None,
    "preview_page_number": 1,
    "preview_zoom": 100,
    "preview_show_markers": True,
    "comparison_preview_mode": "side_by_side",
    "comparison_sync_scroll": True,
    "comparison_sync_zoom": True,
    "results_view_active": True,
}
LATEST_RESULT_KEYS = [
    "latest_result",
    "latest_original_docx_bytes",
    "latest_original_file_name",
    "latest_docx_bytes",
    "latest_report_bytes",
    "latest_run_id",
    "latest_docx_file_name",
    "latest_report_file_name",
    "latest_debug_info",
    "latest_input_signature",
    "latest_document_version",
    "latest_preview_version",
]


def _init_session_state() -> None:
    for key, default in SESSION_DEFAULTS.items():
        if key not in st.session_state:
            st.session_state[key] = default


def _clear_latest_result() -> None:
    for key in LATEST_RESULT_KEYS:
        st.session_state.pop(key, None)
    st.session_state["active_preview_mode"] = None
    st.session_state["preview_page_number"] = 1
    st.session_state["preview_zoom"] = 100


def _render_preview_upload_state_keepers(ui_language: str) -> None:
    with st.container(key="preview_upload_state_keeper"):
        st.file_uploader(t(ui_language, "upload_report"), type=["docx", "pdf"], key="report_file")
        st.file_uploader(t(ui_language, "upload_template"), type=["docx"], key="template_file")


def _coerce_session_value(key: str, valid_values: list[str], default: str) -> None:
    if key not in st.session_state:
        st.session_state[key] = default
    elif st.session_state[key] not in valid_values:
        st.session_state[key] = default


def _coerce_session_enum(key: str, valid_values: list[str], default: str, aliases: dict[str, str]) -> None:
    if key not in st.session_state:
        st.session_state[key] = default
        return
    current = st.session_state[key]
    if current in aliases:
        st.session_state[key] = aliases[current]
    elif current not in valid_values:
        st.session_state[key] = default


def _localized_selectbox(
    label_text: str,
    options: list[str],
    state_key: str,
    default: str,
    ui_language: str,
    format_func,
    aliases: dict[str, str] | None = None,
) -> str:
    """Keep stable values in state while forcing Streamlit to refresh translated labels."""
    _coerce_session_enum(state_key, options, default, aliases or {})
    widget_key = f"{state_key}_{ui_language}_select"
    language_key = f"{state_key}_select_language"
    if st.session_state.get(language_key) != ui_language:
        st.session_state[widget_key] = st.session_state[state_key]
        st.session_state[language_key] = ui_language
    elif widget_key not in st.session_state or st.session_state[widget_key] not in options:
        st.session_state[widget_key] = st.session_state[state_key]

    selected_value = st.selectbox(
        label_text,
        options,
        index=options.index(st.session_state[widget_key]),
        format_func=format_func,
        key=widget_key,
    )
    st.session_state[state_key] = selected_value
    return selected_value


def _state_index(state_key: str, options: list[str], default: str, aliases: dict[str, str] | None = None) -> int:
    _coerce_session_enum(state_key, options, default, aliases or {})
    return options.index(st.session_state[state_key])


def _state_selectbox(
    label_text: str,
    options: list[str],
    state_key: str,
    default: str,
    ui_language: str,
    format_func,
    aliases: dict[str, str] | None = None,
) -> str:
    selected_value = st.selectbox(
        label_text,
        options,
        index=_state_index(state_key, options, default, aliases),
        format_func=format_func,
        key=state_key,
    )
    return selected_value


def _ui_language_selectbox() -> str:
    _coerce_session_enum("ui_language", UI_LANGUAGE_OPTIONS, "en", UI_LANGUAGE_ALIASES)
    current_language = st.session_state["ui_language"]
    ui_language = st.selectbox(
        t(current_language, "ui_language"),
        UI_LANGUAGE_OPTIONS,
        index=UI_LANGUAGE_OPTIONS.index(current_language),
        format_func=lambda value: UI_LANGUAGE_LABELS[value],
        key="ui_language",
    )
    return ui_language


def _language_name(ui_language: str, value: str) -> str:
    return LANGUAGE_NAMES[ui_language].get(value, value)


def _render_summary_cards(cards: list[dict[str, object]]) -> None:
    card_html = []
    for card in cards:
        title = html.escape(str(card["title"]))
        value = html.escape(str(card["value"]))
        caption = html.escape(str(card["caption"]))
        card_html.append(
            '<div class="summary-card">'
            f'<div class="summary-card-title">{title}</div>'
            f'<div class="summary-card-value">{value}</div>'
            f'<div class="summary-card-caption">{caption}</div>'
            "</div>"
        )
    st.markdown(f'<div class="summary-grid">{"".join(card_html)}</div>', unsafe_allow_html=True)


def _stats_card(ui_language: str, card_id: str, value: object) -> dict[str, object]:
    return {
        "title": stat_card_text(ui_language, card_id, "title"),
        "value": value,
        "caption": stat_card_text(ui_language, card_id, "caption"),
    }


def _stats_help_body_html(body: str) -> str:
    paragraphs = []
    bullets = []
    for line in body.splitlines():
        cleaned = line.strip()
        if not cleaned:
            continue
        if cleaned.startswith("- "):
            bullets.append(html.escape(cleaned[2:]))
        else:
            paragraphs.append(html.escape(cleaned))

    paragraph_html = "".join(f"<p>{paragraph}</p>" for paragraph in paragraphs)
    bullet_html = ""
    if bullets:
        bullet_html = "<ul>" + "".join(f"<li>{bullet}</li>" for bullet in bullets) + "</ul>"
    return paragraph_html + bullet_html


def _render_stats_help(ui_language: str) -> None:
    title = html.escape(t(ui_language, "stats_help_title"))
    body_html = _stats_help_body_html(t(ui_language, "stats_help_body"))
    st.markdown(
        f"""
        <details class="stats-help-panel">
            <summary>
                <span class="stats-help-chevron">›</span>
                <span class="stats-help-title">{title}</span>
            </summary>
            <div class="stats-help-body">{body_html}</div>
        </details>
        """,
        unsafe_allow_html=True,
    )


def _render_download_debug_info(debug_info: dict[str, object] | None, ui_language: str) -> None:
    if not debug_info:
        return
    preview = debug_info.get("exported_docx_preview") or []
    metadata = {key: value for key, value in debug_info.items() if key != "exported_docx_preview"}
    title = "本次导出 DOCX 的调试信息" if ui_language == "zh" else "Debug info for this exported DOCX"
    with st.expander(title, expanded=False):
        st.json(metadata)
        preview_lines = "\n".join(f"- {line}" for line in preview) if preview else "- No preview lines found."
        st.code("Expected smoke preview from exported docx:\n" + preview_lines, language="text")


def _render_readable_table(rows: list[dict[str, object]]) -> None:
    if not rows:
        return
    headers = list(rows[0].keys())
    thead = "".join(f"<th>{html.escape(str(header))}</th>" for header in headers)
    body_rows = []
    for row in rows:
        cells = "".join(f"<td>{html.escape(str(row.get(header, '')))}</td>" for header in headers)
        body_rows.append(f"<tr>{cells}</tr>")
    st.markdown(
        '<div class="readable-table-wrap">'
        '<table class="readable-table">'
        f"<thead><tr>{thead}</tr></thead>"
        f'<tbody>{"".join(body_rows)}</tbody>'
        "</table>"
        "</div>",
        unsafe_allow_html=True,
    )


ISSUE_GROUP_ORDER = [
    "table",
    "image",
    "equation",
    "reference_numbering",
    "format_style",
    "manual_review",
    "unresolved_images",
    "other",
]
REVIEW_GROUPS_DEFAULT_COLLAPSED = {"manual_review", "unresolved_images"}


def _render_issue_groups(result, ui_language: str, run_id: str) -> None:
    groups = _build_issue_groups(result, ui_language)
    if not groups:
        st.success(t(ui_language, "no_reference_issues"))
        return

    state_prefix = f"issue_group_open_{run_id or 'current'}_"
    first_open_key = _default_open_issue_group(groups)
    for group in groups:
        state_key = state_prefix + group["key"]
        if state_key not in st.session_state:
            st.session_state[state_key] = group["key"] == first_open_key

    control_cols = st.columns([0.72, 0.72, 4.6])
    with control_cols[0]:
        if st.button(t(ui_language, "expand_all"), use_container_width=True, key=f"{state_prefix}expand_all"):
            for group in groups:
                st.session_state[state_prefix + group["key"]] = True
            st.rerun()
    with control_cols[1]:
        if st.button(t(ui_language, "collapse_all"), use_container_width=True, key=f"{state_prefix}collapse_all"):
            for group in groups:
                st.session_state[state_prefix + group["key"]] = False
            st.rerun()

    for group in groups:
        _render_issue_group(group, state_prefix, ui_language)


def _build_issue_groups(result, ui_language: str) -> list[dict[str, object]]:
    buckets: dict[str, list[dict[str, object]]] = {key: [] for key in ISSUE_GROUP_ORDER}
    issue_meta: dict[str, dict[str, int]] = {key: {"errors": 0, "warnings": 0, "reviews": 0} for key in ISSUE_GROUP_ORDER}

    for issue in getattr(result, "issues", []) or []:
        key = _issue_group_key(issue)
        severity = (getattr(issue, "severity", "warning") or "warning").lower()
        if severity == "error":
            issue_meta[key]["errors"] += 1
        elif severity == "warning":
            issue_meta[key]["warnings"] += 1
        else:
            issue_meta[key]["reviews"] += 1
        buckets[key].append(
            {
                t(ui_language, "severity"): label(SEVERITY_LABELS, ui_language, severity),
                t(ui_language, "category"): label(CATEGORY_LABELS, ui_language, getattr(issue, "category", "")),
                t(ui_language, "paragraph"): (issue.paragraph_index + 1)
                if getattr(issue, "paragraph_index", None) is not None
                else t(ui_language, "not_applicable"),
                t(ui_language, "message"): _friendly_table_text(getattr(issue, "message", "")),
            }
        )

    action_rows = _action_issue_rows(getattr(result, "actions", []), ui_language)
    if action_rows:
        buckets["format_style"].extend(action_rows)
        issue_meta["format_style"]["reviews"] += len(action_rows)

    manual_rows = _manual_review_rows(getattr(result, "manual_review", []), ui_language)
    if manual_rows:
        buckets["manual_review"].extend(manual_rows)
        issue_meta["manual_review"]["reviews"] += len(manual_rows)

    unresolved_rows = _unresolved_image_rows(getattr(result, "images", []), ui_language)
    if unresolved_rows:
        buckets["unresolved_images"].extend(unresolved_rows)
        issue_meta["unresolved_images"]["reviews"] += len(unresolved_rows)

    groups: list[dict[str, object]] = []
    for key in ISSUE_GROUP_ORDER:
        rows = buckets[key]
        if not rows:
            continue
        meta = issue_meta[key]
        groups.append(
            {
                "key": key,
                "title": t(ui_language, f"issue_group_{key}"),
                "rows": rows,
                "count": len(rows),
                "errors": meta["errors"],
                "warnings": meta["warnings"],
                "reviews": meta["reviews"],
            }
        )
    return groups


def _render_issue_group(group: dict[str, object], state_prefix: str, ui_language: str) -> None:
    state_key = state_prefix + str(group["key"])
    is_open = bool(st.session_state.get(state_key, False))
    label_text = _issue_group_button_label(group, is_open, ui_language)
    if st.button(label_text, use_container_width=True, key=f"{state_key}_toggle"):
        st.session_state[state_key] = not is_open
        st.rerun()
    if is_open:
        _render_readable_table(group["rows"])


def _issue_group_button_label(group: dict[str, object], is_open: bool, ui_language: str) -> str:
    parts = [
        "-" if is_open else "+",
        str(group["title"]),
        f"{group['count']} {t(ui_language, 'items')}",
    ]
    if group.get("errors"):
        parts.append(f"{t(ui_language, 'errors')}: {group['errors']}")
    if group.get("warnings"):
        parts.append(f"{t(ui_language, 'warnings')}: {group['warnings']}")
    if group.get("reviews"):
        parts.append(f"{t(ui_language, 'reviews')}: {group['reviews']}")
    return " | ".join(parts)


def _default_open_issue_group(groups: list[dict[str, object]]) -> str | None:
    for group in groups:
        if group["key"] not in REVIEW_GROUPS_DEFAULT_COLLAPSED and int(group.get("errors") or 0) > 0:
            return str(group["key"])
    for group in groups:
        if group["key"] not in REVIEW_GROUPS_DEFAULT_COLLAPSED:
            return str(group["key"])
    return None


def _issue_group_key(issue) -> str:
    category = str(getattr(issue, "category", "") or "").lower()
    message = str(getattr(issue, "message", "") or "").lower()
    if category in {"table", "native_table", "image_table"}:
        return "table"
    if category in {"figure", "image", "picture"}:
        return "image"
    if category == "equation":
        return "equation"
    if category in {"reference", "citation", "numbering"}:
        return "reference_numbering"
    if category in {"style", "format", "font", "paragraph", "layout"}:
        return "format_style"
    if any(token in message for token in ("reference", "citation", "numbering", "renumber", "duplicate number")):
        return "reference_numbering"
    if any(token in message for token in ("style", "format", "font", "layout", "spacing")):
        return "format_style"
    return "other"


def _manual_review_rows(items, ui_language: str) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    seen: set[str] = set()
    for item in items or []:
        message = _friendly_table_text(item)
        marker = _dedupe_marker(message)
        if not marker or marker in seen:
            continue
        seen.add(marker)
        rows.append({t(ui_language, "message"): message})
    return rows


def _action_issue_rows(actions, ui_language: str) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for action in actions or []:
        rows.append(
            {
                t(ui_language, "kind"): label(CATEGORY_LABELS, ui_language, getattr(action, "kind", "")),
                t(ui_language, "paragraph"): (action.paragraph_index + 1)
                if getattr(action, "paragraph_index", None) is not None
                else t(ui_language, "not_applicable"),
                t(ui_language, "before"): _friendly_table_text(getattr(action, "before", "")),
                t(ui_language, "after"): _friendly_table_text(getattr(action, "after", "")),
            }
        )
    return rows


def _unresolved_image_rows(images, ui_language: str) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    seen: set[str] = set()
    for image in images or []:
        unresolved = not getattr(image, "is_suspected_table", False) and not getattr(image, "is_suspected_figure", False)
        ocr_failed = bool(getattr(image, "ocr_error", None)) or not getattr(image, "ocr_available", False)
        if not unresolved and not ocr_failed:
            continue
        paragraph = getattr(image, "paragraph_index", None)
        reason = getattr(image, "ocr_error", None) or (
            t(ui_language, "ocr_unavailable") if ocr_failed else t(ui_language, "unresolved_image")
        )
        marker = _dedupe_marker(f"{getattr(image, 'image_index', '')}|{getattr(image, 'file_name', '')}|{reason}")
        if marker in seen:
            continue
        seen.add(marker)
        rows.append(
            {
                t(ui_language, "image"): getattr(image, "image_index", ""),
                t(ui_language, "paragraph"): (paragraph + 1) if paragraph is not None else t(ui_language, "not_applicable"),
                t(ui_language, "file"): getattr(image, "file_name", ""),
                t(ui_language, "classification"): _image_classification_label(image, ui_language),
                t(ui_language, "message"): _friendly_table_text(reason),
            }
        )
    return rows


def _image_classification_label(image, ui_language: str) -> str:
    if getattr(image, "is_suspected_table", False):
        return t(ui_language, "image_table")
    if getattr(image, "is_suspected_figure", False):
        return t(ui_language, "image_figure")
    return t(ui_language, "unresolved_image")


def _friendly_table_text(value: object, max_length: int = 700) -> str:
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    if len(text) > max_length:
        return text[: max_length - 1].rstrip() + "..."
    return text


def _dedupe_marker(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip().casefold()


def _diagnostic_snapshot() -> dict[str, str]:
    app_path = Path(__file__).resolve()
    localization_path = Path(localization_module.__file__).resolve()
    cwd = Path.cwd().resolve()
    app_mtime = datetime.fromtimestamp(app_path.stat().st_mtime).isoformat(timespec="seconds")
    localization_mtime = datetime.fromtimestamp(localization_path.stat().st_mtime).isoformat(timespec="seconds")
    current_locale = str(st.session_state.get("ui_language", "en"))
    return {
        "app_path": str(app_path),
        "cwd": str(cwd),
        "pid": str(os.getpid()),
        "git_commit": _git_value(["git", "rev-parse", "--short", "HEAD"], cwd),
        "git_branch": _git_value(["git", "branch", "--show-current"], cwd),
        "app_mtime": app_mtime,
        "build_time": app_mtime,
        "frontend_asset_version": app_mtime,
        "localization_path": str(localization_path),
        "localization_mtime": localization_mtime,
        "current_locale": current_locale,
        "return_analysis_results_label": t(current_locale, "return_analysis_results"),
        "return_previous_level_label": t(current_locale, "return_previous_level"),
        "current_theme": "light",
        "loaded_at": datetime.now().isoformat(timespec="seconds"),
        "preview_service": str(Path(preview_service_module.__file__).resolve()),
    }


def _git_value(command: list[str], cwd: Path) -> str:
    try:
        completed = subprocess.run(
            command,
            cwd=str(cwd),
            timeout=5,
            check=False,
            text=True,
            encoding="utf-8",
            errors="replace",
            stdout=subprocess.PIPE,
            stderr=subprocess.DEVNULL,
        )
    except Exception:
        return "unavailable"
    value = (completed.stdout or "").strip()
    return value or "unavailable"


def _render_diagnostics(ui_language: str) -> None:
    with st.expander(t(ui_language, "runtime_diagnostics"), expanded=False):
        st.json(_diagnostic_snapshot())


def _is_fix_mode(mode: str) -> bool:
    return ProcessingMode(mode) in {
        ProcessingMode.FIX_NUMBERING_AND_FORMATTING,
        ProcessingMode.FIX_NUMBERING_ONLY,
    }


def _has_equation_auto_edit_suspended_notice(result) -> bool:
    return not getattr(formatter_engine, "EQUATION_AUTO_EDIT_ENABLED", True) or any(
        "automatic equation renumbering is suspended" in str(item).lower()
        for item in (getattr(result, "manual_review", []) or [])
    )


def _preview_mode_for_processing_mode(mode: str) -> PreviewMode:
    return PreviewMode.COMPARISON if _is_fix_mode(mode) else PreviewMode.ISSUE


def _preview_view_options(ui_language: str) -> dict[str, str]:
    return {
        "side_by_side": t(ui_language, "side_by_side_comparison"),
        "modified_only": t(ui_language, "modified_preview"),
        "diff": t(ui_language, "difference_view"),
    }


@st.cache_data(show_spinner=False)
def _cached_render_preview_bundle(
    *,
    mode_value: str,
    task_id: str,
    document_version: str,
    original_docx_bytes: bytes,
    original_file_name: str,
    modified_docx_bytes: bytes | None,
    modified_file_name: str | None,
    dpi: int,
    max_pages: int,
    page_number: int,
    preload_radius: int,
    converter_available: bool,
    renderer_available: bool,
    converter_label: str,
    renderer_label: str,
    gotenberg_url: str | None,
    libreoffice_path: str | None,
    pdftoppm_path: str | None,
    pdftotext_path: str | None,
    word_available: bool,
    word_status: str,
    word_path: str | None,
    word_error: str | None,
    low_fidelity_available: bool,
    low_fidelity_label: str,
    _result,
):
    environment = _preview_environment_from_kwargs(
        converter_available=converter_available,
        renderer_available=renderer_available,
        converter_label=converter_label,
        renderer_label=renderer_label,
        gotenberg_url=gotenberg_url,
        libreoffice_path=libreoffice_path,
        pdftoppm_path=pdftoppm_path,
        pdftotext_path=pdftotext_path,
        word_available=word_available,
        word_status=word_status,
        word_path=word_path,
        word_error=word_error,
        low_fidelity_available=low_fidelity_available,
        low_fidelity_label=low_fidelity_label,
    )
    return render_preview_bundle(
        mode=mode_value,
        task_id=task_id,
        document_version=document_version,
        original_docx_bytes=original_docx_bytes,
        original_file_name=original_file_name,
        result=_result,
        modified_docx_bytes=modified_docx_bytes,
        modified_file_name=modified_file_name,
        dpi=dpi,
        max_pages=max_pages,
        page_number=page_number,
        preload_radius=preload_radius,
        environment=environment,
    )


def _render_preview_entry(gate, environment: PreviewEnvironment, ui_language: str) -> None:
    title_key = "issue_preview_title" if gate.mode == PreviewMode.ISSUE else "comparison_preview_title"
    desc_key = "open_issue_preview_help" if gate.mode == PreviewMode.ISSUE else "open_comparison_preview_help"
    state_label = t(ui_language, gate.reason_key)
    st.markdown(
        f"""
        <div class="preview-action-band">
            <div>
                <div class="preview-action-title">{html.escape(t(ui_language, title_key))}</div>
                <div class="preview-action-desc">{html.escape(t(ui_language, desc_key))}</div>
            </div>
            <div class="preview-action-meta">
                <span>{html.escape(environment.converter_label)}</span>
                <span>{html.escape(environment.renderer_label)}</span>
                <span>{html.escape(state_label)}</span>
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )
    button_help = t(ui_language, gate.help_key if gate.enabled else gate.reason_key)
    if st.button(
        t(ui_language, gate.button_key),
        type="primary",
        use_container_width=True,
        disabled=not gate.enabled,
        help=button_help,
        key=gate.button_key,
    ):
        st.session_state["active_preview_mode"] = gate.mode.value
        st.session_state["preview_page_number"] = 1
        st.session_state["preview_show_markers"] = True
        st.rerun()


def _render_preview_workspace(
    *,
    ui_language: str,
    mode: PreviewMode,
    report_bytes: bytes,
    report_file_name: str,
    modified_docx_bytes: bytes | None,
    modified_file_name: str | None,
    result,
    task_id: str,
    document_version: str,
    environment: PreviewEnvironment,
) -> None:
    st.markdown('<div class="preview-route-top-spacer" aria-hidden="true"></div>', unsafe_allow_html=True)
    header_cols = st.columns([1.3, 2.4, 1.2])
    with header_cols[0]:
        if st.button(t(ui_language, "return_analysis_results"), use_container_width=True, key="return_from_preview"):
            st.session_state["active_preview_mode"] = None
            st.session_state["results_view_active"] = True
            st.rerun()
    with header_cols[1]:
        st.markdown(
            f"<div class='preview-workspace-title'>{html.escape(_preview_title_for_mode(mode, ui_language))}</div>",
            unsafe_allow_html=True,
        )
    with header_cols[2]:
        st.caption(f"{environment.converter_label} -> {environment.renderer_label}")

    requested_page = max(1, int(st.session_state.get("preview_page_number", 1)))
    try:
        with st.spinner(t(ui_language, "preparing_preview")):
            bundle = _cached_render_preview_bundle(
                mode_value=mode.value,
                task_id=task_id,
                document_version=document_version,
                original_docx_bytes=report_bytes,
                original_file_name=report_file_name,
                modified_docx_bytes=modified_docx_bytes if mode == PreviewMode.COMPARISON else None,
                modified_file_name=modified_file_name,
                dpi=120,
                max_pages=40,
                page_number=requested_page,
                preload_radius=1,
                converter_available=environment.converter_available,
                renderer_available=environment.renderer_available,
                converter_label=environment.converter_label,
                renderer_label=environment.renderer_label,
                gotenberg_url=environment.gotenberg_url,
                libreoffice_path=environment.libreoffice_path,
                pdftoppm_path=environment.pdftoppm_path,
                pdftotext_path=_preview_environment_attr(environment, "pdftotext_path"),
                word_available=environment.word_available,
                word_status=environment.word_status,
                word_path=environment.word_path,
                word_error=environment.word_error,
                low_fidelity_available=environment.low_fidelity_available,
                low_fidelity_label=environment.low_fidelity_label,
                _result=result,
            )
    except PreviewUnavailable as exc:
        st.error(f"{t(ui_language, 'preview_generation_failed')}: {exc}")
        return
    layout_guard_blocked = bool(
        bundle.modified is not None and bundle.original.page_count != bundle.modified.page_count
    )
    if layout_guard_blocked:
        st.session_state["latest_preview_version"] = None
        if ui_language == "zh":
            st.error(
                f"导出已阻断：修改后文件页数从 {bundle.original.page_count} 页变为 "
                f"{bundle.modified.page_count} 页。请保留原版式或改用仅分析模式。"
            )
        else:
            st.error(
                f"Export blocked: the modified file changed from {bundle.original.page_count} "
                f"pages to {bundle.modified.page_count} pages. Preserve the original layout or use analyze-only review."
            )
    else:
        st.session_state["latest_preview_version"] = bundle.preview_version

    max_page = max(bundle.original.page_count, bundle.modified.page_count if bundle.modified else 1)
    current_page = min(max(1, requested_page), max_page)
    st.session_state["preview_page_number"] = current_page

    toolbar_cols = st.columns([0.68, 0.68, 0.9, 0.82, 1.05, 1.0, 1.0, 1.0])
    with toolbar_cols[0]:
        if st.button(
            t(ui_language, "previous_page"),
            disabled=current_page <= 1,
            use_container_width=True,
            key="preview_previous_page",
        ):
            st.session_state["preview_page_number"] = max(1, current_page - 1)
            st.rerun()
    with toolbar_cols[1]:
        if st.button(
            t(ui_language, "next_page"),
            disabled=current_page >= max_page,
            use_container_width=True,
            key="preview_next_page",
        ):
            st.session_state["preview_page_number"] = min(max_page, current_page + 1)
            st.rerun()
    with toolbar_cols[2]:
        selected_page = st.number_input(
            t(ui_language, "page_number"),
            min_value=1,
            max_value=max_page,
            value=current_page,
            step=1,
            key=f"preview_page_input_{mode.value}_{document_version}",
        )
        if int(selected_page) != current_page:
            st.session_state["preview_page_number"] = int(selected_page)
            st.rerun()
    with toolbar_cols[3]:
        st.caption(f"{t(ui_language, 'total_pages')}: {max_page}")
    with toolbar_cols[4]:
        zoom_value = st.slider(
            t(ui_language, "zoom"),
            min_value=50,
            max_value=200,
            value=int(st.session_state.get("preview_zoom", 100)),
            step=10,
            key=f"preview_zoom_control_{mode.value}_{document_version}",
        )
        st.session_state["preview_zoom"] = int(zoom_value)
    with toolbar_cols[5]:
        st.session_state["preview_show_markers"] = st.checkbox(
            t(ui_language, "show_issue_markers" if mode == PreviewMode.ISSUE else "show_structured_changes"),
            value=bool(st.session_state.get("preview_show_markers", True)),
            key="preview_show_markers_control",
        )
    with toolbar_cols[6]:
        if mode == PreviewMode.COMPARISON:
            st.session_state["comparison_sync_scroll"] = st.checkbox(
                t(ui_language, "sync_scroll"),
                value=bool(st.session_state.get("comparison_sync_scroll", True)),
                key="comparison_sync_scroll_control",
            )
    with toolbar_cols[7]:
        if mode == PreviewMode.COMPARISON:
            st.session_state["comparison_sync_zoom"] = st.checkbox(
                t(ui_language, "sync_zoom"),
                value=bool(st.session_state.get("comparison_sync_zoom", True)),
                key="comparison_sync_zoom_control",
            )

    with st.expander(t(ui_language, "preview_performance"), expanded=False):
        st.json(bundle.timings)

    if mode == PreviewMode.ISSUE:
        _render_issue_preview(bundle, current_page, ui_language)
    else:
        _render_comparison_preview(bundle, current_page, ui_language)


def _preview_title_for_mode(mode: PreviewMode, ui_language: str) -> str:
    return t(ui_language, "issue_preview_title" if mode == PreviewMode.ISSUE else "comparison_preview_title")


def _get_rendered_page(document, page_number: int):
    pages = {page.page_number: page for page in document.rendered_pages}
    return pages.get(page_number) or document.rendered_pages[-1]


def _render_zoomed_page_html(page, markers, *, marker_visible: bool, label_text: str) -> str:
    zoom = max(50, min(200, int(st.session_state.get("preview_zoom", 100)))) / 100
    page_html = render_page_html(page, markers, marker_visible=marker_visible, label=label_text)
    return f'<div class="preview-zoom-frame" style="--preview-zoom:{zoom:.2f};">{page_html}</div>'


def _render_issue_preview(bundle, current_page: int, ui_language: str) -> None:
    page = _get_rendered_page(bundle.original, min(current_page, bundle.original.page_count))
    page_items = items_for_page(bundle.issues, page.page_number)
    left, center, right = st.columns([0.8, 2.2, 1.05])
    with left:
        _render_thumbnail_column(bundle.original, current_page, "issue", ui_language)
    with center:
        st.markdown(
            _render_zoomed_page_html(
                page,
                page_items,
                marker_visible=bool(st.session_state.get("preview_show_markers", True)),
                label_text=t(ui_language, "original_file"),
            ),
            unsafe_allow_html=True,
        )
    with right:
        st.subheader(t(ui_language, "current_page_issues"))
        _render_preview_items(page_items, ui_language, empty_key="no_page_issues")


def _render_comparison_preview(bundle, current_page: int, ui_language: str) -> None:
    view_options = _preview_view_options(ui_language)
    selected_view = st.radio(
        t(ui_language, "comparison_view_mode"),
        list(view_options.keys()),
        format_func=lambda value: view_options[value],
        key="comparison_preview_mode",
        horizontal=True,
    )
    original_page = _get_rendered_page(bundle.original, min(current_page, bundle.original.page_count))
    modified_doc = bundle.modified or bundle.original
    modified_page = _get_rendered_page(modified_doc, min(current_page, modified_doc.page_count))
    page_changes = items_for_page(bundle.changes, modified_page.page_number)
    page_pixel_changes = items_for_page(getattr(bundle, "pixel_changes", ()), modified_page.page_number)

    if selected_view == "modified_only":
        cols = st.columns([0.75, 2.25, 1.0])
        with cols[0]:
            _render_thumbnail_column(modified_doc, current_page, "modified", ui_language)
        with cols[1]:
            st.markdown(
                _render_zoomed_page_html(
                    modified_page,
                    page_changes,
                    marker_visible=False,
                    label_text=t(ui_language, "modified_file"),
                ),
                unsafe_allow_html=True,
            )
        with cols[2]:
            st.subheader(t(ui_language, "current_page_changes"))
            _render_preview_items(page_changes, ui_language, empty_key="no_page_changes")
        return

    if selected_view == "diff":
        st.info(t(ui_language, "pixel_diff_active" if page_pixel_changes else "pixel_diff_unavailable"))

    original_label = html.escape(t(ui_language, "original_file"))
    modified_label = html.escape(t(ui_language, "modified_file"))
    diff_markers = page_pixel_changes if selected_view == "diff" else []
    show_diff_markers = bool(st.session_state.get("preview_show_markers", True)) and bool(diff_markers)
    original_html = _render_zoomed_page_html(
        original_page,
        diff_markers,
        marker_visible=show_diff_markers,
        label_text=t(ui_language, "original_file"),
    )
    modified_html = _render_zoomed_page_html(
        modified_page,
        diff_markers,
        marker_visible=show_diff_markers,
        label_text=t(ui_language, "modified_file"),
    )
    st.markdown(
        f"""
        <div class="preview-columns">
            <section class="preview-pane">
                <div class="preview-pane-title">{original_label}</div>
                <div class="preview-canvas-container">{original_html}</div>
            </section>
            <section class="preview-pane">
                <div class="preview-pane-title">{modified_label}</div>
                <div class="preview-canvas-container">{modified_html}</div>
            </section>
        </div>
        """,
        unsafe_allow_html=True,
    )
    st.subheader(t(ui_language, "current_page_changes"))
    _render_preview_items(page_changes, ui_language, empty_key="no_page_changes")


def _render_thumbnail_column(document, current_page: int, prefix: str, ui_language: str) -> None:
    st.caption(t(ui_language, "page_thumbnails"))
    rendered_pages = {page.page_number: page for page in document.rendered_pages}
    first_page = max(1, current_page - 2)
    last_page = min(document.page_count, first_page + 5)
    for page_number in range(first_page, last_page + 1):
        button_type = "primary" if page_number == current_page else "secondary"
        if st.button(str(page_number), key=f"{prefix}_thumb_{page_number}", use_container_width=True, type=button_type):
            st.session_state["preview_page_number"] = page_number
            st.rerun()
        page = rendered_pages.get(page_number)
        if page:
            st.image(page.image_bytes, use_container_width=True)
        else:
            st.caption(t(ui_language, "preview_page_not_cached"))


def _render_preview_items(items, ui_language: str, *, empty_key: str) -> None:
    if not items:
        st.success(t(ui_language, empty_key))
        return
    for item in items:
        if item.location_status == "located":
            badge = t(ui_language, "located")
        else:
            badge = "未能精确定位，未绘制标注框" if ui_language == "zh" else "Could not locate precisely; no marker drawn"
        with st.container(border=True):
            st.markdown(
                f'<div class="preview-change-card" data-change-id="{html.escape(item.item_id)}">'
                f'<strong>{html.escape(item.title)}</strong></div>',
                unsafe_allow_html=True,
            )
            st.caption(f"{badge} | {t(ui_language, 'confidence')}: {item.confidence:.2f}")
            if item.message:
                st.write(item.message)
            if item.before or item.after:
                if item.before:
                    st.caption(t(ui_language, "before"))
                    st.code(item.before, language="text")
                if item.after:
                    st.caption(t(ui_language, "after"))
                    st.code(item.after, language="text")


def _render_active_preview_from_session(ui_language: str) -> bool:
    active_preview = st.session_state.get("active_preview_mode")
    result = st.session_state.get("latest_result")
    report_bytes = st.session_state.get("latest_original_docx_bytes")
    if not active_preview:
        return False
    if result is None or not report_bytes:
        st.session_state["active_preview_mode"] = None
        return False

    _render_preview_upload_state_keepers(ui_language)
    active_mode = PreviewMode(active_preview)
    modified_docx_bytes = st.session_state.get("latest_docx_bytes") if active_mode == PreviewMode.COMPARISON else None
    document_version = st.session_state.get("latest_document_version") or build_document_version(
        input_signature=st.session_state.get("latest_input_signature"),
        original_docx_bytes=report_bytes,
        modified_docx_bytes=modified_docx_bytes,
    )
    _render_preview_workspace(
        ui_language=ui_language,
        mode=active_mode,
        report_bytes=report_bytes,
        report_file_name=st.session_state.get("latest_original_file_name", "report.docx"),
        modified_docx_bytes=modified_docx_bytes,
        modified_file_name=st.session_state.get("latest_docx_file_name"),
        result=result,
        task_id=st.session_state.get("latest_run_id", ""),
        document_version=document_version,
        environment=detect_preview_environment(),
    )
    return True


st.set_page_config(page_title="Lab Report Formatter", layout="wide")
_init_session_state()

st.markdown(
    """
    <style>
    :root {
        --page-bg: #f5f7fb;
        --surface-bg: #ffffff;
        --surface-muted: #f8fafc;
        --text-primary: #0f172a;
        --text-secondary: #475569;
        --text-inverse: #f8fafc;
        --border-subtle: #dbe3ef;
        --border-strong: #94a3b8;
        --accent: #f97316;
        --accent-hover: #ea580c;
        --accent-soft: #fee2e2;
        --accent-text: #991b1b;
        --focus-ring: #fdba74;
        --control-bg: #172033;
        --control-hover-bg: #243047;
        --control-active-bg: #0f172a;
        --control-surface-bg: #ffffff;
        --control-surface-hover-bg: #f8fafc;
        --control-surface-active-bg: #eef2f7;
        --control-text: #0f172a;
        --control-hover-text: #0f172a;
        --control-border: #cbd5e1;
        --control-hover-border: #94a3b8;
        --control-primary-bg: #dc2626;
        --control-primary-hover-bg: #b91c1c;
        --control-primary-active-bg: #991b1b;
        --control-primary-text: #ffffff;
        --disabled-bg: #e2e8f0;
        --disabled-border: #cbd5e1;
        --disabled-text: #475569;
        --danger: #991b1b;
        --warning: #a16207;
        --success: #166534;
        --color-text-primary: var(--text-primary);
        --color-text-secondary: var(--text-secondary);
        --color-text-muted: var(--text-secondary);
        --color-text-inverse: var(--text-inverse);
        --color-surface-page: var(--page-bg);
        --color-surface-card: var(--surface-bg);
        --color-surface-control: var(--surface-muted);
        --color-control-text: var(--control-text);
        --color-control-hover-text: var(--control-hover-text);
        --color-border-default: var(--border-subtle);
        --color-border-strong: var(--border-strong);
        --color-accent: var(--accent);
        --color-accent-hover: var(--accent-hover);
        --color-accent-active: var(--control-active-bg);
        --color-disabled-text: var(--disabled-text);
        --color-disabled-background: var(--disabled-bg);
    }
    .stApp {
        background: var(--page-bg);
        color: var(--text-primary);
    }
    .block-container {
        padding-top: 2.15rem;
        padding-bottom: 2rem;
        max-width: 1320px;
    }
    [data-testid="stHorizontalBlock"] {
        gap: 1rem;
        align-items: stretch;
    }
    [data-testid="stHorizontalBlock"] > div[data-testid="column"] {
        min-width: 0;
    }
    [data-testid="stAppViewContainer"] main,
    [data-testid="stAppViewContainer"] main label,
    [data-testid="stAppViewContainer"] main p,
    [data-testid="stAppViewContainer"] main span {
        color: var(--text-primary);
        -webkit-text-fill-color: currentColor;
    }
    [data-testid="stAppViewContainer"] main div[data-testid="stButton"] button {
        display: inline-flex;
        align-items: center;
        justify-content: center;
        gap: 0.45rem;
        min-height: 42px;
        padding: 0.55rem 0.85rem;
        border-radius: 8px;
        border: 1px solid var(--control-border);
        background: var(--control-surface-bg);
        color: var(--control-text);
        -webkit-text-fill-color: var(--control-text);
        font-weight: 700;
        line-height: 1.2;
        box-shadow: none;
        opacity: 1;
        transition: background-color 120ms ease, border-color 120ms ease, color 120ms ease, transform 120ms ease, box-shadow 120ms ease;
    }
    [data-testid="stAppViewContainer"] main div[data-testid="stButton"] button *,
    [data-testid="stAppViewContainer"] main div[data-testid="stButton"] button p {
        display: inline-flex;
        align-items: center;
        justify-content: center;
        color: inherit !important;
        -webkit-text-fill-color: currentColor !important;
        line-height: inherit !important;
        margin: 0 !important;
    }
    [data-testid="stAppViewContainer"] main div[data-testid="stButton"] button:not(:disabled):hover {
        background: var(--control-surface-hover-bg);
        border-color: var(--control-hover-border);
        color: var(--control-hover-text);
        -webkit-text-fill-color: var(--control-hover-text);
        transform: translateY(-1px);
    }
    [data-testid="stAppViewContainer"] main div[data-testid="stButton"] button:not(:disabled):active {
        background: var(--control-surface-active-bg);
        color: var(--control-text);
        -webkit-text-fill-color: var(--control-text);
        transform: translateY(0);
    }
    [data-testid="stAppViewContainer"] main div[data-testid="stButton"] button:not([data-testid="stBaseButton-primary"]):not([kind="primary"]),
    [data-testid="stAppViewContainer"] main button[data-testid="stBaseButton-secondary"],
    [data-testid="stAppViewContainer"] main button[kind="secondary"] {
        background: var(--control-surface-bg) !important;
        background-color: var(--control-surface-bg) !important;
        border-color: var(--control-border) !important;
        color: var(--control-text) !important;
        -webkit-text-fill-color: var(--control-text) !important;
        opacity: 1 !important;
    }
    [data-testid="stAppViewContainer"] main div[data-testid="stButton"] button:not([data-testid="stBaseButton-primary"]):not([kind="primary"]) *,
    [data-testid="stAppViewContainer"] main button[data-testid="stBaseButton-secondary"] *,
    [data-testid="stAppViewContainer"] main button[kind="secondary"] * {
        color: var(--control-text) !important;
        -webkit-text-fill-color: var(--control-text) !important;
        opacity: 1 !important;
    }
    [data-testid="stAppViewContainer"] main div[data-testid="stButton"] button:not([data-testid="stBaseButton-primary"]):not([kind="primary"]):not(:disabled):hover,
    [data-testid="stAppViewContainer"] main div[data-testid="stButton"] button:not([data-testid="stBaseButton-primary"]):not([kind="primary"]):not(:disabled):focus-visible,
    [data-testid="stAppViewContainer"] main button[data-testid="stBaseButton-secondary"]:not(:disabled):hover,
    [data-testid="stAppViewContainer"] main button[data-testid="stBaseButton-secondary"]:not(:disabled):focus-visible,
    [data-testid="stAppViewContainer"] main button[kind="secondary"]:not(:disabled):hover,
    [data-testid="stAppViewContainer"] main button[kind="secondary"]:not(:disabled):focus-visible {
        background: var(--control-surface-hover-bg) !important;
        background-color: var(--control-surface-hover-bg) !important;
        border-color: var(--control-hover-border) !important;
        color: var(--control-hover-text) !important;
        -webkit-text-fill-color: var(--control-hover-text) !important;
    }
    [data-testid="stAppViewContainer"] main div[data-testid="stButton"] button:not([data-testid="stBaseButton-primary"]):not([kind="primary"]):not(:disabled):active,
    [data-testid="stAppViewContainer"] main button[data-testid="stBaseButton-secondary"]:not(:disabled):active,
    [data-testid="stAppViewContainer"] main button[kind="secondary"]:not(:disabled):active {
        background: var(--control-surface-active-bg) !important;
        background-color: var(--control-surface-active-bg) !important;
        border-color: var(--control-hover-border) !important;
        color: var(--control-text) !important;
        -webkit-text-fill-color: var(--control-text) !important;
    }
    button[data-testid="stBaseButton-primary"],
    [data-testid="stAppViewContainer"] main button[data-testid="stBaseButton-primary"],
    [data-testid="stAppViewContainer"] main div[data-testid="stButton"] button[data-testid="stBaseButton-primary"],
    [data-testid="stAppViewContainer"] main div[data-testid="stButton"] button[kind="primary"] {
        background: var(--control-primary-bg) !important;
        border-color: var(--control-primary-bg) !important;
        color: var(--control-primary-text) !important;
        -webkit-text-fill-color: var(--control-primary-text) !important;
    }
    button[data-testid="stBaseButton-primary"]:not(:disabled):hover,
    [data-testid="stAppViewContainer"] main button[data-testid="stBaseButton-primary"]:not(:disabled):hover,
    [data-testid="stAppViewContainer"] main div[data-testid="stButton"] button[data-testid="stBaseButton-primary"]:not(:disabled):hover,
    [data-testid="stAppViewContainer"] main div[data-testid="stButton"] button[kind="primary"]:not(:disabled):hover {
        background: var(--control-primary-hover-bg) !important;
        border-color: var(--control-primary-hover-bg) !important;
        color: var(--control-primary-text) !important;
        -webkit-text-fill-color: var(--control-primary-text) !important;
    }
    button[data-testid="stBaseButton-primary"]:not(:disabled):active,
    [data-testid="stAppViewContainer"] main button[data-testid="stBaseButton-primary"]:not(:disabled):active,
    [data-testid="stAppViewContainer"] main div[data-testid="stButton"] button[data-testid="stBaseButton-primary"]:not(:disabled):active,
    [data-testid="stAppViewContainer"] main div[data-testid="stButton"] button[kind="primary"]:not(:disabled):active {
        background: var(--control-primary-active-bg) !important;
        border-color: var(--control-primary-active-bg) !important;
    }
    [data-testid="stAppViewContainer"] main div[data-testid="stButton"] button:focus-visible,
    [data-testid="stAppViewContainer"] main input:focus-visible,
    [data-testid="stAppViewContainer"] main textarea:focus-visible,
    [data-testid="stAppViewContainer"] main [role="radio"]:focus-visible,
    [data-testid="stAppViewContainer"] main [role="checkbox"]:focus-visible {
        outline: 2px solid var(--focus-ring) !important;
        outline-offset: 2px !important;
        box-shadow: 0 0 0 4px rgba(253, 186, 116, 0.32) !important;
    }
    [data-testid="stAppViewContainer"] main div[data-testid="stButton"] button:disabled,
    [data-testid="stAppViewContainer"] main div[data-testid="stButton"] button[disabled],
    button[data-testid="stBaseButton-secondary"]:disabled,
    button[data-testid="stBaseButton-secondary"][disabled],
    [data-testid="stAppViewContainer"] main button[data-testid="stBaseButton-secondary"]:disabled,
    [data-testid="stAppViewContainer"] main button[data-testid="stBaseButton-secondary"][disabled] {
        background: var(--disabled-bg) !important;
        background-color: var(--disabled-bg) !important;
        border-color: var(--disabled-border) !important;
        color: var(--disabled-text) !important;
        -webkit-text-fill-color: var(--disabled-text) !important;
        opacity: 1 !important;
        cursor: not-allowed !important;
        transform: none !important;
    }
    [data-testid="stAppViewContainer"] main [data-testid="stCheckbox"] label,
    [data-testid="stAppViewContainer"] main [data-testid="stRadio"] label,
    [data-testid="stAppViewContainer"] main [data-testid="stSlider"] label,
    [data-testid="stAppViewContainer"] main [data-testid="stNumberInput"] label {
        color: var(--text-primary) !important;
        -webkit-text-fill-color: var(--text-primary) !important;
        font-weight: 650;
    }
    [data-testid="stAppViewContainer"] main [data-testid="stWidgetLabel"],
    [data-testid="stAppViewContainer"] main [data-testid="stWidgetLabel"] *,
    [data-testid="stAppViewContainer"] main [data-testid="stNumberInput"] [data-testid="stWidgetLabel"] *,
    [data-testid="stAppViewContainer"] main [data-testid="stSlider"] [data-testid="stWidgetLabel"] *,
    [data-testid="stAppViewContainer"] main [data-testid="stCheckbox"] [data-testid="stWidgetLabel"] * {
        color: var(--text-primary) !important;
        -webkit-text-fill-color: var(--text-primary) !important;
        opacity: 1 !important;
    }
    [data-testid="stAppViewContainer"] main [data-testid="stCheckbox"] span,
    [data-testid="stAppViewContainer"] main [data-testid="stCheckbox"] p,
    [data-testid="stAppViewContainer"] main [data-testid="stRadio"] span,
    [data-testid="stAppViewContainer"] main [data-testid="stRadio"] p,
    [data-testid="stAppViewContainer"] main [data-testid="stSlider"] span,
    [data-testid="stAppViewContainer"] main [data-testid="stSlider"] p {
        color: var(--text-primary) !important;
        -webkit-text-fill-color: var(--text-primary) !important;
        opacity: 1 !important;
    }
    [data-testid="stAppViewContainer"] main [data-testid="stNumberInput"] [data-baseweb="input"],
    [data-testid="stAppViewContainer"] main [data-testid="stNumberInput"] div[data-baseweb="base-input"],
    [data-testid="stAppViewContainer"] main [data-testid="stNumberInput"] input {
        background: var(--surface-bg) !important;
        background-color: var(--surface-bg) !important;
        color: var(--text-primary) !important;
        -webkit-text-fill-color: var(--text-primary) !important;
        border-color: var(--border-subtle) !important;
        opacity: 1 !important;
    }
    [data-testid="stAppViewContainer"] main [data-testid="stNumberInput"] button,
    [data-testid="stAppViewContainer"] main [data-testid="stNumberInput"] button[kind],
    [data-testid="stAppViewContainer"] main [data-testid="stNumberInput"] [role="button"] {
        background: var(--surface-muted) !important;
        background-color: var(--surface-muted) !important;
        border-color: var(--border-subtle) !important;
        color: var(--text-primary) !important;
        -webkit-text-fill-color: var(--text-primary) !important;
        opacity: 1 !important;
    }
    [data-testid="stAppViewContainer"] main [data-testid="stNumberInput"] button *,
    [data-testid="stAppViewContainer"] main [data-testid="stNumberInput"] [role="button"] *,
    [data-testid="stAppViewContainer"] main [data-testid="stNumberInput"] svg {
        color: var(--text-primary) !important;
        fill: currentColor !important;
        -webkit-text-fill-color: var(--text-primary) !important;
        opacity: 1 !important;
    }
    [data-testid="stAppViewContainer"] main input,
    [data-testid="stAppViewContainer"] main textarea {
        background: var(--surface-bg) !important;
        color: var(--text-primary) !important;
        -webkit-text-fill-color: var(--text-primary) !important;
        border-color: var(--border-subtle) !important;
    }
    section[data-testid="stSidebar"] {
        background: #172033;
        border-right: 1px solid #25324c;
    }
    section[data-testid="stSidebar"] * {
        color: #f7f9fc;
    }
    section[data-testid="stSidebar"] label,
    section[data-testid="stSidebar"] p {
        color: #f7f9fc !important;
    }
    section[data-testid="stSidebar"] div[data-baseweb="select"] {
        background-color: #0b1020 !important;
        color: #f8fafc !important;
    }
    section[data-testid="stSidebar"] div[data-baseweb="select"] > div {
        display: flex !important;
        align-items: center !important;
        background-color: #0b1020 !important;
        color: #f8fafc !important;
        border: 1px solid #243047 !important;
        border-radius: 10px !important;
        min-height: 44px !important;
        box-shadow: none !important;
    }
    section[data-testid="stSidebar"] div[data-baseweb="select"] [data-baseweb="value-container"],
    section[data-testid="stSidebar"] div[data-baseweb="select"] [data-baseweb="single-value"],
    section[data-testid="stSidebar"] div[data-baseweb="select"] [data-baseweb="placeholder"] {
        display: flex !important;
        align-items: center !important;
        min-height: 42px !important;
        line-height: 1.2 !important;
        padding-top: 0 !important;
        padding-bottom: 0 !important;
    }
    section[data-testid="stSidebar"] div[data-baseweb="select"] [data-baseweb="select-dropdown"],
    section[data-testid="stSidebar"] div[data-baseweb="select"] [data-baseweb="select-dropdown"] > div {
        display: flex !important;
        align-items: center !important;
        justify-content: center !important;
    }
    section[data-testid="stSidebar"] div[data-baseweb="select"] div,
    section[data-testid="stSidebar"] div[data-baseweb="select"] span {
        color: #f8fafc !important;
    }
    section[data-testid="stSidebar"] div[data-baseweb="select"] input {
        color: #f8fafc !important;
        -webkit-text-fill-color: #f8fafc !important;
        line-height: 1.2 !important;
    }
    section[data-testid="stSidebar"] div[data-baseweb="select"] input::placeholder {
        color: #94a3b8 !important;
        -webkit-text-fill-color: #94a3b8 !important;
        opacity: 1 !important;
    }
    section[data-testid="stSidebar"] div[data-baseweb="select"] svg {
        color: #cbd5e1 !important;
        fill: #cbd5e1 !important;
    }
    div[data-baseweb="popover"] ul,
    div[data-baseweb="popover"] div[role="listbox"] {
        background-color: #111827 !important;
        border: 1px solid #243047 !important;
    }
    div[data-baseweb="popover"] li,
    div[data-baseweb="popover"] div[role="option"] {
        display: flex !important;
        align-items: center !important;
        min-height: 36px !important;
        line-height: 1.25 !important;
        background-color: #111827 !important;
        color: #f8fafc !important;
    }
    div[data-baseweb="popover"] li:hover,
    div[data-baseweb="popover"] div[role="option"]:hover,
    div[data-baseweb="popover"] div[aria-selected="true"] {
        background-color: #1f2937 !important;
        color: #ffffff !important;
    }
    section[data-testid="stSidebar"] [data-testid="stCaptionContainer"] p {
        color: #bdc9dc !important;
        font-size: 0.82rem;
        line-height: 1.35;
    }
    section[data-testid="stSidebar"] [data-testid="stMarkdownContainer"] li {
        color: #dbe5f4 !important;
    }
    section[data-testid="stSidebar"] div[data-testid="stExpander"] {
        border-color: #3a4865;
        background: rgba(255,255,255,0.04);
        border-radius: 8px;
    }
    [data-testid="stFileUploader"] section {
        background: #ffffff !important;
        border: 1px dashed #9fb0c9 !important;
        border-radius: 8px !important;
        color: #111827 !important;
    }
    [data-testid="stFileUploader"] section * {
        color: #111827 !important;
    }
    [data-testid="stFileUploader"] button {
        background: #172033 !important;
        border: 1px solid #2f3c56 !important;
        color: #ffffff !important;
        border-radius: 8px !important;
    }
    [data-testid="stFileUploader"] button * {
        color: #ffffff !important;
    }
    [data-testid="stFileChips"] {
        display: flex !important;
        align-items: center !important;
        gap: 0.75rem !important;
        flex-wrap: wrap !important;
    }
    [data-testid="stFileChip"],
    .stFileChip,
    [data-testid="stFileUploaderFile"],
    .stFileUploaderFile,
    [class*="stFileUploaderFile"] {
        display: flex !important;
        align-items: center !important;
        gap: 0.75rem !important;
        min-width: min(320px, 100%) !important;
        max-width: min(520px, 100%) !important;
        padding: 0.65rem 0.75rem !important;
        background: #ffffff !important;
        border: 1px solid #dbe3ef !important;
        border-radius: 12px !important;
        box-shadow: 0 1px 2px rgba(15, 23, 42, 0.06) !important;
        color: #111827 !important;
        -webkit-text-fill-color: #111827 !important;
    }
    [data-testid="stFileChip"] *,
    .stFileChip *,
    [data-testid="stFileUploaderFile"] *,
    .stFileUploaderFile *,
    [class*="stFileUploaderFile"] * {
        color: #111827 !important;
        -webkit-text-fill-color: #111827 !important;
        opacity: 1 !important;
    }
    [data-testid="stFileChip"] > div:first-child,
    .stFileChip > div:first-child {
        width: 2.75rem !important;
        height: 2.75rem !important;
        background: #f8fafc !important;
        color: #334155 !important;
        -webkit-text-fill-color: #334155 !important;
        border-radius: 12px !important;
        border: 1px solid #cbd5e1 !important;
        box-shadow: inset 0 0 0 1px rgba(255, 255, 255, 0.75) !important;
    }
    [data-testid="stFileChip"] > div:first-child *,
    .stFileChip > div:first-child *,
    [data-testid="stFileUploaderFile"] svg,
    .stFileUploaderFile svg,
    [class*="stFileUploaderFile"] svg {
        color: #334155 !important;
        fill: currentColor !important;
        -webkit-text-fill-color: #334155 !important;
    }
    [data-testid="stFileChip"] > div:first-child svg,
    .stFileChip > div:first-child svg {
        width: 1.35rem !important;
        height: 1.35rem !important;
        color: #334155 !important;
        fill: currentColor !important;
    }
    [data-testid="stFileChipName"],
    .stFileChipName,
    [data-testid="stFileUploaderFileName"],
    .stFileUploaderFileName,
    [class*="stFileUploaderFileName"] {
        color: #111827 !important;
        -webkit-text-fill-color: #111827 !important;
        font-weight: 650 !important;
        min-width: 0 !important;
        max-width: 360px !important;
        overflow: hidden !important;
        text-overflow: ellipsis !important;
        white-space: nowrap !important;
    }
    [data-testid="stFileChipName"] + div,
    .stFileChipName + div,
    [data-testid="stFileChip"] small,
    [data-testid="stFileUploaderFileSize"],
    .stFileUploaderFileSize,
    [class*="stFileUploaderFileSize"] {
        color: #64748b !important;
        -webkit-text-fill-color: #64748b !important;
        font-size: 0.8rem !important;
    }
    [data-testid="stFileChipDeleteBtn"] button,
    [data-testid="stFileUploaderDeleteBtn"],
    [data-testid="stFileUploaderFile"] button,
    .stFileUploaderFile button,
    [class*="stFileUploaderFile"] button {
        background: #f1f5f9 !important;
        border: 1px solid #dbe3ef !important;
        color: #334155 !important;
        border-radius: 999px !important;
    }
    [data-testid="stFileChipDeleteBtn"] button *,
    [data-testid="stFileUploaderDeleteBtn"] *,
    [data-testid="stFileUploaderFile"] button *,
    .stFileUploaderFile button *,
    [class*="stFileUploaderFile"] button * {
        color: #334155 !important;
        -webkit-text-fill-color: #334155 !important;
    }
    [data-testid="stFileChipDeleteBtn"] button:hover,
    [data-testid="stFileUploaderDeleteBtn"]:hover,
    [data-testid="stFileUploaderFile"] button:hover,
    .stFileUploaderFile button:hover,
    [class*="stFileUploaderFile"] button:hover {
        background: #fee2e2 !important;
        border-color: #fecaca !important;
        color: #991b1b !important;
    }
    .summary-grid {
        display: grid;
        grid-template-columns: repeat(auto-fit, minmax(180px, 1fr));
        gap: 0.75rem;
        margin: 1rem 0 0.9rem;
    }
    .summary-card {
        background: #ffffff;
        border: 1px solid #d8dee8;
        border-radius: 8px;
        padding: 0.78rem 0.9rem;
        min-height: 104px;
        box-shadow: 0 1px 2px rgba(15, 23, 42, 0.04);
    }
    .summary-card-title {
        color: #334155;
        font-size: 0.82rem;
        font-weight: 700;
        line-height: 1.3;
        margin-bottom: 0.45rem;
    }
    .summary-card-value {
        color: #0f172a;
        font-size: 1.42rem;
        font-weight: 760;
        line-height: 1.15;
        margin-bottom: 0.5rem;
    }
    .summary-card-caption {
        color: #475569;
        font-size: 0.78rem;
        line-height: 1.35;
    }
    .readable-table-wrap {
        overflow-x: auto;
        margin: 0.55rem 0 1rem;
        border: 1px solid #e2e8f0;
        border-radius: 8px;
        background: #ffffff;
    }
    .readable-table {
        width: 100%;
        table-layout: fixed;
        border-collapse: collapse;
        background: #ffffff;
        color: #0f172a;
        font-size: 0.88rem;
    }
    .readable-table th {
        background: #f8fafc;
        color: #0f172a;
        font-weight: 740;
        text-align: left;
        border-bottom: 1px solid #e2e8f0;
        padding: 0.72rem 0.8rem;
        white-space: normal;
        overflow-wrap: anywhere;
        word-break: break-word;
    }
    .readable-table td {
        color: #0f172a;
        border-bottom: 1px solid #eef2f7;
        padding: 0.68rem 0.8rem;
        vertical-align: top;
        line-height: 1.42;
        white-space: normal;
        overflow-wrap: anywhere;
        word-break: break-word;
    }
    .readable-table tbody tr:last-child td {
        border-bottom: none;
    }
    .readable-table tbody tr:nth-child(even) td {
        background: #fbfdff;
    }
    h1, h2, h3 {
        letter-spacing: 0;
    }
    h1 {
        font-size: 2rem;
        line-height: 1.18;
        margin-top: 0;
        margin-bottom: 0.25rem;
        color: #111827;
        white-space: normal;
        overflow: visible;
    }
    [data-testid="stAppViewContainer"] main h1,
    [data-testid="stAppViewContainer"] main h2,
    [data-testid="stAppViewContainer"] main h3,
    [data-testid="stAppViewContainer"] main [data-testid="stHeader"],
    [data-testid="stAppViewContainer"] main [data-testid="stHeading"],
    [data-testid="stAppViewContainer"] main [data-testid="stMarkdownContainer"] h1,
    [data-testid="stAppViewContainer"] main [data-testid="stMarkdownContainer"] h2,
    [data-testid="stAppViewContainer"] main [data-testid="stMarkdownContainer"] h3,
    section[data-testid="stMain"] h1,
    section[data-testid="stMain"] h2,
    section[data-testid="stMain"] h3,
    section[data-testid="stMain"] [data-testid="stHeading"] h1,
    section[data-testid="stMain"] [data-testid="stHeading"] h2,
    section[data-testid="stMain"] [data-testid="stHeading"] h3 {
        color: #0f172a !important;
        -webkit-text-fill-color: #0f172a !important;
        opacity: 1 !important;
        line-height: 1.18 !important;
        overflow: visible !important;
        white-space: normal !important;
    }
    [data-testid="stMetric"] {
        border: 1px solid #d8dee8;
        border-radius: 8px;
        padding: 0.7rem 0.85rem;
        background: #ffffff;
    }
    [data-testid="stAppViewContainer"] main [data-testid="stMetric"] * {
        color: #111827 !important;
        -webkit-text-fill-color: #111827 !important;
    }
    [data-testid="stAppViewContainer"] main [data-testid="stMetricLabel"] *,
    [data-testid="stAppViewContainer"] main [data-testid="stCaptionContainer"] p {
        color: #475569 !important;
        -webkit-text-fill-color: #475569 !important;
    }
    [data-testid="stMetricValue"] {
        font-size: 1.45rem;
        color: #111827 !important;
        -webkit-text-fill-color: #111827 !important;
    }
    [data-testid="stAppViewContainer"] main [data-testid="stMarkdownContainer"] p,
    [data-testid="stAppViewContainer"] main [data-testid="stMarkdownContainer"] li {
        color: #111827 !important;
        -webkit-text-fill-color: #111827 !important;
    }
    [data-testid="stAppViewContainer"] main [data-testid="stTable"] *,
    [data-testid="stAppViewContainer"] main [data-testid="stDataFrame"] *,
    [data-testid="stAppViewContainer"] main div[data-baseweb="tab"] * {
        color: #111827 !important;
        -webkit-text-fill-color: #111827 !important;
    }
    [data-testid="stAppViewContainer"] main [data-testid="stTabs"] div[role="tablist"] {
        display: inline-flex !important;
        gap: 0.25rem !important;
        background: #ffffff !important;
        border: 1px solid #e2e8f0 !important;
        border-radius: 12px !important;
        padding: 0.38rem !important;
        box-shadow: 0 1px 2px rgba(15, 23, 42, 0.04) !important;
    }
    .stTabs div[role="tablist"],
    div[data-baseweb="tab-list"] {
        display: inline-flex !important;
        gap: 0.25rem !important;
        background: #ffffff !important;
        border: 1px solid #e2e8f0 !important;
        border-radius: 12px !important;
        padding: 0.38rem !important;
        box-shadow: 0 1px 2px rgba(15, 23, 42, 0.04) !important;
    }
    [data-testid="stAppViewContainer"] main [data-testid="stTabs"] button[role="tab"] {
        border-radius: 10px !important;
        border: 1px solid transparent !important;
        padding: 0.42rem 0.78rem !important;
        background: transparent !important;
        min-height: 36px !important;
    }
    .stTabs button[role="tab"],
    div[data-baseweb="tab"] {
        border-radius: 10px !important;
        border: 1px solid transparent !important;
        padding: 0.42rem 0.78rem !important;
        background: transparent !important;
        min-height: 36px !important;
    }
    [data-testid="stAppViewContainer"] main [data-testid="stTabs"] button[role="tab"],
    [data-testid="stAppViewContainer"] main [data-testid="stTabs"] button[role="tab"] p,
    .stTabs button[role="tab"],
    .stTabs button[role="tab"] p,
    div[data-baseweb="tab"],
    div[data-baseweb="tab"] p {
        color: #334155 !important;
        -webkit-text-fill-color: #334155 !important;
        font-weight: 650 !important;
    }
    [data-testid="stAppViewContainer"] main [data-testid="stTabs"] button[role="tab"]:hover,
    [data-testid="stAppViewContainer"] main [data-testid="stTabs"] button[role="tab"]:focus-visible {
        background: #f8fafc !important;
        border-color: #cbd5e1 !important;
        outline: none !important;
    }
    .stTabs button[role="tab"]:hover,
    .stTabs button[role="tab"]:focus-visible,
    div[data-baseweb="tab"]:hover {
        background: #f8fafc !important;
        border-color: #cbd5e1 !important;
        outline: none !important;
    }
    [data-testid="stAppViewContainer"] main [data-testid="stTabs"] button[role="tab"][aria-selected="true"],
    [data-testid="stAppViewContainer"] main [data-testid="stTabs"] button[role="tab"][aria-selected="true"] p,
    .stTabs button[role="tab"][aria-selected="true"],
    .stTabs button[role="tab"][aria-selected="true"] p,
    div[data-baseweb="tab"][aria-selected="true"],
    div[data-baseweb="tab"][aria-selected="true"] p {
        background: #fee2e2 !important;
        border-color: #fecaca !important;
        color: #991b1b !important;
        -webkit-text-fill-color: #991b1b !important;
    }
    [data-testid="stAppViewContainer"] main [data-testid="stTabs"] div[data-baseweb="tab-highlight"] {
        background-color: transparent !important;
    }
    div[data-baseweb="tab-highlight"] {
        background-color: transparent !important;
    }
    [data-testid="stDownloadButton"] button,
    div.stDownloadButton[data-testid="stDownloadButton"] button,
    [data-testid="stAppViewContainer"] [data-testid="stDownloadButton"] button {
        background: #f97316 !important;
        background-color: #f97316 !important;
        border: 1px solid #ea580c !important;
        color: #ffffff !important;
        -webkit-text-fill-color: #ffffff !important;
        border-radius: 8px !important;
        min-height: 44px !important;
        opacity: 1 !important;
        box-shadow: none !important;
        transition: background-color 120ms ease, border-color 120ms ease, transform 120ms ease, box-shadow 120ms ease !important;
    }
    [data-testid="stDownloadButton"] button *,
    div.stDownloadButton[data-testid="stDownloadButton"] button *,
    [data-testid="stAppViewContainer"] [data-testid="stDownloadButton"] button * {
        color: #ffffff !important;
        -webkit-text-fill-color: #ffffff !important;
        opacity: 1 !important;
    }
    [data-testid="stDownloadButton"] button:not(:disabled):hover,
    div.stDownloadButton[data-testid="stDownloadButton"] button:not(:disabled):hover,
    [data-testid="stAppViewContainer"] [data-testid="stDownloadButton"] button:not(:disabled):hover {
        background: #ea580c !important;
        background-color: #ea580c !important;
        border-color: #c2410c !important;
        color: #ffffff !important;
        -webkit-text-fill-color: #ffffff !important;
        box-shadow: 0 8px 18px rgba(234, 88, 12, 0.2) !important;
        transform: scale(1.01) !important;
    }
    [data-testid="stDownloadButton"] button:not(:disabled):active,
    div.stDownloadButton[data-testid="stDownloadButton"] button:not(:disabled):active,
    [data-testid="stAppViewContainer"] [data-testid="stDownloadButton"] button:not(:disabled):active {
        background: #c2410c !important;
        background-color: #c2410c !important;
        border-color: #9a3412 !important;
        color: #ffffff !important;
        -webkit-text-fill-color: #ffffff !important;
        transform: scale(0.98) !important;
    }
    [data-testid="stDownloadButton"] button:focus-visible,
    div.stDownloadButton[data-testid="stDownloadButton"] button:focus-visible,
    [data-testid="stAppViewContainer"] [data-testid="stDownloadButton"] button:focus-visible {
        outline: 2px solid #fdba74 !important;
        outline-offset: 2px !important;
        box-shadow: 0 0 0 4px rgba(253, 186, 116, 0.35) !important;
    }
    [data-testid="stDownloadButton"] button:disabled,
    [data-testid="stDownloadButton"] button[disabled],
    div.stDownloadButton[data-testid="stDownloadButton"] button:disabled,
    div.stDownloadButton[data-testid="stDownloadButton"] button[disabled],
    [data-testid="stAppViewContainer"] [data-testid="stDownloadButton"] button:disabled,
    [data-testid="stAppViewContainer"] [data-testid="stDownloadButton"] button[disabled] {
        background: #cbd5e1 !important;
        background-color: #cbd5e1 !important;
        border-color: #94a3b8 !important;
        color: #334155 !important;
        -webkit-text-fill-color: #334155 !important;
        opacity: 1 !important;
        transform: none !important;
        cursor: not-allowed !important;
        box-shadow: none !important;
    }
    [data-testid="stDownloadButton"] button:disabled:hover,
    [data-testid="stDownloadButton"] button[disabled]:hover,
    div.stDownloadButton[data-testid="stDownloadButton"] button:disabled:hover,
    div.stDownloadButton[data-testid="stDownloadButton"] button[disabled]:hover,
    [data-testid="stAppViewContainer"] [data-testid="stDownloadButton"] button:disabled:hover,
    [data-testid="stAppViewContainer"] [data-testid="stDownloadButton"] button[disabled]:hover {
        background: #cbd5e1 !important;
        background-color: #cbd5e1 !important;
        border-color: #94a3b8 !important;
        color: #334155 !important;
        -webkit-text-fill-color: #334155 !important;
        opacity: 1 !important;
        transform: none !important;
    }
    [data-testid="stDownloadButton"] button:disabled *,
    [data-testid="stDownloadButton"] button[disabled] *,
    div.stDownloadButton[data-testid="stDownloadButton"] button:disabled *,
    div.stDownloadButton[data-testid="stDownloadButton"] button[disabled] *,
    [data-testid="stAppViewContainer"] [data-testid="stDownloadButton"] button:disabled *,
    [data-testid="stAppViewContainer"] [data-testid="stDownloadButton"] button[disabled] * {
        color: #334155 !important;
        -webkit-text-fill-color: #334155 !important;
        opacity: 1 !important;
    }
    body [data-testid="stDownloadButton"] button:disabled,
    body [data-testid="stDownloadButton"] button[disabled],
    body div.stDownloadButton[data-testid="stDownloadButton"] button:disabled,
    body div.stDownloadButton[data-testid="stDownloadButton"] button[disabled] {
        background: #cbd5e1 !important;
        background-color: #cbd5e1 !important;
        background-image: none !important;
        border-color: #94a3b8 !important;
        color: #334155 !important;
        -webkit-text-fill-color: #334155 !important;
        opacity: 1 !important;
        box-shadow: none !important;
        transform: none !important;
    }
    [data-testid="stAppViewContainer"] main [data-testid="stExpander"] {
        background: #ffffff !important;
        border: 1px solid #e2e8f0 !important;
        border-radius: 8px !important;
        overflow: hidden !important;
    }
    [data-testid="stAppViewContainer"] main [data-testid="stExpander"] details,
    [data-testid="stAppViewContainer"] main [data-testid="stExpander"] summary {
        background: #ffffff !important;
        color: #0f172a !important;
        -webkit-text-fill-color: #0f172a !important;
        opacity: 1 !important;
    }
    [data-testid="stAppViewContainer"] main [data-testid="stExpander"] summary:hover {
        background: #f1f5f9 !important;
    }
    [data-testid="stAppViewContainer"] main [data-testid="stExpander"] summary *,
    [data-testid="stAppViewContainer"] main [data-testid="stExpander"] [data-testid="stMarkdownContainer"],
    [data-testid="stAppViewContainer"] main [data-testid="stExpander"] [data-testid="stMarkdownContainer"] *,
    [data-testid="stAppViewContainer"] main [data-testid="stExpander"] p,
    [data-testid="stAppViewContainer"] main [data-testid="stExpander"] li,
    [data-testid="stAppViewContainer"] main [data-testid="stExpander"] span {
        color: #334155 !important;
        -webkit-text-fill-color: #334155 !important;
        opacity: 1 !important;
    }
    [data-testid="stAppViewContainer"] main [data-testid="stExpander"] summary p,
    [data-testid="stAppViewContainer"] main [data-testid="stExpander"] summary span {
        color: #0f172a !important;
        -webkit-text-fill-color: #0f172a !important;
        font-weight: 700 !important;
    }
    .stats-help-panel {
        margin: 1rem 0 1.25rem 0;
        border: 1px solid #cbd5e1;
        border-radius: 12px;
        background: #ffffff;
        box-shadow: 0 1px 2px rgba(15, 23, 42, 0.06);
        overflow: hidden;
        color: #334155;
    }
    .stats-help-panel summary {
        display: flex;
        align-items: center;
        gap: 0.65rem;
        min-height: 48px;
        padding: 0.85rem 1rem;
        list-style: none;
        cursor: pointer;
        background: #f1f5f9;
        color: #1e293b !important;
        -webkit-text-fill-color: #1e293b !important;
        font-weight: 700;
        opacity: 1 !important;
        transition: background-color 120ms ease, color 120ms ease;
    }
    .stats-help-panel summary::-webkit-details-marker {
        display: none;
    }
    .stats-help-panel summary:hover {
        background: #e2e8f0;
        color: #0f172a !important;
        -webkit-text-fill-color: #0f172a !important;
    }
    .stats-help-panel[open] summary {
        background: #e2e8f0;
        border-bottom: 1px solid #cbd5e1;
    }
    .stats-help-chevron {
        display: inline-flex;
        align-items: center;
        justify-content: center;
        width: 1rem;
        color: #475569 !important;
        -webkit-text-fill-color: #475569 !important;
        font-size: 1.35rem;
        line-height: 1;
        transform: rotate(0deg);
        transition: transform 120ms ease;
    }
    .stats-help-panel[open] .stats-help-chevron {
        transform: rotate(90deg);
    }
    .stats-help-title {
        color: #1e293b !important;
        -webkit-text-fill-color: #1e293b !important;
        opacity: 1 !important;
    }
    .stats-help-body {
        padding: 1rem 1.25rem 1.1rem 1.25rem;
        background: #ffffff;
        color: #334155 !important;
        -webkit-text-fill-color: #334155 !important;
    }
    .stats-help-body p,
    .stats-help-body li,
    .stats-help-body span,
    .stats-help-body strong {
        color: #334155 !important;
        -webkit-text-fill-color: #334155 !important;
        opacity: 1 !important;
    }
    .stats-help-body p {
        margin: 0 0 0.8rem 0;
        line-height: 1.65;
    }
    .stats-help-body ul {
        margin: 0;
        padding-left: 1.25rem;
    }
    .stats-help-body li {
        margin: 0.35rem 0;
        line-height: 1.6;
    }
    .status-strip {
        border: 1px solid #cbd5e1;
        border-left: 4px solid #2563eb;
        border-radius: 8px;
        padding: 0.85rem 1rem;
        background: #ffffff;
        color: #162033;
    }
    .preview-action-band {
        display: flex;
        align-items: center;
        justify-content: space-between;
        gap: 1rem;
        border: 1px solid var(--border-subtle);
        border-radius: 8px;
        padding: 0.9rem 1rem;
        background: var(--surface-bg);
        margin: 0.8rem 0 0.55rem;
    }
    .preview-action-title {
        color: var(--text-primary);
        font-size: 1rem;
        font-weight: 760;
        line-height: 1.25;
    }
    .preview-action-desc {
        color: var(--text-secondary);
        font-size: 0.86rem;
        line-height: 1.35;
        margin-top: 0.2rem;
    }
    .preview-action-meta {
        display: flex;
        align-items: center;
        justify-content: flex-end;
        gap: 0.45rem;
        flex-wrap: wrap;
        min-width: 220px;
    }
    .preview-action-meta span {
        display: inline-flex;
        align-items: center;
        min-height: 28px;
        padding: 0.18rem 0.5rem;
        border-radius: 999px;
        border: 1px solid var(--border-subtle);
        background: var(--surface-muted);
        color: var(--text-secondary);
        -webkit-text-fill-color: var(--text-secondary);
        font-size: 0.76rem;
        font-weight: 650;
        white-space: normal;
    }
    .preview-route-top-spacer {
        height: 4rem;
    }
    .preview-workspace-title {
        color: var(--text-primary);
        font-size: 1.25rem;
        font-weight: 780;
        text-align: center;
        line-height: 1.35;
        padding-top: 0.35rem;
    }
    .preview-columns {
        display: grid;
        grid-template-columns: repeat(2, minmax(0, 1fr));
        gap: 24px;
        align-items: start;
        margin-top: 0.8rem;
    }
    .preview-pane {
        min-width: 0;
    }
    .preview-pane-title {
        color: var(--text-primary);
        -webkit-text-fill-color: var(--text-primary);
        font-size: 0.9rem;
        font-weight: 760;
        line-height: 1.35;
        margin: 0 0 0.45rem;
    }
    .preview-canvas-container {
        width: 100%;
        min-width: 0;
        overflow: auto;
    }
    .preview-zoom-frame {
        width: 100%;
        overflow-x: auto;
        overflow-y: visible;
        padding: 0.25rem 0 0.75rem;
    }
    .preview-zoom-frame .preview-page-shell {
        zoom: var(--preview-zoom);
        width: min(calc(100% / var(--preview-zoom)), 980px);
    }
    .preview-page-shell {
        position: relative;
        width: min(100%, 980px);
        margin: 0 auto 1rem;
        background: var(--surface-bg);
        border: 1px solid var(--border-strong);
        border-radius: 4px;
        overflow: hidden;
        box-shadow: 0 14px 34px rgba(15, 23, 42, 0.16);
    }
    .preview-page-image {
        display: block;
        width: 100%;
        height: auto;
        background: #ffffff;
    }
    .preview-overlay {
        position: absolute;
        inset: 0;
        pointer-events: none;
    }
    .preview-marker {
        position: absolute;
        border-radius: 5px;
        border: 2px solid #2563eb;
        background: rgba(37, 99, 235, 0.16);
        box-shadow: 0 0 0 2px rgba(255, 255, 255, 0.72);
        pointer-events: none;
    }
    .preview-marker-error {
        border-color: #dc2626;
        background: rgba(220, 38, 38, 0.17);
    }
    .preview-marker-warning {
        border-color: #f97316;
        background: rgba(249, 115, 22, 0.18);
    }
    .preview-marker-change {
        border-color: #16a34a;
        background: rgba(22, 163, 74, 0.12);
    }
    @media (max-width: 900px) {
        .preview-action-band {
            align-items: flex-start;
            flex-direction: column;
        }
        .preview-action-meta {
            justify-content: flex-start;
            min-width: 0;
        }
    }
    @media (max-width: 1000px) {
        .preview-columns {
            grid-template-columns: 1fr;
        }
    }
    @media (max-width: 900px) {
        [data-testid="stHorizontalBlock"] {
            flex-wrap: wrap;
        }
        [data-testid="stHorizontalBlock"] > div[data-testid="column"] {
            flex: 1 1 220px;
            min-width: min(220px, 100%);
        }
    }
    .qa-band {
        border: 1px solid #d8dee8;
        border-radius: 8px;
        padding: 1rem;
        background: #ffffff;
    }
    .st-key-preview_upload_state_keeper {
        display: none !important;
    }
    .st-key-preview_previous_page button,
    .st-key-preview_previous_page button :is(p, span, div),
    .st-key-preview_previous_page button:disabled,
    .st-key-preview_previous_page button[disabled],
    .st-key-preview_previous_page button[aria-disabled="true"],
    .st-key-preview_previous_page button:disabled :is(p, span, div),
    .st-key-preview_previous_page button[disabled] :is(p, span, div),
    .st-key-preview_previous_page button[aria-disabled="true"] :is(p, span, div) {
        color: #FFFFFF !important;
        -webkit-text-fill-color: #FFFFFF !important;
        opacity: 1 !important;
    }
    .st-key-preview_previous_page button,
    .st-key-preview_previous_page button:disabled,
    .st-key-preview_previous_page button[disabled],
    .st-key-preview_previous_page button[aria-disabled="true"] {
        background: #172033 !important;
        background-color: #172033 !important;
    }
    .st-key-preview_previous_page button:not(:disabled):hover {
        background: #243047 !important;
        background-color: #243047 !important;
    }
    .st-key-preview_previous_page button:not(:disabled):active {
        background: #0F172A !important;
        background-color: #0F172A !important;
    }
    .st-key-preview_page_not_cached button,
    .st-key-preview_page_not_cached button:disabled,
    .st-key-preview_page_not_cached button[disabled],
    .st-key-preview_page_not_cached button[aria-disabled="true"],
    [class*="st-key-"][class*="_thumb_"] button,
    [class*="st-key-"][class*="_thumb_"] button:disabled,
    [class*="st-key-"][class*="_thumb_"] button[disabled],
    [class*="st-key-"][class*="_thumb_"] button[aria-disabled="true"],
    [class*="st-key-"][class*="thumb"] button,
    [class*="st-key-"][class*="thumb"] button:disabled,
    [class*="st-key-"][class*="thumb"] button[disabled],
    [class*="st-key-"][class*="thumb"] button[aria-disabled="true"] {
        display: inline-flex !important;
        align-items: center !important;
        justify-content: center !important;
        min-height: 40px !important;
        padding: 0.55rem 0.85rem !important;
        background: #172033 !important;
        background-color: #172033 !important;
        border-color: #2D3B55 !important;
        color: #FFFFFF !important;
        -webkit-text-fill-color: #FFFFFF !important;
        opacity: 1 !important;
        line-height: 1.2 !important;
        text-shadow: none !important;
    }
    .st-key-preview_page_not_cached button :is(span, p, div, label, small, strong),
    .st-key-preview_page_not_cached button svg,
    .st-key-preview_page_not_cached button:disabled :is(span, p, div, label, small, strong),
    .st-key-preview_page_not_cached button[disabled] :is(span, p, div, label, small, strong),
    .st-key-preview_page_not_cached button[aria-disabled="true"] :is(span, p, div, label, small, strong),
    [class*="st-key-"][class*="_thumb_"] button :is(span, p, div, label, small, strong),
    [class*="st-key-"][class*="_thumb_"] button svg,
    [class*="st-key-"][class*="_thumb_"] button:disabled :is(span, p, div, label, small, strong),
    [class*="st-key-"][class*="_thumb_"] button[disabled] :is(span, p, div, label, small, strong),
    [class*="st-key-"][class*="_thumb_"] button[aria-disabled="true"] :is(span, p, div, label, small, strong),
    [class*="st-key-"][class*="thumb"] button :is(span, p, div, label, small, strong),
    [class*="st-key-"][class*="thumb"] button svg,
    [class*="st-key-"][class*="thumb"] button:disabled :is(span, p, div, label, small, strong),
    [class*="st-key-"][class*="thumb"] button[disabled] :is(span, p, div, label, small, strong),
    [class*="st-key-"][class*="thumb"] button[aria-disabled="true"] :is(span, p, div, label, small, strong) {
        display: inline-flex !important;
        align-items: center !important;
        justify-content: center !important;
        margin: 0 !important;
        color: #FFFFFF !important;
        fill: currentColor !important;
        stroke: currentColor !important;
        -webkit-text-fill-color: #FFFFFF !important;
        opacity: 1 !important;
        line-height: 1.2 !important;
        text-shadow: none !important;
    }
    .st-key-preview_page_not_cached button:not(:disabled):hover,
    .st-key-preview_page_not_cached button:not(:disabled):focus-visible,
    [class*="st-key-"][class*="_thumb_"] button:not(:disabled):hover,
    [class*="st-key-"][class*="_thumb_"] button:not(:disabled):focus-visible,
    [class*="st-key-"][class*="thumb"] button:not(:disabled):hover,
    [class*="st-key-"][class*="thumb"] button:not(:disabled):focus-visible {
        background: #243047 !important;
        background-color: #243047 !important;
        border-color: #41516F !important;
        color: #FFFFFF !important;
        -webkit-text-fill-color: #FFFFFF !important;
    }
    .st-key-preview_page_not_cached button:not(:disabled):active,
    [class*="st-key-"][class*="_thumb_"] button:not(:disabled):active,
    [class*="st-key-"][class*="thumb"] button:not(:disabled):active {
        background: #0F172A !important;
        background-color: #0F172A !important;
        border-color: #93A4C7 !important;
        color: #FFFFFF !important;
        -webkit-text-fill-color: #FFFFFF !important;
    }
    [class*="st-key-preview_page_input_"] [data-testid="stNumberInputContainer"] {
        display: flex !important;
        align-items: stretch !important;
        overflow: hidden !important;
        background: #272832 !important;
        background-color: #272832 !important;
        border: 1px solid #3A3D49 !important;
        border-radius: 12px !important;
        box-shadow: none !important;
        height: 40px !important;
    }
    [class*="st-key-preview_page_input_"] [data-testid="stNumberInputContainer"] > div,
    [class*="st-key-preview_page_input_"] [data-testid="stNumberInput"] [data-baseweb="input"] {
        align-self: stretch !important;
        height: 100% !important;
        background: #272832 !important;
        background-color: #272832 !important;
        border: 0 !important;
        border-radius: 0 !important;
        box-shadow: none !important;
    }
    [class*="st-key-preview_page_input_"] [data-testid="stNumberInput"] [data-baseweb="input"] > div,
    [class*="st-key-preview_page_input_"] [data-testid="stNumberInput"] [data-baseweb="base-input"],
    [class*="st-key-preview_page_input_"] [data-testid="stNumberInput"] input,
    [class*="st-key-preview_page_input_"] [data-testid="stNumberInput"] button,
    [class*="st-key-preview_page_input_"] [data-testid="stNumberInput"] [role="button"] {
        background: #272832 !important;
        background-color: #272832 !important;
        border: 0 !important;
        border-radius: 0 !important;
        box-shadow: none !important;
        color: #FFFFFF !important;
        -webkit-text-fill-color: #FFFFFF !important;
        opacity: 1 !important;
    }
    [class*="st-key-preview_page_input_"] [data-testid="stNumberInput"] input {
        height: 100% !important;
        min-height: 100% !important;
    }
    [class*="st-key-preview_page_input_"] [data-testid="stNumberInput"] button,
    [class*="st-key-preview_page_input_"] [data-testid="stNumberInput"] [role="button"] {
        align-self: stretch !important;
        height: 100% !important;
        min-height: 100% !important;
    }
    [class*="st-key-preview_page_input_"] [data-testid="stNumberInput"] button *,
    [class*="st-key-preview_page_input_"] [data-testid="stNumberInput"] [role="button"] *,
    [class*="st-key-preview_page_input_"] [data-testid="stNumberInput"] svg {
        color: #FFFFFF !important;
        fill: currentColor !important;
        stroke: currentColor !important;
        -webkit-text-fill-color: #FFFFFF !important;
        opacity: 1 !important;
    }
    [class*="st-key-preview_page_input_"] [data-testid="stNumberInput"] button:hover,
    [class*="st-key-preview_page_input_"] [data-testid="stNumberInput"] [role="button"]:hover {
        background: #333540 !important;
        background-color: #333540 !important;
    }
    .st-key-return_to_previous_level_from_results button,
    .st-key-return_from_preview button,
    .st-key-preview_next_page button,
    [class*="st-key-issue_group_open_"][class*="_expand_all"] button,
    [class*="st-key-issue_group_open_"][class*="_collapse_all"] button,
    [class*="st-key-issue_group_open_"][class*="_toggle"] button,
    .st-key-return_to_previous_level_from_results button :is(span, p, div, label, small, strong),
    .st-key-return_from_preview button :is(span, p, div, label, small, strong),
    .st-key-preview_next_page button :is(span, p, div, label, small, strong),
    [class*="st-key-issue_group_open_"][class*="_expand_all"] button :is(span, p, div, label, small, strong),
    [class*="st-key-issue_group_open_"][class*="_collapse_all"] button :is(span, p, div, label, small, strong),
    [class*="st-key-issue_group_open_"][class*="_toggle"] button :is(span, p, div, label, small, strong) {
        color: #FFFFFF !important;
        -webkit-text-fill-color: #FFFFFF !important;
        opacity: 1 !important;
    }
    .st-key-return_to_previous_level_from_results button,
    .st-key-return_from_preview button,
    .st-key-preview_next_page button,
    [class*="st-key-issue_group_open_"][class*="_expand_all"] button,
    [class*="st-key-issue_group_open_"][class*="_collapse_all"] button,
    [class*="st-key-issue_group_open_"][class*="_toggle"] button {
        background: #172033 !important;
        background-color: #172033 !important;
    }
    .st-key-return_to_previous_level_from_results button svg,
    .st-key-return_from_preview button svg,
    .st-key-preview_next_page button svg,
    [class*="st-key-issue_group_open_"][class*="_expand_all"] button svg,
    [class*="st-key-issue_group_open_"][class*="_collapse_all"] button svg,
    [class*="st-key-issue_group_open_"][class*="_toggle"] button svg {
        color: #FFFFFF !important;
        fill: currentColor !important;
        stroke: currentColor !important;
        opacity: 1 !important;
    }
    [class*="st-key-preview_page_input_"] [data-testid="stWidgetLabel"],
    [class*="st-key-preview_page_input_"] [data-testid="stWidgetLabel"] :is(span, p, div, label, small, strong),
    [class*="st-key-preview_zoom_control_"] [data-testid="stWidgetLabel"],
    [class*="st-key-preview_zoom_control_"] [data-testid="stWidgetLabel"] :is(span, p, div, label, small, strong),
    .st-key-preview_show_markers_control label,
    .st-key-preview_show_markers_control label :is(span, p, div, small, strong),
    .st-key-comparison_sync_scroll_control label,
    .st-key-comparison_sync_scroll_control label :is(span, p, div, small, strong),
    .st-key-comparison_sync_zoom_control label,
    .st-key-comparison_sync_zoom_control label :is(span, p, div, small, strong),
    .st-key-comparison_preview_mode [data-testid="stWidgetLabel"],
    .st-key-comparison_preview_mode [data-testid="stWidgetLabel"] :is(span, p, div, label, small, strong),
    .st-key-comparison_preview_mode [data-testid="stRadio"] label,
    .st-key-comparison_preview_mode [data-testid="stRadio"] label :is(span, p, div, small, strong) {
        color: #273044 !important;
        -webkit-text-fill-color: #273044 !important;
        opacity: 1 !important;
    }
    .st-key-comparison_preview_mode [data-testid="stRadio"] label {
        display: inline-flex !important;
        align-items: center !important;
        min-height: 28px !important;
        line-height: 1.2 !important;
    }
    .st-key-comparison_preview_mode [data-testid="stRadio"] label :is(span, p, div, small, strong) {
        line-height: 1.2 !important;
        margin-top: 0 !important;
        margin-bottom: 0 !important;
    }
    </style>
    """,
    unsafe_allow_html=True,
)

with st.sidebar:
    ui_language = _ui_language_selectbox()

    st.subheader(t(ui_language, "settings"))
    report_language_options = ["auto", "en", "zh"]
    report_language_aliases = {
        "Auto": "auto",
        "自动": "auto",
        "English": "en",
        "英文": "en",
        "Chinese": "zh",
        "中文": "zh",
    }
    report_language_override = _localized_selectbox(
        t(ui_language, "report_language"),
        report_language_options,
        "report_language_override",
        "auto",
        ui_language,
        format_func=lambda value: {
            "auto": t(ui_language, "auto"),
            "en": t(ui_language, "english"),
            "zh": t(ui_language, "chinese"),
        }[value],
        aliases=report_language_aliases,
    )
    st.caption(t(ui_language, "default_report_language_note"))

    profile_keys = ["template_matching", "ieee", "custom"]
    profile_aliases = {
        "Template-matching": "template_matching",
        "模板匹配": "template_matching",
        "IEEE": "ieee",
        "Custom": "custom",
        "自定义": "custom",
    }
    format_profile_name = _localized_selectbox(
        t(ui_language, "format_profile"),
        profile_keys,
        "format_profile_name",
        "template_matching",
        ui_language,
        format_func=lambda value: label(PROFILE_LABELS, ui_language, value),
        aliases=profile_aliases,
    )
    st.caption(t(ui_language, "default_profile_note"))

    mode_value_by_key = {
        "analyze_only": ProcessingMode.ANALYZE_ONLY.value,
        "fix_numbering_and_formatting": ProcessingMode.FIX_NUMBERING_AND_FORMATTING.value,
        "fix_numbering_only": ProcessingMode.FIX_NUMBERING_ONLY.value,
        "check_references_only": ProcessingMode.CHECK_REFERENCES_ONLY.value,
    }
    mode_aliases = {
        value: key for key, value in mode_value_by_key.items()
    } | {
        label(MODE_LABELS, "zh", value): key for key, value in mode_value_by_key.items()
    }
    mode_keys = list(mode_value_by_key.keys())
    mode_key = _localized_selectbox(
        t(ui_language, "processing_mode"),
        mode_keys,
        "processing_mode",
        "analyze_only",
        ui_language,
        format_func=lambda value: label(MODE_LABELS, ui_language, mode_value_by_key[value]),
        aliases=mode_aliases,
    )
    mode = mode_value_by_key[mode_key]
    st.caption(t(ui_language, "default_mode_note"))

    custom_options = {}
    if format_profile_name == "custom":
        with st.expander(t(ui_language, "custom_settings"), expanded=True):
            number_keys = ["arabic", "roman_upper", "roman_lower"]
            table_layout_keys = ["single_line", "number_line_plus_caption"]
            equation_keys = ["parentheses", "brackets", "eq_parentheses", "equation_parentheses"]
            custom_options["figure_prefix"] = st.text_input(t(ui_language, "figure_prefix"), key="figure_prefix")
            custom_options["figure_number_style"] = _state_selectbox(
                t(ui_language, "figure_number_style"),
                number_keys,
                "figure_number_style",
                "arabic",
                ui_language,
                format_func=lambda value: label(NUMBER_STYLE_LABELS, ui_language, value),
            )
            custom_options["figure_separator"] = st.text_input(t(ui_language, "figure_separator"), key="figure_separator")
            custom_options["table_prefix"] = st.text_input(t(ui_language, "table_prefix"), key="table_prefix")
            custom_options["table_number_style"] = _state_selectbox(
                t(ui_language, "table_number_style"),
                number_keys,
                "table_number_style",
                "roman_upper",
                ui_language,
                format_func=lambda value: label(NUMBER_STYLE_LABELS, ui_language, value),
            )
            custom_options["table_layout"] = _state_selectbox(
                t(ui_language, "table_layout"),
                table_layout_keys,
                "table_layout",
                "single_line",
                ui_language,
                format_func=lambda value: label(TABLE_LAYOUT_LABELS, ui_language, value),
            )
            custom_options["equation_wrapper"] = _state_selectbox(
                t(ui_language, "equation_format"),
                equation_keys,
                "equation_wrapper",
                "parentheses",
                ui_language,
                format_func=lambda value: label(EQUATION_WRAPPER_LABELS, ui_language, value),
            )

    style_source_value_by_key = {
        "auto_detect": "Auto-detect from uploaded report",
        "uploaded_template": "Use uploaded template as style source",
        "fallback_defaults": "Use fallback defaults",
    }
    style_source_aliases = {
        value: key for key, value in style_source_value_by_key.items()
    } | {
        label(STYLE_SOURCE_LABELS, "zh", value): key for key, value in style_source_value_by_key.items()
    }
    style_source_keys = list(style_source_value_by_key.keys())
    style_source_key = _localized_selectbox(
        t(ui_language, "style_source"),
        style_source_keys,
        "style_source",
        "auto_detect",
        ui_language,
        format_func=lambda value: label(STYLE_SOURCE_LABELS, ui_language, style_source_value_by_key[value]),
        aliases=style_source_aliases,
    )
    style_source = style_source_value_by_key[style_source_key]
    st.caption(t(ui_language, "default_style_source_note"))

    st.subheader(t(ui_language, "options"))
    preserve_cover_page = st.checkbox(t(ui_language, "preserve_cover"), key="preserve_cover_page")
    preserve_two_column_layout = st.checkbox(t(ui_language, "preserve_two_column"), key="preserve_two_column_layout")
    insert_placeholder_captions = st.checkbox(t(ui_language, "insert_placeholders"), key="insert_placeholder_captions")
    renumber_figures = st.checkbox(t(ui_language, "renumber_figures"), key="renumber_figures")
    renumber_tables = st.checkbox(t(ui_language, "renumber_tables"), key="renumber_tables")
    renumber_equations = st.checkbox(t(ui_language, "renumber_equations"), key="renumber_equations")
    check_references = st.checkbox(t(ui_language, "check_references"), key="check_references")
    generate_report = st.checkbox(t(ui_language, "generate_report"), key="generate_report")
    with st.expander(t(ui_language, "default_notes"), expanded=False):
        st.markdown(
            "\n".join(
                [
                    f"- {t(ui_language, 'default_report_language_note')}",
                    f"- {t(ui_language, 'default_profile_note')}",
                    f"- {t(ui_language, 'default_mode_note')}",
                    f"- {t(ui_language, 'default_style_source_note')}",
                    f"- {t(ui_language, 'default_options_note')}",
                ]
            )
        )
    run = st.button(t(ui_language, "run"), type="primary", use_container_width=True, key="run_formatter")
    _render_diagnostics(ui_language)

if _render_active_preview_from_session(ui_language):
    st.stop()

st.title(t(ui_language, "app_title"))
st.caption(t(ui_language, "app_caption"))

upload_cols = st.columns(2)
with upload_cols[0]:
    report_file = st.file_uploader(
        t(ui_language, "upload_report"),
        type=["docx", "pdf"],
        key="report_file",
    )
with upload_cols[1]:
    template_file = st.file_uploader(
        t(ui_language, "upload_template"),
        type=["docx"],
        key="template_file",
    )

report_bytes = report_file.getvalue() if report_file else None
template_bytes = template_file.getvalue() if template_file else None
is_pdf_input = bool(report_file and Path(report_file.name).suffix.lower() == ".pdf")
effective_style_source = style_source
if style_source == "Use uploaded template as style source" and template_bytes is None:
    effective_style_source = "Auto-detect from uploaded report"
if is_pdf_input:
    mode = ProcessingMode.ANALYZE_ONLY.value
    effective_style_source = "Auto-detect from uploaded report"

options = FormattingOptions(
    preserve_cover_page=preserve_cover_page,
    preserve_two_column_layout=preserve_two_column_layout,
    insert_placeholder_captions=insert_placeholder_captions,
    renumber_figures=renumber_figures,
    renumber_tables=renumber_tables,
    renumber_equations=renumber_equations,
    check_in_text_references=check_references,
    generate_check_report=generate_report,
    ui_language=ui_language,
    report_language_override=report_language_override,
    format_profile_name=format_profile_name,
    custom_profile_options=custom_options,
)

current_input_signature = None
if report_bytes is None:
    _clear_latest_result()
else:
    current_input_signature = build_processing_signature(
        report_bytes=report_bytes,
        report_file_name=report_file.name,
        template_bytes=template_bytes,
        template_file_name=template_file.name if template_file else None,
        mode=mode,
        style_source=effective_style_source,
        options=options,
    )
    latest_signature = st.session_state.get("latest_input_signature")
    if latest_signature is not None and latest_signature != current_input_signature:
        _clear_latest_result()

report_doc_for_detection = None
language_detection = None
if report_file and not is_pdf_input:
    report_doc_for_detection = Document(io.BytesIO(report_bytes))
    language_detection = detect_report_language(report_doc_for_detection)
    final_report_language = resolve_report_language(language_detection, report_language_override)
    _render_summary_cards(
        [
            _stats_card(ui_language, "detected_language", _language_name(ui_language, language_detection["language"])),
            _stats_card(ui_language, "final_language", _language_name(ui_language, final_report_language)),
            _stats_card(ui_language, "confidence", f"{language_detection['confidence']:.2f}"),
            _stats_card(ui_language, "english_words", language_detection["english_word_count"]),
            _stats_card(ui_language, "chinese_chars", language_detection["chinese_char_count"]),
        ]
    )
    if language_detection["language"] == "unknown" and report_language_override == "auto":
        st.warning(t(ui_language, "language_unknown_warning"))

if not report_file:
    st.markdown(
        f"""
        <div class="status-strip">
        {t(ui_language, "waiting")}
        </div>
        """,
        unsafe_allow_html=True,
    )
else:
    if is_pdf_input:
        st.info(t(ui_language, "pdf_analysis_only"))
    reuse_latest_result = bool(
        run
        and not st.session_state.get("results_view_active", True)
        and st.session_state.get("latest_result") is not None
        and st.session_state.get("latest_input_signature") == current_input_signature
    )
    if run and not reuse_latest_result:
        if not is_pdf_input and style_source == "Use uploaded template as style source" and template_bytes is None:
            st.warning(t(ui_language, "no_template_warning"))

        with st.spinner(t(ui_language, "processing")):
            if is_pdf_input:
                pipeline_run = process_uploaded_pdf(
                    report_bytes=report_bytes,
                    input_file_name=report_file.name,
                    options=options,
                )
            else:
                pipeline_run = process_uploaded_docx(
                    report_bytes=report_bytes,
                    input_file_name=report_file.name,
                    template_bytes=template_bytes,
                    template_file_name=template_file.name if template_file else None,
                    mode=mode,
                    style_source=effective_style_source,
                    options=options,
                )

        st.session_state["latest_result"] = pipeline_run.result
        st.session_state["latest_original_docx_bytes"] = report_bytes if not is_pdf_input else None
        st.session_state["latest_original_file_name"] = report_file.name
        st.session_state["latest_docx_bytes"] = pipeline_run.docx_bytes if not is_pdf_input else None
        st.session_state["latest_report_bytes"] = pipeline_run.report_bytes
        st.session_state["latest_run_id"] = pipeline_run.run_id
        st.session_state["latest_docx_file_name"] = pipeline_run.docx_file_name if not is_pdf_input else None
        st.session_state["latest_report_file_name"] = pipeline_run.report_file_name
        st.session_state["latest_debug_info"] = pipeline_run.debug_info()
        st.session_state["latest_input_signature"] = current_input_signature
        st.session_state["latest_document_version"] = (
            None
            if is_pdf_input
            else build_document_version(
                input_signature=current_input_signature,
                original_docx_bytes=report_bytes,
                modified_docx_bytes=pipeline_run.docx_bytes if _is_fix_mode(mode) else None,
            )
        )
    if run:
        st.session_state["results_view_active"] = True

    result = st.session_state.get("latest_result")
    if result is None or not st.session_state.get("results_view_active", True):
        _render_stats_help(ui_language)
    else:
        _render_summary_cards(
            [
                _stats_card(ui_language, "figures", len(result.figures)),
                _stats_card(ui_language, "tables", len(result.tables)),
                _stats_card(ui_language, "native_tables", len(getattr(result, "native_tables", []))),
                _stats_card(ui_language, "image_tables", len(getattr(result, "image_tables", []))),
                _stats_card(ui_language, "equation_blocks", len(getattr(result, "equation_blocks", result.equations))),
                _stats_card(ui_language, "equation_numbers", len(result.equations)),
                _stats_card(ui_language, "unnumbered_equations", len(getattr(result, "unnumbered_equations", []))),
                _stats_card(ui_language, "reference_citations", len(getattr(result, "reference_citations", []))),
                _stats_card(ui_language, "skipped_toc", getattr(result, "skipped_toc_entries_count", 0)),
                _stats_card(ui_language, "issues", len(result.issues)),
            ]
        )
        _render_stats_help(ui_language)
        if _is_fix_mode(mode) and _has_equation_auto_edit_suspended_notice(result):
            st.warning(t(ui_language, "equation_auto_edit_suspended_notice"))

        preview_environment = detect_preview_environment()
        preview_mode = _preview_mode_for_processing_mode(mode)
        latest_docx_bytes = st.session_state.get("latest_docx_bytes")
        modified_docx_bytes = latest_docx_bytes if preview_mode == PreviewMode.COMPARISON else None
        document_version = st.session_state.get("latest_document_version") or build_document_version(
            input_signature=current_input_signature,
            original_docx_bytes=report_bytes,
            modified_docx_bytes=modified_docx_bytes,
        )
        preview_gate = evaluate_preview_gate(
            mode=preview_mode,
            report_bytes=report_bytes,
            result_present=result is not None,
            modified_docx_bytes=modified_docx_bytes,
            current_signature=current_input_signature,
            latest_signature=st.session_state.get("latest_input_signature"),
            environment=preview_environment,
        )

        active_preview = st.session_state.get("active_preview_mode")
        if active_preview:
            active_mode = PreviewMode(active_preview)
            _render_preview_workspace(
                ui_language=ui_language,
                mode=active_mode,
                report_bytes=report_bytes,
                report_file_name=report_file.name,
                modified_docx_bytes=latest_docx_bytes if active_mode == PreviewMode.COMPARISON else None,
                modified_file_name=st.session_state.get("latest_docx_file_name"),
                result=result,
                task_id=st.session_state.get("latest_run_id", ""),
                document_version=document_version,
                environment=preview_environment,
            )
            st.stop()

        _render_preview_entry(preview_gate, preview_environment, ui_language)

        result_nav_cols = st.columns([0.82, 4.2])
        with result_nav_cols[0]:
            if st.button(
                t(ui_language, "return_previous_level"),
                use_container_width=True,
                key="return_to_previous_level_from_results",
            ):
                st.session_state["active_preview_mode"] = None
                st.session_state["results_view_active"] = False
                st.rerun()

        tabs = st.tabs(
            [
                t(ui_language, "tab_profile"),
                t(ui_language, "tab_objects"),
                t(ui_language, "tab_issues"),
                t(ui_language, "tab_downloads"),
            ]
        )

        with tabs[0]:
            _render_readable_table(_style_rows(result.style_profile, ui_language))
            st.info(t(ui_language, "template_missing_note"))
            with st.expander(t(ui_language, "developer_debug_info"), expanded=False):
                st.caption(t(ui_language, "raw_regex_patterns"))
                _render_readable_table(_style_debug_rows(result.style_profile, ui_language))
                st.json(result.style_profile.to_dict())

        with tabs[1]:
            cols = st.columns(3)
            cols[0].write(t(ui_language, "figures"))
            with cols[0]:
                _render_readable_table(_caption_rows(result.figures, ui_language))
            cols[1].write(t(ui_language, "tables"))
            with cols[1]:
                _render_readable_table(_caption_rows(result.tables, ui_language))
            cols[2].write(t(ui_language, "equations"))
            with cols[2]:
                _render_readable_table(_caption_rows(result.equations, ui_language))
            st.write(t(ui_language, "image_scan"))
            _render_readable_table(_image_rows(getattr(result, "images", []), ui_language))

        with tabs[2]:
            _render_issue_groups(result, ui_language, st.session_state.get("latest_run_id", ""))

        with tabs[3]:
            run_id = st.session_state.get("latest_run_id", "")
            docx_file_name = st.session_state.get("latest_docx_file_name", f"report_formatted_{run_id}.docx")
            report_file_name = st.session_state.get("latest_report_file_name", f"report_check_report_{run_id}.txt")
            st.caption(f"{t(ui_language, 'current_download_run_id')}: {run_id}")
            if _is_fix_mode(mode):
                expected_preview_version = build_preview_version(document_version, PreviewMode.COMPARISON)
                preview_ready_for_download = st.session_state.get("latest_preview_version") == expected_preview_version
                st.caption(t(ui_language, "download_after_preview_note"))
                st.download_button(
                    f"{t(ui_language, 'download_docx')} - {docx_file_name}",
                    data=st.session_state["latest_docx_bytes"],
                    file_name=docx_file_name,
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    type="primary",
                    use_container_width=True,
                    disabled=not preview_ready_for_download,
                    key=f"download_formatted_docx_{run_id}",
                )
                if not preview_ready_for_download:
                    st.warning(t(ui_language, "download_waiting_for_preview"))
            else:
                st.info(t(ui_language, "analyze_mode_no_docx_download"))
            if result.report_text:
                st.download_button(
                    f"{t(ui_language, 'download_report')} - {report_file_name}",
                    data=st.session_state["latest_report_bytes"],
                    file_name=report_file_name,
                    mime="text/plain",
                    type="primary",
                    use_container_width=True,
                    key=f"download_format_report_{run_id}",
                )
                st.text_area(t(ui_language, "report_preview"), result.report_text, height=360)
            _render_download_debug_info(st.session_state.get("latest_debug_info"), ui_language)
