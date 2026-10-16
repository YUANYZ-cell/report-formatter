from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "samples"


def add_body(paragraph, text: str):
    paragraph.text = text
    paragraph.style = "Normal"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)


def add_caption(paragraph, text: str):
    paragraph.text = text
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(10)


def clean_template_report(path: Path):
    doc = Document()
    doc.add_paragraph("IEEE LAB REPORT TEMPLATE")
    doc.add_page_break()
    doc.add_paragraph("Abstract")
    add_body(doc.add_paragraph(), "This report references Fig. 1, Table I, and Eq. (1).")
    add_caption(doc.add_paragraph(), "Fig. 1. Clean sensor waveform")
    add_caption(doc.add_paragraph(), "TABLE I")
    add_caption(doc.add_paragraph(), "SENSOR SUMMARY")
    doc.add_paragraph("y = mx + b     (1)")
    doc.save(path)


def bad_numbering_report(path: Path):
    doc = Document()
    doc.add_paragraph("IEEE LAB REPORT WITH BAD NUMBERING")
    doc.add_page_break()
    doc.add_paragraph("Abstract")
    add_body(doc.add_paragraph(), "This report references Fig. 1, Table I, and Eq. (1).")
    add_caption(doc.add_paragraph(), "Fig. 1. First result")
    add_caption(doc.add_paragraph(), "Fig. 3. Third result with missing numbering")
    add_caption(doc.add_paragraph(), "Table 1")
    add_caption(doc.add_paragraph(), "MEASURED VALUES")
    doc.add_paragraph("v = ir     (3)")
    doc.save(path)


def reference_errors_report(path: Path):
    doc = Document()
    doc.add_paragraph("IEEE LAB REPORT WITH REFERENCE ERRORS")
    doc.add_page_break()
    doc.add_paragraph("Abstract")
    add_body(
        doc.add_paragraph(),
        "This report cites Figure 1, Fig. 5, Table 2, and Eq. (3), but only some objects exist.",
    )
    add_caption(doc.add_paragraph(), "Fig. 1. Existing figure")
    add_caption(doc.add_paragraph(), "TABLE I")
    add_caption(doc.add_paragraph(), "EXISTING TABLE")
    doc.add_paragraph("p = vi     (1)")
    doc.save(path)


def main():
    SAMPLES.mkdir(exist_ok=True)
    clean_template_report(SAMPLES / "sample_clean_template_report.docx")
    bad_numbering_report(SAMPLES / "sample_bad_numbering_report.docx")
    reference_errors_report(SAMPLES / "sample_reference_errors_report.docx")


if __name__ == "__main__":
    main()
