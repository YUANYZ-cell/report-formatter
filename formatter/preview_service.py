from __future__ import annotations

import hashlib
import html
import io
import json
import mimetypes
import os
import re
import shutil
import subprocess
import tempfile
import textwrap
import time
import urllib.error
import urllib.parse
import urllib.request
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from enum import Enum
from functools import lru_cache
from pathlib import Path
from typing import Callable, Iterable

try:
    from PIL import Image, ImageChops, ImageDraw, ImageFilter, ImageFont
except Exception:  # pragma: no cover - Pillow is optional at import time.
    Image = None
    ImageChops = None
    ImageDraw = None
    ImageFilter = None
    ImageFont = None


PREVIEW_CACHE_VERSION = "preview-cache-2026-06-25-safe-pixel-diff-v2"
DEFAULT_PREVIEW_PRELOAD_RADIUS = 1
PIXEL_DIFF_MAX_MARKERS_PER_PAGE = 8
PIXEL_DIFF_MAX_TOTAL_AREA_RATIO = 0.035


class PreviewMode(str, Enum):
    ISSUE = "issue"
    COMPARISON = "comparison"


class PreviewStatus(str, Enum):
    IDLE = "idle"
    PREPARING = "preparing"
    CONVERTING = "converting"
    READY = "ready"
    FAILED = "failed"
    STALE = "stale"
    CANCELLED = "cancelled"


class PreviewUnavailable(RuntimeError):
    pass


@dataclass(frozen=True)
class PreviewEnvironment:
    converter_available: bool
    renderer_available: bool
    converter_label: str
    renderer_label: str
    gotenberg_url: str | None = None
    libreoffice_path: str | None = None
    pdftoppm_path: str | None = None
    pdftotext_path: str | None = None
    word_available: bool = False
    word_status: str = "not_checked"
    word_path: str | None = None
    word_error: str | None = None
    low_fidelity_available: bool = True
    low_fidelity_label: str = "DOCX browser preview"

    @property
    def high_fidelity_available(self) -> bool:
        return self.converter_available and self.renderer_available

    @property
    def available(self) -> bool:
        return self.high_fidelity_available or self.low_fidelity_available

    @property
    def uses_low_fidelity(self) -> bool:
        return not self.high_fidelity_available and self.low_fidelity_available


@dataclass(frozen=True)
class PreviewGate:
    mode: PreviewMode
    enabled: bool
    reason_key: str
    button_key: str
    help_key: str


@dataclass(frozen=True)
class RenderedPage:
    page_number: int
    image_bytes: bytes
    width: int | None = None
    height: int | None = None


@dataclass(frozen=True)
class RenderedDocument:
    label: str
    docx_sha256: str
    pdf_sha256: str
    page_count: int
    rendered_pages: tuple[RenderedPage, ...]
    converter_label: str
    renderer_label: str
    pdf_bytes: bytes = field(default=b"", repr=False, compare=False)


@dataclass(frozen=True)
class PreviewItem:
    item_id: str
    kind: str
    severity: str
    page_number: int
    title: str
    message: str
    before: str = ""
    after: str = ""
    rule_id: str = ""
    confidence: float = 0.35
    approximate: bool = True
    rect: tuple[float, float, float, float] | None = None
    object_id: str = ""
    source_anchor: str = ""
    coordinate_space: str = "pdf_normalized_top_left"
    location_status: str = "unlocated"
    is_cross_column: bool = False


@dataclass(frozen=True)
class PdfTextLine:
    page_number: int
    page_width: float
    page_height: float
    text: str
    bbox: tuple[float, float, float, float]


@dataclass(frozen=True)
class PreviewBundle:
    mode: PreviewMode
    task_id: str
    document_version: str
    preview_version: str
    original: RenderedDocument
    modified: RenderedDocument | None
    issues: tuple[PreviewItem, ...]
    changes: tuple[PreviewItem, ...]
    pixel_changes: tuple[PreviewItem, ...] = field(default_factory=tuple)
    timings: dict[str, float | int | str] = field(default_factory=dict)


@dataclass(frozen=True)
class _WordCapability:
    available: bool
    status: str
    path: str | None = None
    error: str | None = None


def sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def short_hash(*parts: object, length: int = 16) -> str:
    payload = json.dumps(parts, sort_keys=True, ensure_ascii=True, default=str).encode("utf-8")
    return sha256_bytes(payload)[:length]


def build_document_version(
    *,
    input_signature: str | None,
    original_docx_bytes: bytes | None,
    modified_docx_bytes: bytes | None,
) -> str:
    return short_hash(
        "document",
        input_signature or "",
        sha256_bytes(original_docx_bytes or b""),
        sha256_bytes(modified_docx_bytes or b""),
    )


def build_preview_version(document_version: str, mode: PreviewMode | str, dpi: int = 120) -> str:
    return short_hash("preview", document_version, PreviewMode(mode).value, dpi)


def _cache_root() -> Path:
    configured = os.environ.get("REPORT_FORMATTER_PREVIEW_CACHE_DIR")
    root = Path(configured) if configured else Path(tempfile.gettempdir()) / "rule_based_ieee_preview_cache"
    root.mkdir(parents=True, exist_ok=True)
    return root


def _cache_key(*parts: object) -> str:
    payload = json.dumps(parts, sort_keys=True, ensure_ascii=True, default=str).encode("utf-8")
    return sha256_bytes(payload)


def _cache_path(kind: str, key: str, suffix: str) -> Path:
    directory = _cache_root() / kind / key[:2]
    directory.mkdir(parents=True, exist_ok=True)
    return directory / f"{key}{suffix}"


def _converter_cache_key(docx_bytes: bytes, file_name: str, environment: PreviewEnvironment) -> str:
    return _cache_key(
        PREVIEW_CACHE_VERSION,
        "docx-pdf",
        sha256_bytes(docx_bytes),
        Path(file_name or "document.docx").suffix.lower(),
        environment.converter_label,
        environment.word_status,
        environment.word_path,
        environment.libreoffice_path,
        environment.gotenberg_url,
    )


def _page_cache_key(pdf_bytes: bytes, page_number: int, dpi: int) -> str:
    return _cache_key(PREVIEW_CACHE_VERSION, "pdf-page", sha256_bytes(pdf_bytes), page_number, dpi)


def _record_elapsed(timings: dict[str, float | int | str] | None, key: str, started_at: float) -> None:
    if timings is not None:
        timings[key] = round(time.perf_counter() - started_at, 3)


def _increment_timing(timings: dict[str, float | int | str] | None, key: str, amount: int = 1) -> None:
    if timings is not None:
        timings[key] = int(timings.get(key, 0) or 0) + amount


def _set_timing(timings: dict[str, float | int | str] | None, key: str, value: float | int | str) -> None:
    if timings is not None:
        timings[key] = value


def _valid_cached_pdf(path: Path) -> bool:
    if not path.exists() or path.stat().st_size < 512:
        return False
    try:
        with path.open("rb") as handle:
            return handle.read(4) == b"%PDF"
    except OSError:
        return False


def _selected_page_numbers(page_count: int, page_number: int | None, preload_radius: int) -> list[int]:
    page_count = max(1, page_count)
    if page_number is None:
        return list(range(1, page_count + 1))
    current_page = min(max(1, int(page_number)), page_count)
    radius = max(0, int(preload_radius))
    start_page = max(1, current_page - radius)
    end_page = min(page_count, current_page + radius)
    return list(range(start_page, end_page + 1))


def _find_poppler_executable(command: str) -> str | None:
    """Prefer a native Poppler executable over a potentially broken .cmd shim on Windows."""
    return shutil.which(f"{command}.exe") or shutil.which(command)


@lru_cache(maxsize=1)
def detect_preview_environment() -> PreviewEnvironment:
    word = _detect_word_capability()
    libreoffice_path = _find_libreoffice()
    gotenberg_url = _clean_gotenberg_url(os.environ.get("GOTENBERG_URL") or os.environ.get("GOTENBERG_ENDPOINT"))
    pdftoppm_path = _find_poppler_executable("pdftoppm")
    pdftotext_path = _find_poppler_executable("pdftotext") or _sibling_executable(pdftoppm_path, "pdftotext")
    high_converter_available = bool(word.available or libreoffice_path or gotenberg_url)
    renderer_available = bool(pdftoppm_path)
    low_fidelity_available = _low_fidelity_available()

    if high_converter_available and renderer_available:
        converter_label = _high_fidelity_converter_label(word, libreoffice_path, gotenberg_url)
        renderer_label = "pdftoppm"
    elif low_fidelity_available:
        converter_label = "DOCX browser preview"
        renderer_label = "Pillow text renderer"
    else:
        converter_label = _high_fidelity_converter_label(word, libreoffice_path, gotenberg_url)
        renderer_label = "pdftoppm" if renderer_available else "Unavailable"

    return PreviewEnvironment(
        converter_available=high_converter_available,
        renderer_available=renderer_available,
        converter_label=converter_label,
        renderer_label=renderer_label,
        gotenberg_url=gotenberg_url,
        libreoffice_path=libreoffice_path,
        pdftoppm_path=pdftoppm_path,
        pdftotext_path=pdftotext_path,
        word_available=word.available,
        word_status=word.status,
        word_path=word.path,
        word_error=word.error,
        low_fidelity_available=low_fidelity_available,
    )


def evaluate_preview_gate(
    *,
    mode: PreviewMode | str,
    report_bytes: bytes | None,
    result_present: bool,
    modified_docx_bytes: bytes | None,
    current_signature: str | None,
    latest_signature: str | None,
    environment: PreviewEnvironment,
) -> PreviewGate:
    preview_mode = PreviewMode(mode)
    if preview_mode == PreviewMode.ISSUE:
        button_key = "open_issue_preview"
        help_key = "open_issue_preview_help"
    else:
        button_key = "open_comparison_preview"
        help_key = "open_comparison_preview_help"

    if not report_bytes:
        return PreviewGate(preview_mode, False, "preview_missing_original", button_key, help_key)
    if not result_present:
        return PreviewGate(preview_mode, False, "preview_run_required", button_key, help_key)
    if latest_signature is None or current_signature != latest_signature:
        return PreviewGate(preview_mode, False, "preview_stale_settings", button_key, help_key)
    if preview_mode == PreviewMode.COMPARISON and not modified_docx_bytes:
        return PreviewGate(preview_mode, False, "preview_missing_modified", button_key, help_key)
    if not environment.available:
        return PreviewGate(preview_mode, False, "preview_route_unavailable", button_key, help_key)
    reason_key = "preview_low_fidelity_ready" if environment.uses_low_fidelity else "preview_ready"
    return PreviewGate(preview_mode, True, reason_key, button_key, help_key)


def render_preview_bundle(
    *,
    mode: PreviewMode | str,
    task_id: str,
    document_version: str,
    original_docx_bytes: bytes,
    original_file_name: str,
    result,
    modified_docx_bytes: bytes | None = None,
    modified_file_name: str | None = None,
    dpi: int = 120,
    max_pages: int = 40,
    page_number: int | None = None,
    preload_radius: int = DEFAULT_PREVIEW_PRELOAD_RADIUS,
    environment: PreviewEnvironment | None = None,
) -> PreviewBundle:
    started_at = time.perf_counter()
    timings: dict[str, float | int | str] = {
        "cache_version": PREVIEW_CACHE_VERSION,
        "requested_page": int(page_number or 0),
        "preload_radius": int(preload_radius),
        "dpi": int(dpi),
    }
    preview_mode = PreviewMode(mode)
    env = environment or detect_preview_environment()
    if not env.available:
        raise PreviewUnavailable("No usable preview route is available.")
    _set_timing(timings, "converter", env.converter_label)
    _set_timing(timings, "renderer", env.renderer_label)

    if preview_mode == PreviewMode.COMPARISON and modified_docx_bytes:
        warm_docx_pdf_cache_batch(
            [
                ("original", original_docx_bytes, original_file_name),
                ("modified", modified_docx_bytes, modified_file_name or "modified.docx"),
            ],
            env,
            timings=timings,
        )

    original = render_source_to_pages(
        original_docx_bytes,
        original_file_name,
        label="original",
        dpi=dpi,
        max_pages=max_pages,
        page_number=page_number,
        preload_radius=preload_radius,
        environment=env,
        timings=timings,
    )
    modified = None
    if preview_mode == PreviewMode.COMPARISON:
        if not modified_docx_bytes:
            raise PreviewUnavailable("Comparison preview requires modified DOCX bytes.")
        modified = render_source_to_pages(
            modified_docx_bytes,
            modified_file_name or "modified.docx",
            label="modified",
            dpi=dpi,
            max_pages=max_pages,
            page_number=page_number,
            preload_radius=preload_radius,
            environment=env,
            timings=timings,
        )
    page_count = max(original.page_count, modified.page_count if modified else 1)
    issues = tuple(build_issue_items(getattr(result, "issues", []), page_count))
    geometry_pdf = modified.pdf_bytes if modified else original.pdf_bytes
    changes = tuple(
        build_change_items(
            getattr(result, "actions", []),
            page_count,
            pdf_bytes=geometry_pdf,
            pdftotext_path=env.pdftotext_path,
        )
    )
    pixel_diagnostics: dict[int, dict[str, float | int | str]] = {}
    pixel_changes = (
        tuple(build_pixel_diff_items(original.rendered_pages, modified.rendered_pages, diagnostics=pixel_diagnostics))
        if modified
        else ()
    )
    if modified and original.page_count != modified.page_count:
        _set_timing(timings, "layout_guard.status", "page_count_mismatch")
        _set_timing(timings, "layout_guard.original_page_count", original.page_count)
        _set_timing(timings, "layout_guard.modified_page_count", modified.page_count)
    else:
        _set_timing(timings, "layout_guard.status", "ok")
    _set_timing(timings, "pixel_diff.marker_count", len(pixel_changes))
    for page_number, page_diagnostic in pixel_diagnostics.items():
        for key, value in page_diagnostic.items():
            _set_timing(timings, f"pixel_diff.page_{page_number}.{key}", value)
    _record_elapsed(timings, "bundle.total_seconds", started_at)
    return PreviewBundle(
        mode=preview_mode,
        task_id=task_id,
        document_version=document_version,
        preview_version=build_preview_version(document_version, preview_mode, dpi),
        original=original,
        modified=modified,
        issues=issues,
        changes=changes,
        pixel_changes=pixel_changes,
        timings=timings,
    )


def render_source_to_pages(
    source_bytes: bytes,
    file_name: str,
    *,
    label: str,
    dpi: int,
    max_pages: int,
    page_number: int | None = None,
    preload_radius: int = DEFAULT_PREVIEW_PRELOAD_RADIUS,
    environment: PreviewEnvironment,
    timings: dict[str, float | int | str] | None = None,
) -> RenderedDocument:
    started_at = time.perf_counter()
    suffix = Path(file_name or "").suffix.lower()
    if suffix == ".pdf":
        document = render_pdf_source_to_pages(
            source_bytes,
            label=label,
            dpi=dpi,
            max_pages=max_pages,
            page_number=page_number,
            preload_radius=preload_radius,
            environment=environment,
            timings=timings,
        )
    else:
        document = render_docx_to_pages(
            source_bytes,
            file_name,
            label=label,
            dpi=dpi,
            max_pages=max_pages,
            page_number=page_number,
            preload_radius=preload_radius,
            environment=environment,
            timings=timings,
        )
    _record_elapsed(timings, f"{label}.source_total_seconds", started_at)
    return document


def render_pdf_source_to_pages(
    pdf_bytes: bytes,
    *,
    label: str,
    dpi: int,
    max_pages: int,
    page_number: int | None = None,
    preload_radius: int = DEFAULT_PREVIEW_PRELOAD_RADIUS,
    environment: PreviewEnvironment,
    timings: dict[str, float | int | str] | None = None,
) -> RenderedDocument:
    if not environment.renderer_available:
        raise PreviewUnavailable("pdftoppm is required to render PDF pages.")
    with tempfile.TemporaryDirectory(prefix="report_preview_") as tmp_name:
        tmp_dir = Path(tmp_name)
        pages, page_count = render_pdf_bytes_to_png_pages(
            pdf_bytes,
            tmp_dir,
            dpi=dpi,
            max_pages=max_pages,
            page_number=page_number,
            preload_radius=preload_radius,
            environment=environment,
            timings=timings,
            label=label,
        )
    return RenderedDocument(
        label=label,
        docx_sha256=sha256_bytes(pdf_bytes),
        pdf_sha256=sha256_bytes(pdf_bytes),
        page_count=page_count,
        rendered_pages=tuple(pages),
        converter_label="PDF direct preview",
        renderer_label="pdftoppm",
        pdf_bytes=pdf_bytes,
    )


def render_docx_to_pages(
    docx_bytes: bytes,
    file_name: str,
    *,
    label: str,
    dpi: int,
    max_pages: int,
    page_number: int | None = None,
    preload_radius: int = DEFAULT_PREVIEW_PRELOAD_RADIUS,
    environment: PreviewEnvironment,
    timings: dict[str, float | int | str] | None = None,
) -> RenderedDocument:
    if not environment.high_fidelity_available:
        return render_docx_low_fidelity_to_pages(
            docx_bytes,
            file_name,
            label=label,
            max_pages=max_pages,
            page_number=page_number,
            preload_radius=preload_radius,
            environment=environment,
            timings=timings,
        )

    with tempfile.TemporaryDirectory(prefix="report_preview_") as tmp_name:
        tmp_dir = Path(tmp_name)
        pdf_bytes = convert_docx_bytes_to_pdf_cached(
            docx_bytes,
            file_name,
            tmp_dir,
            environment,
            timings=timings,
            label=label,
        )
        pages, page_count = render_pdf_bytes_to_png_pages(
            pdf_bytes,
            tmp_dir,
            dpi=dpi,
            max_pages=max_pages,
            page_number=page_number,
            preload_radius=preload_radius,
            environment=environment,
            timings=timings,
            label=label,
        )
    return RenderedDocument(
        label=label,
        docx_sha256=sha256_bytes(docx_bytes),
        pdf_sha256=sha256_bytes(pdf_bytes),
        page_count=page_count,
        rendered_pages=tuple(pages),
        converter_label=environment.converter_label,
        renderer_label=environment.renderer_label,
        pdf_bytes=pdf_bytes,
    )


def render_docx_low_fidelity_to_pages(
    docx_bytes: bytes,
    file_name: str,
    *,
    label: str,
    max_pages: int,
    page_number: int | None = None,
    preload_radius: int = DEFAULT_PREVIEW_PRELOAD_RADIUS,
    environment: PreviewEnvironment,
    timings: dict[str, float | int | str] | None = None,
) -> RenderedDocument:
    if not environment.low_fidelity_available or Image is None or ImageDraw is None:
        raise PreviewUnavailable("Low-fidelity DOCX preview requires Pillow and python-docx.")
    started_at = time.perf_counter()
    lines = _extract_docx_preview_lines(docx_bytes)
    _record_elapsed(timings, f"{label}.docx_read_seconds", started_at)
    pages, page_count = _render_text_pages_to_png(
        lines,
        title=file_name or f"{label}.docx",
        max_pages=max_pages,
        page_number=page_number,
        preload_radius=preload_radius,
    )
    low_hash = sha256_bytes(b"low-fidelity-docx-preview\0" + docx_bytes)
    return RenderedDocument(
        label=label,
        docx_sha256=sha256_bytes(docx_bytes),
        pdf_sha256=low_hash,
        page_count=page_count,
        rendered_pages=tuple(pages),
        converter_label=environment.low_fidelity_label,
        renderer_label="Pillow text renderer",
        pdf_bytes=b"",
    )


def warm_docx_pdf_cache_batch(
    documents: list[tuple[str, bytes, str]],
    environment: PreviewEnvironment,
    *,
    timings: dict[str, float | int | str] | None = None,
) -> None:
    if not environment.high_fidelity_available or not environment.word_available:
        return

    misses: list[tuple[str, bytes, str, Path]] = []
    for label, docx_bytes, file_name in documents:
        cache_key = _converter_cache_key(docx_bytes, file_name, environment)
        cache_path = _cache_path("pdf", cache_key, ".pdf")
        if _valid_cached_pdf(cache_path):
            _increment_timing(timings, "word_batch.pdf_cache_hits")
            continue
        misses.append((label, docx_bytes, file_name, cache_path))

    if len(misses) < 2:
        return

    powershell_path = shutil.which("powershell") or shutil.which("pwsh")
    if not powershell_path:
        return

    started_at = time.perf_counter()
    with tempfile.TemporaryDirectory(prefix="report_preview_word_batch_") as tmp_name:
        tmp_dir = Path(tmp_name)
        manifest: list[dict[str, str]] = []
        for index, (label, docx_bytes, file_name, cache_path) in enumerate(misses, start=1):
            input_path = tmp_dir / f"{index}-{safe_file_name(file_name or f'{label}.docx')}"
            output_path = tmp_dir / f"{index}-{label}.pdf"
            input_path.write_bytes(docx_bytes)
            manifest.append(
                {
                    "label": label,
                    "input": str(input_path),
                    "output": str(output_path),
                    "cache": str(cache_path),
                }
            )

        manifest_path = tmp_dir / "word_batch_manifest.json"
        manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
        script_path = tmp_dir / "word_export_pdf_batch.ps1"
        script_path.write_text(
            r"""
param([Parameter(Mandatory=$true)][string]$ManifestPath)
$ErrorActionPreference = "Stop"
$items = Get-Content -LiteralPath $ManifestPath -Raw -Encoding UTF8 | ConvertFrom-Json
$word = $null
$document = $null
try {
    $word = New-Object -ComObject Word.Application
    $word.Visible = $false
    $word.DisplayAlerts = 0
    try { $word.AutomationSecurity = 3 } catch {}
    foreach ($item in $items) {
        $document = $word.Documents.Open($item.input, $false, $true)
        $document.ExportAsFixedFormat($item.output, 17)
        $document.Close($false)
        try { [void][System.Runtime.InteropServices.Marshal]::FinalReleaseComObject($document) } catch {}
        $document = $null
        if (-not (Test-Path -LiteralPath $item.output)) {
            throw "Microsoft Word did not create a PDF for $($item.label)."
        }
    }
} finally {
    if ($document -ne $null) {
        try { $document.Close($false) } catch {}
        try { [void][System.Runtime.InteropServices.Marshal]::FinalReleaseComObject($document) } catch {}
    }
    if ($word -ne $null) {
        try { $word.Quit() } catch {}
        try { [void][System.Runtime.InteropServices.Marshal]::FinalReleaseComObject($word) } catch {}
    }
    [GC]::Collect()
    [GC]::WaitForPendingFinalizers()
}
""".strip(),
            encoding="utf-8",
        )
        try:
            _run_command(
                [
                    powershell_path,
                    "-NoProfile",
                    "-NonInteractive",
                    "-ExecutionPolicy",
                    "Bypass",
                    "-File",
                    str(script_path),
                    "-ManifestPath",
                    str(manifest_path),
                ],
                tmp_dir,
                timeout=180,
            )
        except PreviewUnavailable:
            return

        for item in manifest:
            output_path = Path(item["output"])
            cache_path = Path(item["cache"])
            if output_path.exists():
                pdf_bytes = output_path.read_bytes()
                _validate_pdf_bytes(pdf_bytes, "Microsoft Word")
                temp_cache = cache_path.with_suffix(".tmp")
                temp_cache.write_bytes(pdf_bytes)
                os.replace(temp_cache, cache_path)
                _increment_timing(timings, "word_batch.pdf_cache_writes")
    _record_elapsed(timings, "word_batch.total_seconds", started_at)


def convert_docx_bytes_to_pdf_cached(
    docx_bytes: bytes,
    file_name: str,
    working_dir: Path,
    environment: PreviewEnvironment,
    *,
    timings: dict[str, float | int | str] | None = None,
    label: str = "document",
) -> bytes:
    cache_key = _converter_cache_key(docx_bytes, file_name, environment)
    cache_path = _cache_path("pdf", cache_key, ".pdf")
    if _valid_cached_pdf(cache_path):
        _increment_timing(timings, f"{label}.pdf_cache_hit")
        return cache_path.read_bytes()

    _increment_timing(timings, f"{label}.pdf_cache_miss")
    started_at = time.perf_counter()
    input_path = working_dir / safe_file_name(file_name or f"{label}.docx")
    input_path.write_bytes(docx_bytes)
    pdf_bytes = convert_docx_to_pdf(input_path, working_dir, environment)
    _validate_pdf_bytes(pdf_bytes, environment.converter_label)
    cache_temp = cache_path.with_suffix(".tmp")
    cache_temp.write_bytes(pdf_bytes)
    os.replace(cache_temp, cache_path)
    _record_elapsed(timings, f"{label}.docx_to_pdf_seconds", started_at)
    return pdf_bytes


def convert_docx_to_pdf(input_path: Path, working_dir: Path, environment: PreviewEnvironment) -> bytes:
    errors: list[str] = []
    converters: list[tuple[str, Callable[[], bytes]]] = []
    if environment.word_available:
        converters.append(("Microsoft Word", lambda: _convert_docx_with_word(input_path, working_dir)))
    if environment.libreoffice_path:
        converters.append(
            (
                "LibreOffice",
                lambda: _convert_docx_with_libreoffice(input_path, working_dir, environment.libreoffice_path or ""),
            )
        )
    if environment.gotenberg_url:
        converters.append(("Gotenberg", lambda: _convert_docx_with_gotenberg(input_path, environment.gotenberg_url or "")))

    for converter_name, converter in converters:
        try:
            pdf_bytes = converter()
            _validate_pdf_bytes(pdf_bytes, converter_name)
            return pdf_bytes
        except PreviewUnavailable as exc:
            errors.append(f"{converter_name}: {exc}")

    if errors:
        raise PreviewUnavailable("No high-fidelity DOCX converter succeeded. " + " | ".join(errors))
    raise PreviewUnavailable("No high-fidelity DOCX converter is configured.")


def render_pdf_to_png_pages(
    pdf_bytes: bytes,
    working_dir: Path,
    *,
    dpi: int,
    max_pages: int,
    environment: PreviewEnvironment,
) -> list[RenderedPage]:
    pages, _ = render_pdf_bytes_to_png_pages(
        pdf_bytes,
        working_dir,
        dpi=dpi,
        max_pages=max_pages,
        page_number=None,
        preload_radius=0,
        environment=environment,
    )
    return pages


def render_pdf_bytes_to_png_pages(
    pdf_bytes: bytes,
    working_dir: Path,
    *,
    dpi: int,
    max_pages: int,
    page_number: int | None = None,
    preload_radius: int = DEFAULT_PREVIEW_PRELOAD_RADIUS,
    environment: PreviewEnvironment,
    timings: dict[str, float | int | str] | None = None,
    label: str = "document",
) -> tuple[list[RenderedPage], int]:
    if not environment.pdftoppm_path:
        raise PreviewUnavailable("pdftoppm is required to render PDF pages.")
    pdf_path = working_dir / "preview.pdf"
    pdf_path.write_bytes(pdf_bytes)
    started_at = time.perf_counter()
    page_count = min(_pdf_page_count(pdf_path, working_dir), max_pages)
    _record_elapsed(timings, f"{label}.pdf_page_count_seconds", started_at)
    _set_timing(timings, f"{label}.pdf_page_count", page_count)
    pages: list[RenderedPage] = []
    for selected_page in _selected_page_numbers(page_count, page_number, preload_radius):
        cache_key = _page_cache_key(pdf_bytes, selected_page, dpi)
        cache_path = _cache_path("page", cache_key, ".png")
        if cache_path.exists() and cache_path.stat().st_size > 0:
            image_bytes = cache_path.read_bytes()
            width, height = _image_size(image_bytes)
            pages.append(RenderedPage(page_number=selected_page, image_bytes=image_bytes, width=width, height=height))
            _increment_timing(timings, f"{label}.page_cache_hits")
            continue

        _increment_timing(timings, f"{label}.page_cache_misses")
        started_at = time.perf_counter()
        out_prefix = working_dir / f"page-{selected_page}"
        command = [
            environment.pdftoppm_path,
            "-png",
            "-r",
            str(dpi),
            "-f",
            str(selected_page),
            "-l",
            str(selected_page),
            "-singlefile",
            str(pdf_path),
            str(out_prefix),
        ]
        _run_command(command, working_dir, timeout=45)
        image_path = out_prefix.with_suffix(".png")
        if not image_path.exists():
            raise PreviewUnavailable(f"Rendered page image was not created for page {selected_page}.")
        image_bytes = image_path.read_bytes()
        cache_temp = cache_path.with_suffix(".tmp")
        cache_temp.write_bytes(image_bytes)
        os.replace(cache_temp, cache_path)
        _record_elapsed(timings, f"{label}.page_{selected_page}_render_seconds", started_at)
        width, height = _image_size(image_bytes)
        pages.append(RenderedPage(page_number=selected_page, image_bytes=image_bytes, width=width, height=height))
    _set_timing(timings, f"{label}.rendered_page_numbers", ",".join(str(page.page_number) for page in pages))
    return pages, page_count


def build_issue_items(issues: Iterable, page_count: int) -> list[PreviewItem]:
    issues = list(issues or [])
    items: list[PreviewItem] = []
    for index, issue in enumerate(issues, start=1):
        paragraph_index = getattr(issue, "paragraph_index", None)
        severity = getattr(issue, "severity", "warning") or "warning"
        category = getattr(issue, "category", "document") or "document"
        message = getattr(issue, "message", "") or ""
        items.append(
            PreviewItem(
                item_id=f"issue-{index}",
                kind=category,
                severity=severity,
                page_number=1,
                title=f"{category.title()} issue",
                message=message,
                rule_id=f"{category}.{severity}",
                confidence=0.0,
                approximate=True,
                rect=None,
                source_anchor=f"paragraph:{paragraph_index}" if paragraph_index is not None else "document",
                location_status="unlocated",
            )
        )
    return items


def build_change_items(
    actions: Iterable,
    page_count: int,
    *,
    pdf_bytes: bytes = b"",
    pdftotext_path: str | None = None,
) -> list[PreviewItem]:
    actions = list(actions or [])
    lines = extract_pdf_text_lines(pdf_bytes, pdftotext_path) if pdf_bytes and pdftotext_path else []
    items: list[PreviewItem] = []
    for index, action in enumerate(actions, start=1):
        paragraph_index = getattr(action, "paragraph_index", None)
        kind = getattr(action, "kind", "change") or "change"
        message = getattr(action, "message", "") or ""
        before = getattr(action, "before", "") or ""
        after = getattr(action, "after", "") or ""
        located = locate_action_in_pdf(action, lines)
        page_number, rect = located if located else (1, None)
        location_status = "located" if rect is not None else "unlocated"
        confidence = float(getattr(action, "confidence", 1.0)) if rect is not None else 0.0
        try:
            action.page_index = page_number - 1 if rect is not None else None
            action.modified_bbox = rect
            action.location_status = location_status
        except Exception:
            pass
        items.append(
            PreviewItem(
                item_id=getattr(action, "change_id", "") or f"change-{index}",
                kind=kind,
                severity="change",
                page_number=page_number,
                title=f"{kind.title()} change",
                message=message,
                before=before,
                after=after,
                rule_id=f"{kind}.format",
                confidence=confidence,
                approximate=rect is None,
                rect=rect,
                object_id=getattr(action, "object_id", "") or f"change-{index}",
                source_anchor=getattr(action, "source_anchor", "")
                or (f"paragraph:{paragraph_index}" if paragraph_index is not None else "document"),
                coordinate_space=getattr(action, "coordinate_space", "pdf_normalized_top_left"),
                location_status=location_status,
                is_cross_column=bool(getattr(action, "is_cross_column", False)),
            )
        )
    return merge_preview_items(items)


def build_pixel_diff_items(
    original_pages: Iterable[RenderedPage],
    modified_pages: Iterable[RenderedPage],
    *,
    threshold: int = 42,
    min_area_px: int = 24,
    max_area_ratio: float = 0.04,
    max_markers_per_page: int = PIXEL_DIFF_MAX_MARKERS_PER_PAGE,
    max_total_area_ratio: float = PIXEL_DIFF_MAX_TOTAL_AREA_RATIO,
    diagnostics: dict[int, dict[str, float | int | str]] | None = None,
) -> list[PreviewItem]:
    if Image is None or ImageChops is None or ImageFilter is None:
        return []
    original_by_page = {page.page_number: page for page in original_pages}
    modified_by_page = {page.page_number: page for page in modified_pages}
    items: list[PreviewItem] = []
    for page_number in sorted(set(original_by_page).intersection(modified_by_page)):
        boxes = _pixel_diff_boxes(
            original_by_page[page_number].image_bytes,
            modified_by_page[page_number].image_bytes,
            threshold=threshold,
            min_area_px=min_area_px,
            max_area_ratio=max_area_ratio,
        )
        raw_candidate_count = len(boxes)
        total_area_ratio = round(sum(width * height for _x, _y, width, height in boxes), 6)
        status = "located"
        if len(boxes) > max_markers_per_page:
            status = "suppressed_too_many_boxes"
            boxes = []
        elif total_area_ratio > max_total_area_ratio:
            status = "suppressed_total_area_too_large"
            boxes = []
        if diagnostics is not None:
            diagnostics[page_number] = {
                "status": status,
                "candidate_count": raw_candidate_count,
                "final_count": len(boxes),
                "candidate_total_area_ratio": total_area_ratio,
                "max_markers_per_page": max_markers_per_page,
                "max_total_area_ratio": max_total_area_ratio,
            }
        for index, box in enumerate(boxes, start=1):
            x, y, width, height = box
            items.append(
                PreviewItem(
                    item_id=f"pixel-diff-{page_number}-{index}",
                    kind="pixel_diff",
                    severity="change",
                    page_number=page_number,
                    title="Pixel-level difference",
                    message="Rendered page pixels differ in this region.",
                    rule_id="pixel.diff",
                    confidence=1.0,
                    approximate=False,
                    rect=(x, y, width, height),
                    object_id=f"pixel-diff-{page_number}-{index}",
                    source_anchor=f"rendered-page:{page_number}",
                    coordinate_space="image_normalized_top_left",
                    location_status="located",
                )
            )
    return items


def _pixel_diff_boxes(
    original_png: bytes,
    modified_png: bytes,
    *,
    threshold: int,
    min_area_px: int,
    max_area_ratio: float,
) -> list[tuple[float, float, float, float]]:
    try:
        original = Image.open(io.BytesIO(original_png)).convert("RGB")
        modified = Image.open(io.BytesIO(modified_png)).convert("RGB")
    except Exception:
        return []
    if original.size != modified.size:
        modified = modified.resize(original.size)
    width, height = original.size
    if width <= 0 or height <= 0:
        return []
    diff = ImageChops.difference(original, modified).convert("L")
    mask = diff.point(lambda value: 255 if value >= threshold else 0)
    mask = mask.filter(ImageFilter.MaxFilter(5))
    boxes_px = _connected_component_boxes(mask, min_area_px=min_area_px)
    boxes_px = _merge_nearby_pixel_boxes(boxes_px, gap_px=max(6, int(min(width, height) * 0.006)))
    normalized: list[tuple[float, float, float, float]] = []
    for x0, y0, x1, y1 in boxes_px:
        box_width = max(0, x1 - x0 + 1)
        box_height = max(0, y1 - y0 + 1)
        area_ratio = (box_width * box_height) / float(width * height)
        if box_width < 3 or box_height < 3 or area_ratio <= 0:
            continue
        if area_ratio > max_area_ratio or (box_width / width > 0.92 and box_height / height > 0.75):
            continue
        padding = 4
        left = max(0, x0 - padding)
        top = max(0, y0 - padding)
        right = min(width, x1 + 1 + padding)
        bottom = min(height, y1 + 1 + padding)
        normalized.append((left / width, top / height, (right - left) / width, (bottom - top) / height))
    return normalized


def _connected_component_boxes(mask, *, min_area_px: int) -> list[tuple[int, int, int, int]]:
    width, height = mask.size
    pixels = mask.load()
    visited = bytearray(width * height)
    boxes: list[tuple[int, int, int, int]] = []
    for y in range(height):
        for x in range(width):
            offset = y * width + x
            if visited[offset] or not pixels[x, y]:
                continue
            stack = [(x, y)]
            visited[offset] = 1
            area = 0
            x0 = x1 = x
            y0 = y1 = y
            while stack:
                cx, cy = stack.pop()
                area += 1
                x0, x1 = min(x0, cx), max(x1, cx)
                y0, y1 = min(y0, cy), max(y1, cy)
                for nx, ny in ((cx - 1, cy), (cx + 1, cy), (cx, cy - 1), (cx, cy + 1)):
                    if nx < 0 or nx >= width or ny < 0 or ny >= height:
                        continue
                    n_offset = ny * width + nx
                    if visited[n_offset] or not pixels[nx, ny]:
                        continue
                    visited[n_offset] = 1
                    stack.append((nx, ny))
            if area >= min_area_px:
                boxes.append((x0, y0, x1, y1))
    return boxes


def _merge_nearby_pixel_boxes(
    boxes: list[tuple[int, int, int, int]],
    *,
    gap_px: int,
) -> list[tuple[int, int, int, int]]:
    merged: list[tuple[int, int, int, int]] = []
    for box in boxes:
        current = box
        did_merge = True
        while did_merge:
            did_merge = False
            for index, existing in enumerate(merged):
                if _pixel_boxes_close(current, existing, gap_px):
                    current = (
                        min(current[0], existing[0]),
                        min(current[1], existing[1]),
                        max(current[2], existing[2]),
                        max(current[3], existing[3]),
                    )
                    merged.pop(index)
                    did_merge = True
                    break
        merged.append(current)
    return merged


def _pixel_boxes_close(
    left: tuple[int, int, int, int],
    right: tuple[int, int, int, int],
    gap_px: int,
) -> bool:
    return not (
        left[2] + gap_px < right[0]
        or right[2] + gap_px < left[0]
        or left[3] + gap_px < right[1]
        or right[3] + gap_px < left[1]
    )


def extract_pdf_text_lines(pdf_bytes: bytes, pdftotext_path: str | None) -> list[PdfTextLine]:
    if not pdf_bytes or not pdftotext_path:
        return []
    with tempfile.TemporaryDirectory(prefix="preview_geometry_") as tmp_name:
        input_path = Path(tmp_name) / "document.pdf"
        output_path = Path(tmp_name) / "layout.xhtml"
        input_path.write_bytes(pdf_bytes)
        completed = subprocess.run(
            [pdftotext_path, "-bbox-layout", str(input_path), str(output_path)],
            cwd=tmp_name,
            capture_output=True,
            timeout=45,
            check=False,
        )
        if completed.returncode != 0 or not output_path.exists():
            return []
        try:
            return parse_pdftotext_bbox_layout(output_path.read_bytes())
        except (ET.ParseError, ValueError, OSError):
            return []


def parse_pdftotext_bbox_layout(payload: bytes | str) -> list[PdfTextLine]:
    root = ET.fromstring(payload)
    lines: list[PdfTextLine] = []
    pages = [node for node in root.iter() if _xml_local_name(node.tag) == "page"]
    for page_number, page in enumerate(pages, start=1):
        page_width = float(page.attrib.get("width", 0) or 0)
        page_height = float(page.attrib.get("height", 0) or 0)
        if page_width <= 0 or page_height <= 0:
            continue
        for line in (node for node in page.iter() if _xml_local_name(node.tag) == "line"):
            words = [node for node in line.iter() if _xml_local_name(node.tag) == "word"]
            text = normalize_pdf_text(" ".join("".join(word.itertext()) for word in words))
            if not text:
                continue
            try:
                bbox = (
                    float(line.attrib["xMin"]),
                    float(line.attrib["yMin"]),
                    float(line.attrib["xMax"]),
                    float(line.attrib["yMax"]),
                )
            except (KeyError, ValueError):
                continue
            lines.append(PdfTextLine(page_number, page_width, page_height, text, bbox))
    return lines


def locate_action_in_pdf(action, lines: list[PdfTextLine]):
    kind = getattr(action, "kind", "")
    after = normalize_pdf_text(getattr(action, "after", "") or getattr(action, "modified_text", ""))
    if not after:
        return None
    candidates: list[tuple[int, PdfTextLine]] = []
    for line in lines:
        score = _line_match_score(kind, after, line.text)
        if score > 0:
            candidates.append((score, line))
    if not candidates:
        return None
    candidates.sort(key=lambda item: item[0], reverse=True)
    best_score = candidates[0][0]
    best = [line for score, line in candidates if score == best_score]
    if len(best) != 1:
        return None
    line = best[0]
    rect = normalize_pdf_bbox(line.bbox, line.page_width, line.page_height, padding=0.004)
    is_cross_column = bool(getattr(action, "is_cross_column", False))
    if not is_cross_column and rect[2] > 0.48:
        return None
    return line.page_number, rect


def _line_match_score(kind: str, target: str, candidate: str) -> int:
    left = target.casefold()
    right = normalize_pdf_text(candidate).casefold()
    if not left or not right:
        return 0
    if right == left:
        return 100 + len(left)
    if kind == "equation" and re.fullmatch(r"\(\d+\)", left):
        return 90 if right.endswith(left) else 0
    if left in right:
        return 70 + len(left)
    if right in left and len(right) >= 8:
        return 40 + len(right)
    return 0


def normalize_pdf_bbox(
    bbox: tuple[float, float, float, float],
    page_width: float,
    page_height: float,
    *,
    padding: float = 0.0,
) -> tuple[float, float, float, float]:
    x_min, y_min, x_max, y_max = bbox
    x = max(0.0, x_min / page_width - padding)
    y = max(0.0, y_min / page_height - padding)
    right = min(1.0, x_max / page_width + padding)
    bottom = min(1.0, y_max / page_height + padding)
    return (x, y, max(0.0, right - x), max(0.0, bottom - y))


def transform_pdf_bbox(
    bbox: tuple[float, float, float, float],
    crop_box: tuple[float, float, float, float],
    rotation: int = 0,
) -> tuple[float, float, float, float]:
    """Normalize a PDF user-space box for CropBox and page rotation."""
    x0, y0, x1, y1 = bbox
    cx0, cy0, cx1, cy1 = crop_box
    width, height = cx1 - cx0, cy1 - cy0
    points = [(x0 - cx0, y0 - cy0), (x1 - cx0, y1 - cy0)]
    rotation = rotation % 360
    if rotation == 90:
        points = [(y, width - x) for x, y in points]
        width, height = height, width
    elif rotation == 180:
        points = [(width - x, height - y) for x, y in points]
    elif rotation == 270:
        points = [(height - y, x) for x, y in points]
        width, height = height, width
    xs, ys = zip(*points)
    x_min, x_max = min(xs), max(xs)
    y_min, y_max = min(ys), max(ys)
    return (x_min / width, (height - y_max) / height, (x_max - x_min) / width, (y_max - y_min) / height)


def merge_preview_items(items: list[PreviewItem]) -> list[PreviewItem]:
    merged: list[PreviewItem] = []
    by_object: dict[tuple[str, int], int] = {}
    for item in items:
        if item.rect is None:
            merged.append(item)
            continue
        key = (item.object_id, item.page_number)
        existing_index = by_object.get(key)
        if existing_index is None:
            by_object[key] = len(merged)
            merged.append(item)
            continue
        existing = merged[existing_index]
        union = _rect_union(existing.rect, item.rect)
        merged[existing_index] = PreviewItem(
            **{**existing.__dict__, "rect": union, "after": "\n".join(filter(None, [existing.after, item.after]))}
        )
    return merged


def _rect_union(left, right):
    lx, ly, lw, lh = left
    rx, ry, rw, rh = right
    x0, y0 = min(lx, rx), min(ly, ry)
    x1, y1 = max(lx + lw, rx + rw), max(ly + lh, ry + rh)
    return (x0, y0, x1 - x0, y1 - y0)


def normalize_pdf_text(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def _xml_local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def estimate_page_number(paragraph_index: int | None, *, max_index: int, page_count: int) -> int:
    if page_count <= 1 or paragraph_index is None or max_index <= 0:
        return 1
    ratio = min(max(paragraph_index / max_index, 0.0), 0.999)
    return min(page_count, int(ratio * page_count) + 1)


def estimated_rect(paragraph_index: int | None, *, max_index: int) -> tuple[float, float, float, float]:
    if paragraph_index is None or max_index <= 0:
        return (0.08, 0.16, 0.84, 0.07)
    ratio = min(max(paragraph_index / max_index, 0.0), 1.0)
    y = 0.14 + (ratio * 0.7)
    return (0.08, min(y, 0.84), 0.84, 0.065)


def items_for_page(items: Iterable[PreviewItem], page_number: int) -> list[PreviewItem]:
    return [item for item in items if item.page_number == page_number]


def image_data_uri(image_bytes: bytes) -> str:
    import base64

    return "data:image/png;base64," + base64.b64encode(image_bytes).decode("ascii")


def render_page_html(
    page: RenderedPage,
    markers: Iterable[PreviewItem] = (),
    *,
    marker_visible: bool = True,
    label: str = "",
) -> str:
    marker_html = ""
    if marker_visible:
        for marker in markers:
            if marker.rect is None or marker.location_status != "located":
                continue
            x, y, width, height = marker.rect
            severity = html.escape(marker.severity)
            title = html.escape(marker.title)
            marker_html += (
                f'<span class="preview-marker preview-marker-{severity}" '
                f'data-change-id="{html.escape(marker.item_id)}" '
                f'style="left:{x * 100:.2f}%; top:{y * 100:.2f}%; '
                f'width:{width * 100:.2f}%; height:{height * 100:.2f}%;" '
                f'title="{title}"></span>'
            )
    data_uri = image_data_uri(page.image_bytes)
    alt = html.escape(label or f"Page {page.page_number}")
    return (
        '<div class="preview-page-shell">'
        f'<img class="preview-page-image" src="{data_uri}" alt="{alt}" />'
        f'<div class="preview-overlay" aria-hidden="true">{marker_html}</div>'
        "</div>"
    )


def _sibling_executable(path: str | None, stem: str) -> str | None:
    if not path:
        return None
    source = Path(path)
    suffix = source.suffix
    candidate = source.with_name(stem + suffix)
    return str(candidate) if candidate.exists() else None


def safe_file_name(file_name: str) -> str:
    stem = Path(file_name or "document.docx").stem.strip() or "document"
    suffix = Path(file_name or "document.docx").suffix.lower() or ".docx"
    if suffix not in {".docx", ".pdf"}:
        suffix = ".docx"
    stem = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "_", stem)
    stem = re.sub(r"\s+", "_", stem)
    return f"{stem[:90]}{suffix}"


def _clean_gotenberg_url(value: str | None) -> str | None:
    if not value:
        return None
    return value.strip().rstrip("/")


def _low_fidelity_available() -> bool:
    if Image is None or ImageDraw is None:
        return False
    try:
        import docx  # noqa: F401
    except Exception:
        return False
    return True


def _high_fidelity_converter_label(
    word: _WordCapability,
    libreoffice_path: str | None,
    gotenberg_url: str | None,
) -> str:
    if word.available:
        return "Microsoft Word"
    if libreoffice_path:
        return "LibreOffice headless"
    if gotenberg_url:
        return "Gotenberg"
    return "Unavailable"


def _detect_word_capability() -> _WordCapability:
    if os.name != "nt":
        return _WordCapability(False, "not_installed", error="Microsoft Word preview is Windows-only.")

    discovered_path = _first_existing_path(_find_word_paths())
    powershell_path = shutil.which("powershell") or shutil.which("pwsh")
    if not powershell_path:
        status = "installed_com_unavailable" if discovered_path else "not_installed"
        return _WordCapability(False, status, path=discovered_path, error="PowerShell was not found.")

    script = r"""
$ErrorActionPreference = "Stop"
$word = $null
try {
    $word = New-Object -ComObject Word.Application
    $word.Visible = $false
    $word.DisplayAlerts = 0
    $version = $word.Version
    Write-Output "OK|$version"
} finally {
    if ($word -ne $null) {
        try { $word.Quit() } catch {}
        try { [void][System.Runtime.InteropServices.Marshal]::FinalReleaseComObject($word) } catch {}
    }
    [GC]::Collect()
    [GC]::WaitForPendingFinalizers()
}
"""
    try:
        completed = subprocess.run(
            [powershell_path, "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass", "-Command", script],
            timeout=12,
            check=False,
            text=True,
            encoding="utf-8",
            errors="replace",
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )
    except subprocess.TimeoutExpired:
        return _WordCapability(False, "com_timeout", path=discovered_path, error="COM validation timed out.")

    output = (completed.stdout or "").strip()
    if completed.returncode == 0 and output.startswith("OK|"):
        return _WordCapability(True, "usable", path=discovered_path)

    error = (completed.stderr or completed.stdout or "").strip() or "COM validation failed."
    status = "installed_com_unavailable" if discovered_path else "not_installed"
    return _WordCapability(False, status, path=discovered_path, error=error)


def _find_word_paths() -> list[Path]:
    candidates: list[Path] = []
    for env_name in ("WORD_PATH", "WINWORD_PATH"):
        value = os.environ.get(env_name)
        if value:
            candidates.append(Path(value))
    for command_name in ("WINWORD.EXE", "winword.exe", "winword"):
        resolved = shutil.which(command_name)
        if resolved:
            candidates.append(Path(resolved))
    candidates.extend(_word_paths_from_registry())

    program_dirs = [
        os.environ.get("PROGRAMFILES"),
        os.environ.get("PROGRAMFILES(X86)"),
        os.environ.get("LOCALAPPDATA"),
    ]
    office_roots = [
        ("Microsoft Office", "root", "Office16", "WINWORD.EXE"),
        ("Microsoft Office", "Office16", "WINWORD.EXE"),
        ("Microsoft Office", "Office15", "WINWORD.EXE"),
        ("Microsoft Office", "Office14", "WINWORD.EXE"),
        ("Microsoft Office", "Office12", "WINWORD.EXE"),
    ]
    for base in program_dirs:
        if not base:
            continue
        base_path = Path(base)
        for parts in office_roots:
            candidates.append(base_path.joinpath(*parts))
    return _dedupe_paths(candidates)


def _word_paths_from_registry() -> list[Path]:
    powershell_path = shutil.which("powershell") or shutil.which("pwsh")
    if not powershell_path or os.name != "nt":
        return []
    script = r"""
$keys = @(
  "HKLM:\SOFTWARE\Microsoft\Windows\CurrentVersion\App Paths\WINWORD.EXE",
  "HKLM:\SOFTWARE\WOW6432Node\Microsoft\Windows\CurrentVersion\App Paths\WINWORD.EXE",
  "HKCU:\SOFTWARE\Microsoft\Windows\CurrentVersion\App Paths\WINWORD.EXE"
)
foreach ($key in $keys) {
  $item = Get-ItemProperty -Path $key -ErrorAction SilentlyContinue
  if ($item -ne $null) {
    $default = $item.'(default)'
    if ($default) { Write-Output $default }
    if ($item.Path) { Write-Output (Join-Path $item.Path "WINWORD.EXE") }
  }
}
"""
    try:
        completed = subprocess.run(
            [powershell_path, "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass", "-Command", script],
            timeout=8,
            check=False,
            text=True,
            encoding="utf-8",
            errors="replace",
            stdout=subprocess.PIPE,
            stderr=subprocess.DEVNULL,
        )
    except subprocess.TimeoutExpired:
        return []
    return [Path(line.strip()) for line in (completed.stdout or "").splitlines() if line.strip()]


def _find_libreoffice() -> str | None:
    for env_name in ("LIBREOFFICE_PATH", "SOFFICE_PATH"):
        value = os.environ.get(env_name)
        if value and Path(value).exists():
            return value
    for candidate in ("soffice", "soffice.exe", "libreoffice", "libreoffice.exe"):
        resolved = shutil.which(candidate)
        if resolved:
            return resolved
    common_paths = [
        Path(os.environ.get("PROGRAMFILES", "")) / "LibreOffice" / "program" / "soffice.exe",
        Path(os.environ.get("PROGRAMFILES(X86)", "")) / "LibreOffice" / "program" / "soffice.exe",
        Path("/usr/bin/soffice"),
        Path("/usr/local/bin/soffice"),
        Path("/usr/bin/libreoffice"),
        Path("/usr/local/bin/libreoffice"),
        Path("/Applications/LibreOffice.app/Contents/MacOS/soffice"),
    ]
    for path in common_paths:
        if path and path.exists():
            return str(path)
    return None


def _convert_docx_with_word(input_path: Path, working_dir: Path) -> bytes:
    powershell_path = shutil.which("powershell") or shutil.which("pwsh")
    if not powershell_path:
        raise PreviewUnavailable("PowerShell is required for Microsoft Word COM export.")

    output_path = working_dir / f"{input_path.stem}.word.pdf"
    script_path = working_dir / "word_export_pdf.ps1"
    script_path.write_text(
        r"""
param(
    [Parameter(Mandatory=$true)][string]$InputPath,
    [Parameter(Mandatory=$true)][string]$OutputPath
)
$ErrorActionPreference = "Stop"
$word = $null
$document = $null
try {
    $word = New-Object -ComObject Word.Application
    $word.Visible = $false
    $word.DisplayAlerts = 0
    try { $word.AutomationSecurity = 3 } catch {}
    $document = $word.Documents.Open($InputPath, $false, $true)
    $document.ExportAsFixedFormat($OutputPath, 17)
} finally {
    if ($document -ne $null) {
        try { $document.Close($false) } catch {}
        try { [void][System.Runtime.InteropServices.Marshal]::FinalReleaseComObject($document) } catch {}
    }
    if ($word -ne $null) {
        try { $word.Quit() } catch {}
        try { [void][System.Runtime.InteropServices.Marshal]::FinalReleaseComObject($word) } catch {}
    }
    [GC]::Collect()
    [GC]::WaitForPendingFinalizers()
}
if (-not (Test-Path -LiteralPath $OutputPath)) {
    throw "Microsoft Word did not create a PDF."
}
""".strip(),
        encoding="utf-8",
    )
    _run_command(
        [
            powershell_path,
            "-NoProfile",
            "-NonInteractive",
            "-ExecutionPolicy",
            "Bypass",
            "-File",
            str(script_path),
            "-InputPath",
            str(input_path),
            "-OutputPath",
            str(output_path),
        ],
        working_dir,
        timeout=120,
    )
    if not output_path.exists():
        raise PreviewUnavailable("Microsoft Word did not create a PDF.")
    return output_path.read_bytes()


def _convert_docx_with_libreoffice(input_path: Path, working_dir: Path, soffice_path: str) -> bytes:
    output_dir = working_dir / "pdf"
    output_dir.mkdir(parents=True, exist_ok=True)
    profile_dir = working_dir / "lo_profile"
    profile_dir.mkdir(parents=True, exist_ok=True)
    command = [
        soffice_path,
        "--headless",
        "--nologo",
        "--nofirststartwizard",
        f"-env:UserInstallation=file:///{profile_dir.as_posix()}",
        "--convert-to",
        "pdf",
        "--outdir",
        str(output_dir),
        str(input_path),
    ]
    _run_command(command, working_dir, timeout=90)
    pdf_path = output_dir / f"{input_path.stem}.pdf"
    if not pdf_path.exists():
        candidates = list(output_dir.glob("*.pdf"))
        if not candidates:
            raise PreviewUnavailable("LibreOffice did not create a PDF.")
        pdf_path = candidates[0]
    return pdf_path.read_bytes()


def _convert_docx_with_gotenberg(input_path: Path, gotenberg_url: str) -> bytes:
    boundary = "----codex-preview-" + hashlib.sha1(os.urandom(16)).hexdigest()
    content_type = mimetypes.guess_type(str(input_path))[0] or "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    body = io.BytesIO()
    body.write(f"--{boundary}\r\n".encode("ascii"))
    body.write(
        (
            f'Content-Disposition: form-data; name="files"; filename="{input_path.name}"\r\n'
            f"Content-Type: {content_type}\r\n\r\n"
        ).encode("utf-8")
    )
    body.write(input_path.read_bytes())
    body.write(f"\r\n--{boundary}--\r\n".encode("ascii"))
    endpoint = urllib.parse.urljoin(gotenberg_url + "/", "forms/libreoffice/convert")
    request = urllib.request.Request(
        endpoint,
        data=body.getvalue(),
        headers={"Content-Type": f"multipart/form-data; boundary={boundary}"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=120) as response:
            return response.read()
    except urllib.error.URLError as exc:
        raise PreviewUnavailable(f"Gotenberg conversion failed: {exc}") from exc


def _extract_docx_preview_lines(docx_bytes: bytes) -> list[str]:
    try:
        from docx import Document

        document = Document(io.BytesIO(docx_bytes))
    except Exception as exc:
        raise PreviewUnavailable(f"Could not read DOCX for low-fidelity preview: {exc}") from exc

    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = _normalize_preview_line(paragraph.text)
        if text:
            lines.extend(textwrap.wrap(text, width=92, replace_whitespace=False) or [text])
    for table in document.tables:
        lines.append("")
        lines.append("[Table]")
        for row in table.rows[:12]:
            cells = [_normalize_preview_line(cell.text) for cell in row.cells]
            line = " | ".join(cell for cell in cells if cell)
            if line:
                lines.extend(textwrap.wrap(line, width=92, replace_whitespace=False) or [line])
    if not lines:
        lines = ["No readable DOCX text was found for browser fallback preview."]
    return lines


def _render_text_pages_to_png(
    lines: list[str],
    *,
    title: str,
    max_pages: int,
    page_number: int | None = None,
    preload_radius: int = DEFAULT_PREVIEW_PRELOAD_RADIUS,
) -> tuple[list[RenderedPage], int]:
    if Image is None or ImageDraw is None:
        raise PreviewUnavailable("Pillow is required for low-fidelity preview.")
    font = _load_preview_font(22)
    meta_font = _load_preview_font(17)
    width, height = 1100, 1500
    margin_x, margin_y = 86, 86
    line_height = 32
    lines_per_page = max(1, (height - margin_y * 2 - 70) // line_height)
    chunks = [lines[index : index + lines_per_page] for index in range(0, len(lines), lines_per_page)]
    chunks = chunks[: max(1, max_pages)]
    page_count = max(1, len(chunks))
    page_numbers = _selected_page_numbers(page_count, page_number, preload_radius)
    pages: list[RenderedPage] = []
    for page_index in page_numbers:
        chunk = chunks[page_index - 1] if page_index - 1 < len(chunks) else []
        image = Image.new("RGB", (width, height), "white")
        draw = ImageDraw.Draw(image)
        draw.rectangle((42, 42, width - 42, height - 42), outline="#d9e2ef", width=2)
        draw.text((margin_x, 52), f"{title} - low-fidelity preview", fill="#334155", font=meta_font)
        draw.text((width - 170, 52), f"Page {page_index}", fill="#64748b", font=meta_font)
        y = margin_y + 30
        for line in chunk:
            draw.text((margin_x, y), line, fill="#0f172a", font=font)
            y += line_height
        output = io.BytesIO()
        image.save(output, format="PNG", optimize=True)
        pages.append(RenderedPage(page_number=page_index, image_bytes=output.getvalue(), width=width, height=height))
    return pages, page_count


def _load_preview_font(size: int):
    if ImageFont is None:
        return None
    font_candidates = [
        Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "msyh.ttc",
        Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "simsun.ttc",
        Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "arial.ttf",
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/Library/Fonts/Arial Unicode.ttf"),
    ]
    for path in font_candidates:
        if path.exists():
            try:
                return ImageFont.truetype(str(path), size=size)
            except Exception:
                continue
    return ImageFont.load_default()


def _normalize_preview_line(value: str | None) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def _validate_pdf_bytes(pdf_bytes: bytes, converter_name: str) -> None:
    if len(pdf_bytes or b"") < 512:
        raise PreviewUnavailable(f"{converter_name} returned an empty or incomplete PDF.")
    if not pdf_bytes.startswith(b"%PDF"):
        raise PreviewUnavailable(f"{converter_name} did not return a PDF file.")


def _pdf_page_count(pdf_path: Path, working_dir: Path) -> int:
    pdfinfo_path = _find_poppler_executable("pdfinfo")
    if not pdfinfo_path:
        return 1
    completed = _run_command([pdfinfo_path, str(pdf_path)], working_dir, timeout=30, capture_output=True)
    match = re.search(r"^Pages:\s+(\d+)\s*$", completed.stdout or "", re.MULTILINE)
    if not match:
        return 1
    return max(1, int(match.group(1)))


def _run_command(
    command: list[str],
    cwd: Path,
    *,
    timeout: int,
    capture_output: bool = True,
) -> subprocess.CompletedProcess[str]:
    try:
        completed = subprocess.run(
            command,
            cwd=str(cwd),
            timeout=timeout,
            check=False,
            text=True,
            encoding="utf-8",
            errors="replace",
            stdout=subprocess.PIPE if capture_output else None,
            stderr=subprocess.PIPE if capture_output else None,
        )
    except subprocess.TimeoutExpired as exc:
        raise PreviewUnavailable(f"Command timed out after {timeout}s: {command[0]}") from exc
    if completed.returncode != 0:
        message = (completed.stderr or completed.stdout or "").strip()
        raise PreviewUnavailable(message or f"Command failed: {command[0]}")
    return completed


def _image_size(image_bytes: bytes) -> tuple[int | None, int | None]:
    if Image is None:
        return None, None
    with Image.open(io.BytesIO(image_bytes)) as image:
        return image.size


def _max_paragraph_index(items: Iterable) -> int:
    max_index = 0
    for item in items:
        paragraph_index = getattr(item, "paragraph_index", None)
        if isinstance(paragraph_index, int):
            max_index = max(max_index, paragraph_index)
    return max_index


def _dedupe_paths(paths: Iterable[Path]) -> list[Path]:
    seen: set[str] = set()
    result: list[Path] = []
    for path in paths:
        key = str(path).lower()
        if key in seen:
            continue
        seen.add(key)
        result.append(path)
    return result


def _first_existing_path(paths: Iterable[Path]) -> str | None:
    for path in paths:
        if path.exists():
            return str(path)
    return None
