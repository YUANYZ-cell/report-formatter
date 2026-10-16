from __future__ import annotations

import json
from collections import Counter
from pathlib import Path
from typing import Iterable

from docx.text.paragraph import Paragraph

from .caption_detector import (
    detect_figure_captions,
    detect_table_captions,
    match_figure_caption,
    match_table_number,
)
from .equation_detector import detect_equation_numbers
from .object_scanner import detect_main_report_start, looks_like_heading
from .style_profile import StyleItem, StyleProfile
from .utils import ALIGNMENT_TO_STRING, normalize_space, points


FALLBACK_STYLE_PATH = Path(__file__).resolve().parents[1] / "templates" / "fallback_ieee_style.json"


def extract_style_profile(document, template_document=None) -> StyleProfile:
    source_document = template_document or document
    source_label = "uploaded template" if template_document is not None else "uploaded report"
    start_index = detect_main_report_start(source_document)
    fallback = load_fallback_profile()

    body = _infer_body_style(source_document, start_index) or fallback.body_text
    heading = _infer_heading_style(source_document, start_index) or fallback.section_heading
    figure = _infer_figure_caption_style(source_document, start_index) or _merge_fallback(
        fallback.figure_caption, body, "figure_caption"
    )
    table_number = _infer_table_number_style(source_document, start_index) or _merge_fallback(
        fallback.table_number, body, "table_number"
    )
    table_caption = _infer_table_caption_style(source_document, start_index) or _merge_fallback(
        fallback.table_caption, body, "table_caption"
    )
    equation = _infer_equation_style(source_document, start_index) or _merge_fallback(
        fallback.equation, body, "equation"
    )
    references = _infer_references_style(source_document, start_index) or fallback.references

    profile = StyleProfile(
        body_text=body,
        section_heading=heading,
        figure_caption=figure,
        table_number=table_number,
        table_caption=table_caption,
        equation=equation,
        equation_number=fallback.equation_number,
        references=references,
        source=source_label,
        warnings=[],
    )

    for item in [
        profile.body_text,
        profile.section_heading,
        profile.figure_caption,
        profile.table_number,
        profile.table_caption,
        profile.equation,
        profile.equation_number,
    ]:
        item.source = item.source or source_label
    return profile


def load_fallback_profile() -> StyleProfile:
    data = json.loads(FALLBACK_STYLE_PATH.read_text(encoding="utf-8"))

    def item(role: str) -> StyleItem:
        return StyleItem(role=role, source="fallback", **data.get(role, {}))

    return StyleProfile(
        body_text=item("body_text"),
        section_heading=item("section_heading"),
        figure_caption=item("figure_caption"),
        table_number=item("table_number"),
        table_caption=item("table_caption"),
        equation=item("equation"),
        equation_number=item("equation_number"),
        references=item("references"),
        source="fallback",
    )


def _infer_body_style(document, start_index: int) -> StyleItem | None:
    candidates = []
    for index, paragraph in enumerate(document.paragraphs[start_index:], start=start_index):
        text = normalize_space(paragraph.text)
        if len(text) < 40:
            continue
        if _is_special_paragraph(text) or looks_like_heading(text):
            continue
        candidates.append(_style_from_paragraph(paragraph, "body_text", source="inferred body"))
    return _most_common_style(candidates)


def _infer_heading_style(document, start_index: int) -> StyleItem | None:
    for paragraph in document.paragraphs[start_index:]:
        text = normalize_space(paragraph.text)
        if looks_like_heading(text):
            return _style_from_paragraph(paragraph, "section_heading", source="inferred heading")
    return None


def _infer_figure_caption_style(document, start_index: int) -> StyleItem | None:
    styles = []
    for match in detect_figure_captions(document, start_index):
        if match.paragraph is not None:
            paragraph = match.paragraph
        elif match.paragraph_index is not None:
            paragraph = document.paragraphs[match.paragraph_index]
        else:
            continue
        styles.append(_style_from_paragraph(paragraph, "figure_caption", "inferred figure", match.pattern))
    return _most_common_style(styles)


def _infer_table_number_style(document, start_index: int) -> StyleItem | None:
    styles = [
        _style_from_paragraph(document.paragraphs[match.paragraph_index], "table_number", "inferred table number", match.pattern)
        for match in detect_table_captions(document, start_index)
    ]
    return _most_common_style(styles)


def _infer_table_caption_style(document, start_index: int) -> StyleItem | None:
    styles = []
    for match in detect_table_captions(document, start_index):
        if match.caption_paragraph_index is None:
            continue
        styles.append(
            _style_from_paragraph(
                document.paragraphs[match.caption_paragraph_index],
                "table_caption",
                "inferred table caption",
                "next paragraph after TABLE N",
            )
        )
    return _most_common_style(styles)


def _infer_equation_style(document, start_index: int) -> StyleItem | None:
    equations = detect_equation_numbers(document, start_index)
    if not equations:
        return None
    paragraph = equations[0].paragraph
    if paragraph is None and equations[0].paragraph_index is not None:
        paragraph = document.paragraphs[equations[0].paragraph_index]
    if paragraph is None:
        return None
    return _style_from_paragraph(
        paragraph,
        "equation",
        "inferred equation",
        equations[0].pattern,
    )


def _infer_references_style(document, start_index: int) -> StyleItem | None:
    paragraphs = document.paragraphs
    for index, paragraph in enumerate(paragraphs[start_index:], start=start_index):
        if normalize_space(paragraph.text).lower() in {"references", "reference"}:
            next_index = index + 1
            while next_index < len(paragraphs):
                if normalize_space(paragraphs[next_index].text):
                    return _style_from_paragraph(paragraphs[next_index], "references", "inferred references")
                next_index += 1
    return None


def _is_special_paragraph(text: str) -> bool:
    if match_figure_caption(text) or match_table_number(text):
        return True
    lowered = text.lower()
    return lowered in {"references", "reference", "appendix"} or lowered.startswith("[")


def _most_common_style(styles: Iterable[StyleItem]) -> StyleItem | None:
    styles = list(styles)
    if not styles:
        return None
    counter = Counter(_style_key(style) for style in styles)
    key, _ = counter.most_common(1)[0]
    for style in styles:
        if _style_key(style) == key:
            return style
    return styles[0]


def _style_key(style: StyleItem) -> tuple:
    return (
        style.paragraph_style,
        style.font_name,
        style.font_size,
        style.bold,
        style.italic,
        style.alignment,
        style.left_indent,
        style.space_before,
        style.space_after,
        style.line_spacing,
    )


def _style_from_paragraph(
    paragraph: Paragraph, role: str, source: str, pattern: str | None = None
) -> StyleItem:
    run = _first_run(paragraph)
    style_font = getattr(paragraph.style, "font", None)
    paragraph_format = paragraph.paragraph_format

    font_name = _first_non_none(
        getattr(getattr(run, "font", None), "name", None),
        getattr(style_font, "name", None),
    )
    font_size = _first_non_none(
        points(getattr(getattr(run, "font", None), "size", None)),
        points(getattr(style_font, "size", None)),
    )
    bold = _first_non_none(getattr(run, "bold", None), getattr(style_font, "bold", None))
    italic = _first_non_none(getattr(run, "italic", None), getattr(style_font, "italic", None))
    underline = _first_non_none(getattr(run, "underline", None), getattr(style_font, "underline", None))
    all_caps = _first_non_none(
        getattr(getattr(run, "font", None), "all_caps", None),
        getattr(style_font, "all_caps", None),
    )

    alignment = ALIGNMENT_TO_STRING.get(paragraph.alignment)
    if alignment is None and getattr(paragraph.style, "paragraph_format", None):
        alignment = ALIGNMENT_TO_STRING.get(paragraph.style.paragraph_format.alignment)

    item = StyleItem(
        role=role,
        paragraph_style=getattr(paragraph.style, "name", None),
        font_name=font_name,
        font_size=font_size,
        bold=bold,
        italic=italic,
        underline=underline,
        all_caps=all_caps,
        alignment=alignment,
        left_indent=points(paragraph_format.left_indent),
        right_indent=points(paragraph_format.right_indent),
        first_line_indent=points(paragraph_format.first_line_indent),
        space_before=points(paragraph_format.space_before),
        space_after=points(paragraph_format.space_after),
        line_spacing=paragraph_format.line_spacing if isinstance(paragraph_format.line_spacing, float) else None,
        source=source,
        pattern=pattern,
        run_level={
            "font_name": font_name,
            "font_size": font_size,
            "bold": bold,
            "italic": italic,
            "underline": underline,
            "all_caps": all_caps,
        },
        paragraph_level={
            "alignment": alignment,
            "left_indent": points(paragraph_format.left_indent),
            "right_indent": points(paragraph_format.right_indent),
            "first_line_indent": points(paragraph_format.first_line_indent),
            "space_before": points(paragraph_format.space_before),
            "space_after": points(paragraph_format.space_after),
            "line_spacing": paragraph_format.line_spacing if isinstance(paragraph_format.line_spacing, float) else None,
        },
    )
    return item


def _merge_fallback(fallback_item: StyleItem, body: StyleItem, role: str) -> StyleItem:
    return StyleItem(
        role=role,
        paragraph_style=body.paragraph_style,
        font_name=fallback_item.font_name or body.font_name,
        font_size=fallback_item.font_size or body.font_size,
        bold=fallback_item.bold,
        italic=fallback_item.italic,
        underline=fallback_item.underline,
        all_caps=fallback_item.all_caps,
        alignment=fallback_item.alignment or body.alignment,
        left_indent=body.left_indent,
        right_indent=body.right_indent,
        first_line_indent=body.first_line_indent,
        space_before=body.space_before,
        space_after=body.space_after,
        line_spacing=body.line_spacing,
        source="fallback merged with body style",
        pattern=fallback_item.pattern,
    )


def _first_run(paragraph: Paragraph):
    for run in paragraph.runs:
        if run.text.strip():
            return run
    return paragraph.runs[0] if paragraph.runs else None


def _first_non_none(*values):
    for value in values:
        if value is not None:
            return value
    return None
