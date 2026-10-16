from __future__ import annotations

import re

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Inches
from lxml import etree

from .utils import iter_block_items, normalize_space, should_skip_for_body_inventory


MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

FINAL_EQUATION_LABEL_RE = re.compile(
    r"(?P<label>(?:(?:Eq\.?|Equation|equation)\s*)?\((?P<number>[A-Za-z]?\d+[A-Za-z]?)\))\s*$",
    re.IGNORECASE,
)
EQUATION_LABEL_TOKEN_PATTERN = r"(?:(?:Eq\.?|Equation|equation)\s*)?\((?P<number>[A-Za-z]?\d+[A-Za-z]?)\)"
EQUATION_LABEL_CLUSTER_TOKEN_PATTERN = r"(?:(?:Eq\.?|Equation|equation)\s*)?\([A-Za-z]?\d+[A-Za-z]?\)"
EQUATION_LABEL_TOKEN_RE = re.compile(EQUATION_LABEL_TOKEN_PATTERN, re.IGNORECASE)
TRAILING_EQUATION_LABELS_RE = re.compile(
    rf"(?P<labels>{EQUATION_LABEL_CLUSTER_TOKEN_PATTERN}(?:\s+{EQUATION_LABEL_CLUSTER_TOKEN_PATTERN})*)\s*$",
    re.IGNORECASE,
)
_IEEE_CITATION_SINGLE = (
    r"\["
    r"[0-9]+(?:\s*[,]\s*[0-9]+)*(?:\s*[-\u2013]\s*[0-9]+)?"
    r"(?:,\s*(?:Fig\.|Figure|fig\.|figure|eq\.|Eq\.|Equation|equation|Sec\.|Ch\.|Table|Th\.)\s*[^\]]+)?"
    r"\]"
)
IEEE_CITATION_RE = re.compile(
    rf"{_IEEE_CITATION_SINGLE}(?:\s*[-\u2013]\s*{_IEEE_CITATION_SINGLE})?",
    re.IGNORECASE,
)
EQUATION_REFERENCE_RE = re.compile(r"\b(?P<prefix>Eq\.?|Equation|equation)\s*\((?P<number>\d+)\)")
CAPTION_START_RE = re.compile(
    r"^\s*(?:(?:Fig\.?|Figure|Table|TABLE)\b|[\u56fe\u8868]\s*\d+|\u8868\u683c\s*\d+)",
    re.IGNORECASE,
)
REFERENCE_START_RE = re.compile(r"^\s*(?:Equation|Eq\.?|equation|公式|式)\s*\(\d+\)", re.IGNORECASE)
MATH_SYMBOL_RE = re.compile(
    r"(?:=|[\u002b\u002a\u002f\u00d7\u00f7\u2211\u221a\u222b\u2264\u2265\u2248\u2260\u00b1\u005e_]"
    r"|\\(?:frac|sum|int|sqrt|begin\{(?:equation|align|gather)\})"
    r")"
)
PROSE_HINT_RE = re.compile(
    r"\b(defines|referenced|shown|mentioned|exists|relationship|result|equation|figure|table|"
    r"uses?|used|formula|model|where|therefore|because|consequently|calculated|vectorially|"
    r"decomposed|responsible|sustaining|isolated|derived|allowing|within|under|"
    r"shown\s+in|as\s+shown)\b",
    re.IGNORECASE,
)
NATURAL_LANGUAGE_WORD_RE = re.compile(r"\b[A-Za-z]{2,}\b")
FORMULA_OPERATOR_RE = re.compile(
    r"(?:=|[\u00d7\u00f7\u2211\u221a\u222b\u2264\u2265\u2248\u2260\u00b1<>]|\\(?:frac|sum|int|sqrt))"
)


def has_omml_math(paragraph) -> bool:
    return any(_local_name(node) in {"oMath", "oMathPara"} for node in paragraph._p.iter())


def has_display_omml_math(paragraph) -> bool:
    """Return true only for display math, never for OMML embedded in prose."""
    if any(_local_name(node) == "oMathPara" for node in paragraph._p.iter()):
        return True
    if not has_omml_math(paragraph):
        return False
    word_text = get_visible_text_from_xml(paragraph)
    word_text = TRAILING_EQUATION_LABELS_RE.sub("", word_text or "").strip(" \t")
    return not word_text


def detect_ieee_citation_tokens(document, start_index: int = 0) -> list[dict[str, object]]:
    citations: list[dict[str, object]] = []
    for index, paragraph in enumerate(document.paragraphs[start_index:], start=start_index):
        if should_skip_for_body_inventory(paragraph, index, start_index):
            continue
        text = get_visible_text_from_xml(paragraph) or paragraph.text
        for match in IEEE_CITATION_RE.finditer(text):
            citations.append(
                {
                    "paragraph_index": index,
                    "text": match.group(0),
                    "start_char": match.start(),
                    "end_char": match.end(),
                }
            )
    return citations


def protect_ieee_citation_tokens(text: str) -> tuple[str, dict[str, str]]:
    """Replace IEEE square-bracket citations with placeholders before risky regex passes."""
    protected: dict[str, str] = {}

    def replace(match) -> str:
        key = f"__IEEE_CITATION_{len(protected)}__"
        protected[key] = match.group(0)
        return key

    return IEEE_CITATION_RE.sub(replace, text or ""), protected


def restore_ieee_citation_tokens(text: str, protected: dict[str, str]) -> str:
    restored = text or ""
    for key, original in protected.items():
        restored = restored.replace(key, original)
    return restored


def get_visible_text_from_xml(paragraph) -> str:
    texts: list[str] = []
    for run in paragraph._p.xpath(".//w:r"):
        for node in run.iter():
            local = _local_name(node)
            tag = str(node.tag)
            if local == "t" and tag.startswith(f"{{{W_NS}}}"):
                texts.append(node.text or "")
            elif local == "tab" and tag.startswith(f"{{{W_NS}}}"):
                texts.append("\t")
    return "".join(texts)


def match_equation_label_text(text: str, paragraph=None):
    text = text or ""
    protected_text, _protected = protect_ieee_citation_tokens(text)
    if protected_text != text and FINAL_EQUATION_LABEL_RE.search(protected_text) is None:
        return None
    match = FINAL_EQUATION_LABEL_RE.search(text)
    if not match:
        return None
    if _is_protected_ieee_citation(text, match):
        return None
    if _equation_label_score(text, match, paragraph) >= 3:
        return match
    return None


def match_trailing_equation_label_cluster(text: str):
    return TRAILING_EQUATION_LABELS_RE.search(text or "")


def extract_trailing_equation_numbers(text: str) -> list[int]:
    match = match_trailing_equation_label_cluster(text)
    if not match:
        return []
    numbers: list[int] = []
    for token_match in EQUATION_LABEL_TOKEN_RE.finditer(match.group("labels")):
        raw = token_match.group("number")
        number_match = re.search(r"\d+", raw or "")
        if number_match:
            numbers.append(int(number_match.group(0)))
    return numbers


def replace_trailing_equation_numbers_in_text(text: str, formatted_number: str) -> str:
    match = match_trailing_equation_label_cluster(text or "")
    if not match:
        return text
    return text[: match.start("labels")] + formatted_number + text[match.end("labels") :]


def iter_equation_text_paragraphs(document, start_index: int = 0):
    seen: set[int] = set()
    for block in iter_block_items(document):
        if block.kind == "paragraph":
            paragraph = block.element
            identity = id(paragraph._p)
            if identity in seen:
                continue
            seen.add(identity)
            if should_skip_for_body_inventory(paragraph, block.paragraph_index, start_index):
                continue
            yield block.paragraph_index, paragraph
            continue

        if block.kind != "table" or (block.paragraph_index is not None and block.paragraph_index < start_index):
            continue
        for row in block.element.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    identity = id(paragraph._p)
                    if identity in seen:
                        continue
                    seen.add(identity)
                    if should_skip_for_body_inventory(paragraph):
                        continue
                    yield None, paragraph


def collapse_trailing_equation_label_cluster(paragraph, replacement: str | None = None) -> tuple[bool, str, str]:
    before = get_visible_text_from_xml(paragraph) or paragraph.text
    match = match_trailing_equation_label_cluster(before)
    if not match:
        return False, before, before

    labels = list(EQUATION_LABEL_TOKEN_RE.finditer(match.group("labels")))
    if len(labels) < 2:
        return False, before, before

    final_label = replacement or labels[-1].group(0)
    expected = before[: match.start("labels")] + final_label + before[match.end("labels") :]
    if before == expected:
        return False, before, before
    if not replace_text_span_in_runs(paragraph, match.start("labels"), match.end("labels"), final_label):
        return False, before, before

    after = get_visible_text_from_xml(paragraph) or paragraph.text
    return after == expected, before, after


def is_equation_label_paragraph(paragraph) -> bool:
    return match_equation_label_paragraph(paragraph) is not None


def match_equation_label_paragraph(paragraph):
    if should_skip_for_body_inventory(paragraph):
        return None
    text = get_visible_text_from_xml(paragraph) or paragraph.text
    return match_equation_label_text(text, paragraph)


def is_equation_block_paragraph(paragraph) -> bool:
    if should_skip_for_body_inventory(paragraph):
        return False
    text = get_visible_text_from_xml(paragraph) or paragraph.text
    return match_equation_label_text(text, paragraph) is not None or _unnumbered_equation_score(text, paragraph) >= 3


def is_unnumbered_equation_block(paragraph) -> bool:
    if should_skip_for_body_inventory(paragraph):
        return False
    text = get_visible_text_from_xml(paragraph) or paragraph.text
    if match_equation_label_text(text, paragraph):
        return False
    return _unnumbered_equation_score(text, paragraph) >= 3


def replace_final_equation_number(
    paragraph,
    formatted_number: str,
    allow_standalone_label: bool = False,
) -> bool:
    text = get_visible_text_from_xml(paragraph) or paragraph.text
    if not allow_standalone_label and not match_equation_label_text(text, paragraph):
        return False
    match = match_trailing_equation_label_cluster(text)
    if not match:
        return False

    expected = text[: match.start("labels")] + formatted_number + text[match.end("labels") :]
    if _replace_final_number_in_text_node(paragraph, formatted_number, allow_standalone_label):
        if (get_visible_text_from_xml(paragraph) or paragraph.text) == expected:
            return True

    return replace_text_span_in_runs(paragraph, match.start("labels"), match.end("labels"), formatted_number)


def append_equation_number(paragraph, formatted_number: str) -> bool:
    text = get_visible_text_from_xml(paragraph) or paragraph.text
    if match_equation_label_text(text, paragraph):
        return False
    if paragraph.runs:
        paragraph.add_run("\t")
        paragraph.add_run(formatted_number)
    else:
        paragraph.add_run(f"\t{formatted_number}")
    _apply_right_equation_tab_stop(paragraph)
    return True


def replace_text_span_in_runs(paragraph, start: int, end: int, replacement: str) -> bool:
    cursor = 0
    inserted = False
    touched = False
    for run in paragraph.runs:
        text = run.text
        run_start = cursor
        run_end = cursor + len(text)
        cursor = run_end
        if run_end <= start or run_start >= end:
            continue

        touched = True
        overlap_start = max(start, run_start)
        overlap_end = min(end, run_end)
        prefix = text[: overlap_start - run_start]
        suffix = text[overlap_end - run_start :]
        if not inserted:
            run.text = prefix + replacement + suffix
            inserted = True
        else:
            run.text = prefix + suffix
    return touched and inserted


def _replace_final_number_in_text_node(
    paragraph,
    formatted_number: str,
    allow_standalone_label: bool = False,
) -> bool:
    for node in reversed([node for node in paragraph._p.iter() if _local_name(node) == "t"]):
        if not node.text:
            continue
        if not allow_standalone_label and not match_equation_label_text(node.text, paragraph):
            continue
        match = match_trailing_equation_label_cluster(node.text)
        if match:
            node.text = node.text[: match.start("labels")] + formatted_number + node.text[match.end("labels") :]
            return True
    return False


def _equation_label_score(text: str, match, paragraph) -> int:
    before = text[: match.start("label")].strip()
    paragraph_has_math = paragraph is not None and has_display_omml_math(paragraph)
    if paragraph is not None and has_omml_math(paragraph) and not paragraph_has_math:
        return -10
    if not before and not paragraph_has_math:
        return -10
    if before and not paragraph_has_math and _looks_like_explanatory_prose(before):
        return -10
    if CAPTION_START_RE.match(text):
        return -10
    if REFERENCE_START_RE.match(text):
        return -10
    if re.search(r"[\(\[]\d+[\)\]]\s+\w+", text):
        return -10

    score = 0
    if paragraph_has_math and not before:
        score += 4
    if "=" in before:
        score += 3
    if MATH_SYMBOL_RE.search(before):
        score += 2
    if paragraph is not None:
        if paragraph_has_math:
            score += 4
        if _style_mentions_equation(paragraph):
            score += 2
        if "\t" in text[: match.start("label")]:
            score += 2
        if paragraph.alignment in {WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT}:
            score += 1
    if re.search(r"\.{2,}\s*$", before):
        score += 2
    if len(before) > 120:
        score -= 2
    if PROSE_HINT_RE.search(before):
        score -= 4
    return score


def _unnumbered_equation_score(text: str, paragraph) -> int:
    text = normalize_space(text)
    if paragraph is not None and has_omml_math(paragraph) and not has_display_omml_math(paragraph):
        return -10
    if not text and paragraph is not None and has_display_omml_math(paragraph):
        return 5
    if not text:
        return -10
    if _looks_like_explanatory_prose(text):
        return -10
    if CAPTION_START_RE.match(text) or REFERENCE_START_RE.match(text):
        return -10
    if IEEE_CITATION_RE.search(text):
        return -6

    score = 0
    if "=" in text:
        score += 3
    if MATH_SYMBOL_RE.search(text):
        score += 2
    if paragraph is not None:
        if has_display_omml_math(paragraph):
            score += 4
        if _style_mentions_equation(paragraph):
            score += 2
        if paragraph.alignment in {WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT}:
            score += 1
    if len(text) > 100:
        score -= 3
    if len(re.findall(r"[A-Za-z]+", text)) > 10:
        score -= 4
    if PROSE_HINT_RE.search(text):
        score -= 4
    if re.search(r"[.!?]\s*$", text):
        score -= 2
    if "/" in text and "=" not in text and not has_display_omml_math(paragraph):
        score -= 5
    return score


def _looks_like_explanatory_prose(text: str) -> bool:
    """Reject natural-language explanation lines before they enter numbering."""
    cleaned = TRAILING_EQUATION_LABELS_RE.sub("", normalize_space(text or "")).strip()
    if not cleaned:
        return False
    words = NATURAL_LANGUAGE_WORD_RE.findall(cleaned)
    word_count = len(words)
    has_formula_operator = FORMULA_OPERATOR_RE.search(cleaned) is not None
    if word_count >= 4 and cleaned.endswith(":"):
        return True
    if word_count >= 6 and PROSE_HINT_RE.search(cleaned) and not has_formula_operator:
        return True
    if word_count >= 8 and not has_formula_operator:
        return True
    return False


def _apply_right_equation_tab_stop(paragraph) -> None:
    try:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(3.35), WD_TAB_ALIGNMENT.RIGHT)
    except Exception:
        pass


def _is_protected_ieee_citation(text: str, match) -> bool:
    label = match.group("label")
    if label.startswith("["):
        return True
    return False


def _style_mentions_equation(paragraph) -> bool:
    try:
        style = paragraph.style
    except Exception:
        return False
    style_name = getattr(style, "name", "") or ""
    style_id = getattr(style, "style_id", "") or ""
    return "equation" in f"{style_name} {style_id}".lower()


def _local_name(node) -> str:
    try:
        return etree.QName(node).localname
    except Exception:
        return ""
