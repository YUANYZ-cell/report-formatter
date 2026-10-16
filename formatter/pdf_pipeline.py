from __future__ import annotations

import io
from dataclasses import dataclass, replace
from typing import Any

import fitz
from docx import Document

from . import caption_rules, equation_detector, equation_rules, formatter_engine
from .formatter_engine import FormattingOptions, ProcessResult, ProcessingMode
from .streamlit_pipeline import build_output_file_names, make_run_id, sha256_bytes


@dataclass
class PdfAnalysisRun:
    result: ProcessResult
    report_bytes: bytes
    run_id: str
    report_file_name: str
    input_file_name: str
    input_sha256: str
    page_count: int
    page_text_preview: list[str]

    def debug_info(self) -> dict[str, Any]:
        return {
            "run_id": self.run_id,
            "input_file_name": self.input_file_name,
            "input_file_sha256": self.input_sha256,
            "input_type": "pdf_analysis_only",
            "pdf_page_count": self.page_count,
            "formatter_engine.__file__": str(formatter_engine.__file__),
            "caption_rules.__file__": str(caption_rules.__file__),
            "equation_rules.__file__": str(equation_rules.__file__),
            "equation_detector.__file__": str(equation_detector.__file__),
            "extracted_text_preview": self.page_text_preview,
        }


def process_uploaded_pdf(
    *,
    report_bytes: bytes,
    input_file_name: str,
    options: FormattingOptions | None = None,
) -> PdfAnalysisRun:
    """Analyze PDF text without modifying the uploaded PDF or producing a DOCX."""
    if not report_bytes.startswith(b"%PDF"):
        raise ValueError("The uploaded file is not a valid PDF.")

    try:
        pdf = fitz.open(stream=report_bytes, filetype="pdf")
    except fitz.FileDataError as exc:
        raise ValueError("The uploaded PDF could not be opened.") from exc

    try:
        if pdf.needs_pass:
            raise ValueError("Password-protected PDFs cannot be analyzed until they are unlocked.")
        analysis_document, page_by_paragraph, page_texts = _build_analysis_document(pdf)
    finally:
        pdf.close()

    base_options = options or FormattingOptions()
    analysis_options = replace(
        base_options,
        preserve_cover_page=False,
        insert_placeholder_captions=False,
        renumber_figures=False,
        renumber_tables=False,
        renumber_equations=False,
        renumber_headings=False,
    )
    result = formatter_engine.process_document(
        analysis_document,
        mode=ProcessingMode.ANALYZE_ONLY,
        style_source="Auto-detect from uploaded report",
        options=analysis_options,
    )
    _attach_pdf_page_numbers(result, page_by_paragraph)
    result.actions.clear()
    result.manual_review.insert(0, _pdf_analysis_only_message(analysis_options.ui_language, len(page_texts)))
    if not any(page_texts):
        result.manual_review.append(_no_extractable_text_message(analysis_options.ui_language))
    result.report_text = _append_pdf_report_details(
        result.report_text,
        result,
        analysis_options.ui_language,
        len(page_texts),
    )

    run_id = make_run_id()
    _, report_file_name = build_output_file_names(input_file_name, run_id)
    return PdfAnalysisRun(
        result=result,
        report_bytes=result.report_text.encode("utf-8"),
        run_id=run_id,
        report_file_name=report_file_name,
        input_file_name=input_file_name,
        input_sha256=sha256_bytes(report_bytes),
        page_count=len(page_texts),
        page_text_preview=[text for text in page_texts if text][:16],
    )


def _build_analysis_document(pdf: fitz.Document) -> tuple[Document, dict[int, int], list[str]]:
    document = Document()
    page_by_paragraph: dict[int, int] = {}
    page_texts: list[str] = []
    for page_number, page in enumerate(pdf, start=1):
        page_text = page.get_text("text").strip()
        page_texts.append(page_text)
        marker = document.add_paragraph(f"[PDF page {page_number}]")
        page_by_paragraph[_paragraph_index(document, marker)] = page_number
        for line in page_text.splitlines():
            text = line.strip()
            if not text:
                continue
            paragraph = document.add_paragraph(text)
            page_by_paragraph[_paragraph_index(document, paragraph)] = page_number
    return document, page_by_paragraph, page_texts


def _paragraph_index(document: Document, paragraph) -> int:
    return next(index for index, item in enumerate(document.paragraphs) if item._p is paragraph._p)


def _attach_pdf_page_numbers(result: ProcessResult, page_by_paragraph: dict[int, int]) -> None:
    for collection in (
        result.figures,
        result.tables,
        result.equation_blocks,
        result.equations,
        result.unnumbered_equations,
    ):
        for item in collection:
            paragraph_index = getattr(item, "caption_paragraph_index", None)
            if paragraph_index is None:
                paragraph_index = getattr(item, "paragraph_index", None)
            setattr(item, "pdf_page_number", page_by_paragraph.get(paragraph_index))
    for issue in result.issues:
        setattr(issue, "pdf_page_number", page_by_paragraph.get(issue.paragraph_index))


def _pdf_analysis_only_message(ui_language: str, page_count: int) -> str:
    if ui_language == "zh":
        return f"PDF 分析模式：已读取 {page_count} 页文本；不会修改或导出格式化后的 PDF/DOCX。"
    return f"PDF analysis only: read text from {page_count} pages; no PDF or DOCX is modified or exported."


def _no_extractable_text_message(ui_language: str) -> str:
    if ui_language == "zh":
        return "未从 PDF 提取到可分析文本；扫描件需要先完成 OCR。"
    return "No extractable PDF text was found; scanned PDFs require OCR before analysis."


def _append_pdf_report_details(report: str, result: ProcessResult, ui_language: str, page_count: int) -> str:
    if ui_language == "zh":
        lines = ["", "PDF 分析范围：", f"- 已分析页数：{page_count}", "- 仅分析，不修改 PDF 或生成格式化 DOCX。", "", "PDF 页码定位："]
        empty = "- 未定位到带页码的问题。"
    else:
        lines = ["", "PDF analysis only:", f"- Pages analyzed: {page_count}", "- No PDF is modified and no formatted DOCX is generated.", "", "PDF page locations:"]
        empty = "- No issues with a page location were detected."
    located = 0
    for issue in result.issues:
        page_number = getattr(issue, "pdf_page_number", None)
        if page_number is None:
            continue
        located += 1
        prefix = f"第 {page_number} 页" if ui_language == "zh" else f"Page {page_number}"
        lines.append(f"- {prefix}: {issue.message}")
    if not located:
        lines.append(empty)
    return report.rstrip() + "\n" + "\n".join(lines)
