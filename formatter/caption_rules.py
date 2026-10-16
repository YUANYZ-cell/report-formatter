from __future__ import annotations

import re
import hashlib
from dataclasses import dataclass

from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

from .caption_detector import (
    detect_figure_captions,
    detect_table_captions,
    match_figure_caption,
)
from .equation_utils import replace_text_span_in_runs
from .format_profiles import (
    format_figure_caption,
    format_table_caption,
    preferred_reference,
)
from .localization import CATEGORY_LABELS, label
from .style_profile import StyleProfile
from .utils import (
    apply_paragraph_style_from_profile,
    normalize_space,
    roman_to_int,
    set_paragraph_text,
    should_skip_for_body_inventory,
)


TABLE_REFERENCE_RE = re.compile(r"\b(?:TABLE|Table)\s+(?P<number>[IVXLCDM]+|\d+)\b", re.IGNORECASE)


@dataclass
class FormattingAction:
    kind: str
    paragraph_index: int | None
    before: str
    after: str
    message: str
    change_id: str = ""
    object_id: str = ""
    object_type: str = ""
    page_index: int | None = None
    source_anchor: str = ""
    original_text: str = ""
    modified_text: str = ""
    original_bbox: tuple[float, float, float, float] | None = None
    modified_bbox: tuple[float, float, float, float] | None = None
    coordinate_space: str = "pdf_normalized_top_left"
    is_cross_column: bool = False
    heading_level: int | None = None
    old_number: str | int | None = None
    new_number: str | int | None = None
    change_reason: str = ""
    confidence: float = 1.0
    location_status: str = "pending"
    equation_id: str = ""
    table_id: str = ""

    def __post_init__(self) -> None:
        self.object_type = self.object_type or self.kind
        self.original_text = self.original_text or self.before
        self.modified_text = self.modified_text or self.after
        self.change_reason = self.change_reason or self.message
        self.source_anchor = self.source_anchor or (
            f"paragraph:{self.paragraph_index}" if self.paragraph_index is not None else "document"
        )
        seed = f"{self.object_type}|{self.source_anchor}|{self.before}|{self.after}"
        digest = hashlib.sha256(seed.encode("utf-8")).hexdigest()[:16]
        self.object_id = self.object_id or f"{self.object_type}-{digest}"
        self.change_id = self.change_id or f"change-{digest}"


def renumber_figure_captions(
    document,
    style_profile: StyleProfile,
    format_profile: dict,
    report_language: str,
    ui_language: str = "en",
    start_index: int = 0,
    apply_formatting: bool = True,
) -> tuple[list[FormattingAction], dict[int, int]]:
    actions: list[FormattingAction] = []
    number_map: dict[int, int] = {}
    matches = detect_figure_captions(document, start_index)
    for new_number, match in enumerate(matches, start=1):
        if match.paragraph is not None:
            paragraph = match.paragraph
        elif match.paragraph_index is not None:
            paragraph = document.paragraphs[match.paragraph_index]
        else:
            continue
        before = normalize_space(paragraph.text)
        caption_text = match.text or _extract_figure_caption_text(before)
        after = format_figure_caption(new_number, caption_text, format_profile, report_language)
        number_map[match.number] = new_number
        changed = before != after
        if changed:
            if match.start_char not in (None, 0) or (
                match.end_char is not None and match.end_char != len(before)
            ):
                continue
            set_paragraph_text(paragraph, after)
            actions.append(
                FormattingAction(
                    "figure",
                    match.paragraph_index,
                    before,
                    after,
                    _renumber_message("figure", match.number, new_number, ui_language),
                )
            )
        if _should_apply_existing_caption_formatting(apply_formatting, format_profile, changed):
            apply_paragraph_style_from_profile(paragraph, style_profile.figure_caption)
    return actions, number_map


def renumber_table_captions(
    document,
    style_profile: StyleProfile,
    format_profile: dict,
    report_language: str,
    ui_language: str = "en",
    start_index: int = 0,
    apply_formatting: bool = True,
) -> tuple[list[FormattingAction], dict[int, int]]:
    actions: list[FormattingAction] = []
    number_map: dict[int, int] = {}
    matches = detect_table_captions(document, start_index)
    caption_paragraphs: set[object] = set()
    for new_number, match in enumerate(matches, start=1):
        if match.paragraph is not None:
            paragraph = match.paragraph
        elif match.paragraph_index is not None:
            paragraph = document.paragraphs[match.paragraph_index]
        else:
            continue
        caption_paragraphs.add(paragraph._p)
        before = normalize_space(paragraph.text)
        rendered = format_table_caption(new_number, match.text, format_profile, report_language)
        number_map[match.number] = new_number
        main_caption_changed = False

        if isinstance(rendered, list):
            number_line, caption_line = rendered
            after = number_line
            table_id = f"table-{new_number}"
            number_line_changed = before != number_line
            main_caption_changed = number_line_changed
            if number_line_changed:
                set_paragraph_text(paragraph, number_line)
                actions.append(
                    FormattingAction(
                        "table",
                        match.paragraph_index,
                        before,
                        number_line,
                        _renumber_message("table", match.number, number_line.split()[-1], ui_language),
                        object_id=table_id,
                        table_id=table_id,
                        old_number=match.number,
                        new_number=number_line.split()[-1],
                    )
                )
            if match.caption_paragraph_index is not None and match.caption_paragraph_index != match.paragraph_index:
                caption_paragraph = document.paragraphs[match.caption_paragraph_index]
                caption_paragraphs.add(caption_paragraph._p)
                caption_before = normalize_space(caption_paragraph.text)
                caption_changed = caption_before != caption_line
                if caption_changed:
                    set_paragraph_text(caption_paragraph, caption_line)
                    actions.append(
                        FormattingAction(
                            "table",
                            match.caption_paragraph_index,
                            caption_before,
                            caption_line,
                            _table_caption_text_message(ui_language),
                            object_id=table_id,
                            table_id=table_id,
                        )
                    )
                if _should_apply_existing_caption_formatting(apply_formatting, format_profile, caption_changed):
                    apply_paragraph_style_from_profile(caption_paragraph, style_profile.table_caption)
                    _keep_with_next(caption_paragraph)
            else:
                caption_paragraph = _insert_paragraph_after(paragraph, caption_line)
                caption_paragraphs.add(caption_paragraph._p)
                actions.append(
                    FormattingAction(
                        "table",
                        None,
                        "",
                        caption_line,
                        _table_caption_text_message(ui_language),
                        object_id=table_id,
                        table_id=table_id,
                        source_anchor=f"after-paragraph:{match.paragraph_index}",
                    )
                )
                if apply_formatting:
                    apply_paragraph_style_from_profile(caption_paragraph, style_profile.table_caption)
                    _keep_with_next(caption_paragraph)
        else:
            after = rendered
            changed = before != after
            main_caption_changed = changed
            if changed:
                set_paragraph_text(paragraph, after)
                actions.append(
                    FormattingAction(
                        "table",
                        match.paragraph_index,
                        before,
                        after,
                        _renumber_message("table", match.number, new_number, ui_language),
                    )
                )
            if match.caption_paragraph_index is not None and match.caption_paragraph_index != match.paragraph_index:
                caption_paragraph = document.paragraphs[match.caption_paragraph_index]
                caption_before = normalize_space(caption_paragraph.text)
                if caption_before and caption_before == match.text:
                    set_paragraph_text(caption_paragraph, "")
                    actions.append(
                        FormattingAction(
                            "table",
                            match.caption_paragraph_index,
                            caption_before,
                            "",
                            _table_caption_text_message(ui_language),
                        )
                    )

        if _should_apply_existing_caption_formatting(
            apply_formatting,
            format_profile,
            main_caption_changed,
        ):
            apply_paragraph_style_from_profile(paragraph, style_profile.table_number)
            _keep_with_next(paragraph)
    actions.extend(
        _update_table_references(document, number_map, caption_paragraphs, format_profile, ui_language, start_index)
    )
    return actions, number_map


def _should_apply_existing_caption_formatting(
    apply_formatting: bool,
    format_profile: dict,
    changed: bool,
) -> bool:
    if not apply_formatting:
        return False
    if format_profile.get("profile_key") == "template_matching" and not changed:
        return False
    return True


def _extract_figure_caption_text(text: str) -> str:
    match = match_figure_caption(text)
    if match:
        return normalize_space(match.group(3))
    return ""


def _renumber_message(kind: str, old_number, new_number, ui_language: str) -> str:
    label_text = label(CATEGORY_LABELS, ui_language, kind)
    if ui_language == "zh":
        return f"已将{label_text}编号从 {old_number} 修正为 {new_number}。"
    return f"Renumbered {label_text.lower()} {old_number} -> {new_number}."


def _update_table_references(
    document,
    number_map: dict[int, int],
    caption_paragraphs: set[object],
    format_profile: dict,
    ui_language: str,
    start_index: int,
) -> list[FormattingAction]:
    actions: list[FormattingAction] = []
    if not number_map:
        return actions

    for index, paragraph in enumerate(document.paragraphs[start_index:], start=start_index):
        if should_skip_for_body_inventory(paragraph, index, start_index):
            continue
        if paragraph._p in caption_paragraphs:
            continue
        before = paragraph.text
        replacements: list[tuple[int, int, str]] = []
        for match in TABLE_REFERENCE_RE.finditer(before):
            raw_number = match.group("number")
            old_number = int(raw_number) if raw_number.isdigit() else roman_to_int(raw_number) or 0
            new_number = number_map.get(old_number)
            if new_number is None:
                continue
            replacement = preferred_reference("table", new_number, format_profile)
            if replacement == match.group(0):
                continue
            replacements.append((match.start(), match.end(), replacement))

        if not replacements:
            continue

        expected = before
        replaced_all = True
        for start, end, replacement in reversed(replacements):
            replaced_all = replace_text_span_in_runs(paragraph, start, end, replacement) and replaced_all
            expected = expected[:start] + replacement + expected[end:]
        after = paragraph.text
        if not replaced_all or after != expected or after == before:
            continue
        actions.append(
            FormattingAction(
                "table",
                index,
                normalize_space(before),
                normalize_space(after),
                _table_reference_update_message(ui_language),
            )
        )
    return actions


def _insert_paragraph_after(paragraph, text: str) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    inserted = Paragraph(new_p, paragraph._parent)
    inserted.add_run(text)
    return inserted


def _keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True


def _table_reference_update_message(ui_language: str) -> str:
    if ui_language == "zh":
        return "已根据表格重编号同步更新正文表格引用。"
    return "Updated in-text table reference after table renumbering."


def _table_caption_text_message(ui_language: str) -> str:
    if ui_language == "zh":
        return "已根据当前格式规则更新表题文本。"
    return "Updated table caption text according to the selected format profile."
