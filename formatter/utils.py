from __future__ import annotations

import re
from dataclasses import dataclass
from typing import Iterator

from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.shared import Pt

from .style_profile import StyleItem


ALIGNMENT_TO_STRING = {
    WD_ALIGN_PARAGRAPH.LEFT: "left",
    WD_ALIGN_PARAGRAPH.CENTER: "center",
    WD_ALIGN_PARAGRAPH.RIGHT: "right",
    WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
}

STRING_TO_ALIGNMENT = {value: key for key, value in ALIGNMENT_TO_STRING.items()}

ROMAN_NUMERALS = [
    (1000, "M"),
    (900, "CM"),
    (500, "D"),
    (400, "CD"),
    (100, "C"),
    (90, "XC"),
    (50, "L"),
    (40, "XL"),
    (10, "X"),
    (9, "IX"),
    (5, "V"),
    (4, "IV"),
    (1, "I"),
]
TOC_HEADING_TEXTS = {
    "contents",
    "table of contents",
    "list of figures",
    "list of tables",
    "list of equations",
    "\u76ee\u5f55",
    "\u56fe\u76ee\u5f55",
    "\u8868\u76ee\u5f55",
    "\u516c\u5f0f\u76ee\u5f55",
}
TOC_STYLE_RE = re.compile(
    r"^(?:toc\s*\d+|toc heading|table of figures|list of figures|list of tables|list of equations)$",
    re.IGNORECASE,
)
FIELD_INSTRUCTION_RE = re.compile(r"\b(?:TOC|PAGEREF|REF|HYPERLINK|SEQ|XE)\b", re.IGNORECASE)
DOT_LEADER_RE = re.compile(r"(?:\.{3,}|\t).*(?:\b\d+\b|\(\d+\)|\[\d+\])\s*$")
MATH_IN_LEADER_RE = re.compile(r"(?:=|[\u2211\u221a\u222b\u2264\u2265\u2248\u2260\u00b1]|\\(?:frac|sum|int|sqrt))")


@dataclass
class Block:
    kind: str
    element: Paragraph | Table
    paragraph_index: int | None = None


def int_to_roman(value: int) -> str:
    if value <= 0:
        raise ValueError("Roman numerals require a positive integer")
    remainder = value
    result = []
    for integer, numeral in ROMAN_NUMERALS:
        count, remainder = divmod(remainder, integer)
        result.append(numeral * count)
    return "".join(result)


def roman_to_int(value: str) -> int | None:
    value = value.upper().strip()
    if not value or not re.fullmatch(r"[IVXLCDM]+", value):
        return None
    index = 0
    total = 0
    lookup = {numeral: integer for integer, numeral in ROMAN_NUMERALS}
    while index < len(value):
        two = value[index : index + 2]
        if two in lookup:
            total += lookup[two]
            index += 2
        else:
            total += lookup[value[index]]
            index += 1
    return total


def normalize_space(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def paragraph_has_page_break(paragraph: Paragraph) -> bool:
    return bool(paragraph._p.xpath(".//w:br[@w:type='page']"))


def paragraph_has_drawing(paragraph: Paragraph) -> bool:
    return bool(paragraph._p.xpath(".//w:drawing|.//w:pict"))


def is_generated_toc_paragraph(paragraph: Paragraph) -> bool:
    style = _paragraph_style_name(paragraph)
    if TOC_STYLE_RE.match(style):
        return True
    xml = getattr(getattr(paragraph, "_p", None), "xml", "")
    return bool(FIELD_INSTRUCTION_RE.search(xml))


def is_front_matter_index_paragraph(paragraph: Paragraph) -> bool:
    text = normalize_space(paragraph.text)
    lower = text.lower()
    if lower in TOC_HEADING_TEXTS:
        return True
    if is_generated_toc_paragraph(paragraph):
        return True
    if DOT_LEADER_RE.search(text) and not MATH_IN_LEADER_RE.search(text):
        return True
    return False


def should_skip_for_body_inventory(paragraph: Paragraph, paragraph_index: int | None = None, start_index: int = 0) -> bool:
    if paragraph_index is not None and paragraph_index < start_index:
        return True
    return is_front_matter_index_paragraph(paragraph)


def count_skipped_toc_entries(document, start_index: int = 0) -> int:
    return sum(
        1
        for index, paragraph in enumerate(document.paragraphs)
        if should_skip_for_body_inventory(paragraph, index, start_index)
    )


isGeneratedTocParagraph = is_generated_toc_paragraph
isFrontMatterIndexParagraph = is_front_matter_index_paragraph
shouldSkipForBodyInventory = should_skip_for_body_inventory


def iter_block_items(document: DocumentObject) -> Iterator[Block]:
    """Yield paragraphs and tables in document order with paragraph indices."""

    body = document.element.body
    paragraph_index = 0
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Block("paragraph", Paragraph(child, document), paragraph_index)
            paragraph_index += 1
        elif child.tag == qn("w:tbl"):
            yield Block("table", Table(child, document), paragraph_index)


def insert_paragraph_after(paragraph: Paragraph, text: str = "") -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = Paragraph(new_p, paragraph._parent)
    if text:
        new_paragraph.add_run(text)
    return new_paragraph


def insert_paragraph_before_table(table: Table, text: str = "") -> Paragraph:
    new_p = OxmlElement("w:p")
    table._tbl.addprevious(new_p)
    paragraph = Paragraph(new_p, table._parent)
    if text:
        paragraph.add_run(text)
    return paragraph


def clear_paragraph(paragraph: Paragraph) -> None:
    for child in list(paragraph._p):
        paragraph._p.remove(child)


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    clear_paragraph(paragraph)
    paragraph.add_run(text)


def points(value) -> float | None:
    if value is None:
        return None
    try:
        return float(value.pt)
    except AttributeError:
        return None


def apply_paragraph_style_from_profile(paragraph: Paragraph, style_item: StyleItem) -> None:
    if style_item.paragraph_style:
        try:
            paragraph.style = style_item.paragraph_style
        except Exception:
            pass

    fmt = paragraph.paragraph_format
    if style_item.alignment:
        paragraph.alignment = STRING_TO_ALIGNMENT.get(style_item.alignment, paragraph.alignment)
    if style_item.left_indent is not None:
        fmt.left_indent = Pt(style_item.left_indent)
    if style_item.right_indent is not None:
        fmt.right_indent = Pt(style_item.right_indent)
    if style_item.first_line_indent is not None:
        fmt.first_line_indent = Pt(style_item.first_line_indent)
    if style_item.space_before is not None:
        fmt.space_before = Pt(style_item.space_before)
    if style_item.space_after is not None:
        fmt.space_after = Pt(style_item.space_after)
    if style_item.line_spacing is not None:
        fmt.line_spacing = style_item.line_spacing

    for run in paragraph.runs:
        apply_run_style_from_profile(run, style_item)


def apply_run_style_from_profile(run, style_item: StyleItem) -> None:
    if style_item.font_name:
        run.font.name = style_item.font_name
        run._element.rPr.rFonts.set(qn("w:eastAsia"), style_item.font_name)
    if style_item.font_size:
        run.font.size = Pt(style_item.font_size)
    if style_item.bold is not None:
        run.bold = style_item.bold
    if style_item.italic is not None:
        run.italic = style_item.italic
    if style_item.underline is not None:
        run.underline = style_item.underline
    if style_item.all_caps is not None:
        run.font.all_caps = style_item.all_caps


def paragraph_index_map(document: DocumentObject) -> dict[int, Paragraph]:
    return {i: p for i, p in enumerate(document.paragraphs)}


def _paragraph_style_name(paragraph: Paragraph) -> str:
    try:
        style = paragraph.style
    except Exception:
        return ""
    style_name = getattr(style, "name", "") or ""
    style_id = getattr(style, "style_id", "") or ""
    return f"{style_name} {style_id}".strip()
