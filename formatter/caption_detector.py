from __future__ import annotations

import re
from dataclasses import dataclass, field
from io import BytesIO
from typing import Any

from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from .equation_utils import get_visible_text_from_xml
from .utils import normalize_space, roman_to_int, should_skip_for_body_inventory


FIGURE_CAPTION_PATTERNS = [
    re.compile(r"^((?:Fig\.?\s*){1,2})(\d+)\s*[.:]?\s*(.*)", re.IGNORECASE),
    re.compile(r"^(Figure)\s+(\d+)\s*[.:]?\s*(.*)", re.IGNORECASE),
    re.compile(r"^([\u56fe\u9365])\s*(\d+)\s*[.:]?\s*(.*)", re.IGNORECASE),
]
FIGURE_MENTION_RE = re.compile(
    r"\b(?P<label>(?:Fig\.?\s*){1,2}|Figures?|Figure)(?P<space>\s*)(?P<number>\d{1,3})",
    re.IGNORECASE,
)
REFERENCE_AFTER_RE = re.compile(
    r"^(?:[.:,;)]\s*)?(?:"
    r"shows?|demonstrates?|illustrates?|validates?|highlights?|depicts?|"
    r"is\s+shown|are\s+shown|was\s+shown|were\s+shown|"
    r"is\s+(?:also\s+)?referenced|are\s+(?:also\s+)?referenced|"
    r"was\s+(?:also\s+)?referenced|were\s+(?:also\s+)?referenced|"
    r"referenced|cited|mentioned|"
    r"through|to|and|&|in\s+\w+\s+shows?"
    r")\b",
    re.IGNORECASE,
)
REFERENCE_BEFORE_RE = re.compile(
    r"(?:"
    r"as\s+illustrated\s+in|as\s+shown\s+in|shown\s+in|see|"
    r"diagram\s+is\s+shown\s+in|uses?|cites?|refer(?:s|red)?\s+to"
    r")\s*$",
    re.IGNORECASE,
)
REFERENCE_CONNECTOR_BEFORE_RE = re.compile(r"(?:&|and)\s*$", re.IGNORECASE)
REFERENCE_VERB_ANYWHERE_RE = re.compile(
    r"\b(?:shows?|demonstrates?|illustrates?|validates?|highlights?|depicts?)\b",
    re.IGNORECASE,
)
FIGURE_LIST_NUMBER_RE = re.compile(r"\d{1,3}")
PICTURE_MARKERS = (
    "<pic:pic",
    "<w:drawing",
    "<wp:inline",
    "<wp:anchor",
    "<w:pict",
    "<v:imagedata",
)

TABLE_CAPTION_PATTERNS = [
    re.compile(r"^(TABLE)\s+([IVXLCDM]+)\.?\s*$", re.IGNORECASE),
    re.compile(r"^(Table)\s+([IVXLCDM]+)\.?\s*$", re.IGNORECASE),
    re.compile(r"^(Table)\s+(\d+)\.?\s*$", re.IGNORECASE),
    re.compile(r"^(Table|TABLE)\s+([IVXLCDM]+|\d+)\s*[.:\uff1a]?\s+(.+)$", re.IGNORECASE),
    re.compile(r"^(\u8868(?:\u683c)?)\s*([IVXLCDM]+|\d+)\.?\s*(.*)$", re.IGNORECASE),
]
TABLE_PROSE_START_RE = re.compile(
    r"^(?:is|are|was|were|has|have|shows?|shown|presents?|references?|referenced|"
    r"cites?|cited|mentions?|mentioned|systematically|clearly|briefly|also|then|"
    r"therefore|thus|hence|compares?|compared|summarizes?|summarized|lists?|listed|"
    r"contains?|contained|describes?|described|delineates?|delineated)\b",
    re.IGNORECASE,
)
TABLE_MENTION_RE = re.compile(
    r"\b(?:TABLE|Table)\s+([IVXLCDM]+|\d+)\b|\u8868(?:\u683c)?\s*([IVXLCDM]+|\d+)",
    re.IGNORECASE,
)


@dataclass
class CaptionMatch:
    kind: str
    paragraph_index: int | None
    number: int
    text: str
    raw_text: str
    pattern: str
    caption_paragraph_index: int | None = None
    prefix: str | None = None
    number_token: str | None = None
    inline_caption: bool = False
    start_char: int | None = None
    end_char: int | None = None
    source: str = "paragraph"
    confidence: float = 0.0
    paragraph_style: str | None = None
    block_index: int | None = None
    paragraph: Any | None = field(default=None, repr=False, compare=False)


@dataclass
class FigureReference:
    number: int
    text: str
    paragraph_index: int | None
    start_char: int
    end_char: int
    source: str
    paragraph_style: str | None = None
    block_index: int | None = None


@dataclass
class DocumentBlock:
    block_index: int
    paragraph_index: int | None
    source: str
    text: str
    paragraph_style: str | None
    has_picture: bool
    paragraph: Any | None = field(default=None, repr=False, compare=False)


@dataclass
class FigureMention:
    number: int
    text: str
    paragraph_index: int | None
    start_char: int
    end_char: int
    source: str
    classification: str
    confidence: float
    reason: str
    caption_text: str = ""
    paragraph_style: str | None = None
    block_index: int | None = None
    reference_like: bool = False
    paragraph: Any | None = field(default=None, repr=False, compare=False)


def match_figure_caption(text: str):
    normalized = normalize_space(text)
    for pattern in FIGURE_CAPTION_PATTERNS:
        match = pattern.match(normalized)
        if match:
            return match
    return None


def match_table_number(text: str):
    normalized = normalize_space(text)
    for pattern in TABLE_CAPTION_PATTERNS:
        match = pattern.match(normalized)
        if match and _is_table_caption_match(match):
            return match
    return None


def detect_figure_captions(document, start_index: int = 0) -> list[CaptionMatch]:
    matches = [
        CaptionMatch(
            kind="figure",
            paragraph_index=mention.paragraph_index,
            number=mention.number,
            text=mention.caption_text,
            raw_text=mention.text,
            pattern=FIGURE_MENTION_RE.pattern,
            prefix=_canonical_figure_prefix(mention.text),
            number_token=str(mention.number),
            start_char=mention.start_char,
            end_char=mention.end_char,
            source=mention.source,
            confidence=mention.confidence,
            paragraph_style=mention.paragraph_style,
            block_index=mention.block_index,
            paragraph=mention.paragraph,
        )
        for mention in analyze_figure_mentions(document, start_index)
        if mention.classification == "figure_caption"
    ]
    return _dedupe_figure_captions(matches)


def detect_figure_references(document, start_index: int = 0) -> list[FigureReference]:
    references: list[FigureReference] = []
    for mention in analyze_figure_mentions(document, start_index):
        if mention.classification != "figure_reference":
            continue
        for number in _reference_numbers_from_mention(mention):
            references.append(
                FigureReference(
                    number=number,
                    text=_reference_text_for_number(mention.text, number),
                    paragraph_index=mention.paragraph_index,
                    start_char=mention.start_char,
                    end_char=mention.end_char,
                    source=mention.source,
                    paragraph_style=mention.paragraph_style,
                    block_index=mention.block_index,
                )
            )
    return _dedupe_figure_references(references)


def analyze_figure_mentions(document, start_index: int = 0) -> list[FigureMention]:
    blocks = get_document_blocks_in_order(document)
    picture_block_indices = {block.block_index for block in blocks if block.has_picture}
    mentions: list[FigureMention] = []

    for block in blocks:
        if block.paragraph is not None and should_skip_for_body_inventory(block.paragraph, block.paragraph_index, start_index):
            continue
        text = block.text
        if not text or not re.search(r"\bFig(?:ure|ures|\.)?", text, re.IGNORECASE):
            continue

        starts = list(FIGURE_MENTION_RE.finditer(text))
        for position, match in enumerate(starts):
            end = starts[position + 1].start() if position + 1 < len(starts) else len(text)
            segment = normalize_space(text[match.start() : end])
            if not segment:
                continue

            has_nearby_picture = any(
                abs(block.block_index - picture_index) <= 3 for picture_index in picture_block_indices
            )
            classification, confidence, reason, reference_like = _classify_figure_segment(
                full_text=text,
                segment=segment,
                match=match,
                block=block,
                has_nearby_picture=has_nearby_picture,
                follows_figure_candidate=position > 0,
            )
            mentions.append(
                FigureMention(
                    number=int(match.group("number")),
                    text=segment,
                    paragraph_index=block.paragraph_index,
                    start_char=match.start(),
                    end_char=end,
                    source=block.source,
                    classification=classification,
                    confidence=confidence,
                    reason=reason,
                    caption_text=_extract_caption_text_from_segment(segment),
                    paragraph_style=block.paragraph_style,
                    block_index=block.block_index,
                    reference_like=reference_like or match.group("label").lower() == "figures",
                    paragraph=block.paragraph,
                )
            )
    return mentions


def get_document_blocks_in_order(document) -> list[DocumentBlock]:
    blocks: list[DocumentBlock] = []
    paragraph_index = 0

    def add_paragraph(paragraph, index: int | None, source: str) -> None:
        text = normalize_space(get_visible_text_from_xml(paragraph) or paragraph.text)
        if text or has_picture(paragraph):
            blocks.append(
                DocumentBlock(
                    block_index=len(blocks),
                    paragraph_index=index,
                    source=source,
                    text=text,
                    paragraph_style=_paragraph_style_name(paragraph),
                    has_picture=has_picture(paragraph),
                    paragraph=paragraph,
                )
            )
        for textbox_paragraph in _textbox_paragraphs(paragraph):
            textbox_text = normalize_space(get_visible_text_from_xml(textbox_paragraph) or textbox_paragraph.text)
            if not textbox_text and not has_picture(textbox_paragraph):
                continue
            blocks.append(
                DocumentBlock(
                    block_index=len(blocks),
                    paragraph_index=index,
                    source="textbox" if source == "paragraph" else f"{source}_textbox",
                    text=textbox_text,
                    paragraph_style=_paragraph_style_name(textbox_paragraph),
                    has_picture=has_picture(textbox_paragraph),
                    paragraph=textbox_paragraph,
                )
            )

    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            add_paragraph(Paragraph(child, document), paragraph_index, "paragraph")
            paragraph_index += 1
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        add_paragraph(paragraph, None, "table_cell")
    return blocks


def has_picture(paragraph) -> bool:
    xml = getattr(getattr(paragraph, "_p", None), "xml", "")
    return any(marker in xml for marker in PICTURE_MARKERS)


def detect_table_captions(document, start_index: int = 0) -> list[CaptionMatch]:
    return [
        match
        for match in detect_table_inventory(document, start_index, include_uncertain=False)
        if _is_confirmed_table_caption(match)
    ]


def detect_table_inventory(
    document,
    start_index: int = 0,
    include_uncertain: bool = True,
) -> list[CaptionMatch]:
    matches: list[CaptionMatch] = []
    paragraphs = document.paragraphs
    blocks = get_document_blocks_in_order(document)
    native_caption_indices = _native_table_caption_indices(document, start_index)
    image_caption_indices = _image_table_caption_indices(blocks, start_index)
    for index, paragraph in enumerate(paragraphs[start_index:], start=start_index):
        if should_skip_for_body_inventory(paragraph, index, start_index):
            continue
        raw = paragraph.text
        match = match_table_number(raw)
        if not match:
            continue
        prefix = match.group(1)
        token = match.group(2)
        inline_caption = bool(len(match.groups()) >= 3 and normalize_space(match.group(3)))
        number = int(token) if token.isdigit() else roman_to_int(token) or 0
        if inline_caption:
            caption_text = normalize_space(match.group(3))
            caption_index = index
        else:
            caption_index = _next_table_caption_text_index(paragraphs, index + 1)
            caption_text = normalize_space(paragraphs[caption_index].text) if caption_index is not None else ""
        matches.append(
            CaptionMatch(
                kind="table",
                paragraph_index=index,
                number=number,
                text=caption_text,
                raw_text=normalize_space(raw),
                pattern=match.re.pattern,
                caption_paragraph_index=caption_index,
                prefix=prefix,
                number_token=token,
                inline_caption=inline_caption,
                source=_table_source(index, caption_index, native_caption_indices, image_caption_indices),
                confidence=_table_confidence(index, caption_index, native_caption_indices, image_caption_indices),
                paragraph_style=_paragraph_style_name(paragraph),
                block_index=_block_index_for_paragraph(blocks, index),
                paragraph=paragraph,
            )
        )
    if include_uncertain:
        matches.extend(_uncertain_image_table_candidates(blocks, matches, start_index))
    return _dedupe_table_inventory(matches)


def _classify_figure_segment(
    *,
    full_text: str,
    segment: str,
    match,
    block: DocumentBlock,
    has_nearby_picture: bool,
    follows_figure_candidate: bool,
) -> tuple[str, float, str, bool]:
    label = match.group("label")
    label_lower = label.lower().replace(" ", "")
    compact_figure = label_lower == "figure" and not match.group("space")
    plural_reference = label_lower == "figures"
    caption_text = _extract_caption_text_from_segment(segment)
    before = full_text[max(0, match.start() - 90) : match.start()]
    after = full_text[match.end() : min(len(full_text), match.end() + 90)]
    starts_at_boundary = _starts_at_caption_boundary(full_text, match.start())
    preceded_by_running_text = bool(before.strip()) and not starts_at_boundary and not follows_figure_candidate
    reference_like = (
        plural_reference
        or compact_figure
        or bool(REFERENCE_AFTER_RE.match(after.strip()))
        or bool(REFERENCE_BEFORE_RE.search(before))
        or bool(REFERENCE_CONNECTOR_BEFORE_RE.search(before))
        or _reference_verb_soon_after_number(after)
    )

    score = 0
    reasons: list[str] = []
    if segment.lower().startswith(("fig", "figure")):
        score += 3
        reasons.append("figure prefix")
    if starts_at_boundary:
        score += 2
        reasons.append("segment boundary")
    if follows_figure_candidate:
        score += 2
        reasons.append("caption chain")
    if has_nearby_picture:
        score += 3
        reasons.append("near picture")
    if _style_mentions_caption(block.paragraph_style):
        score += 4
        reasons.append("caption style")
    if 5 <= len(caption_text) <= 160:
        score += 1
        reasons.append("caption length")
    if reference_like:
        score -= 5
        reasons.append("reference context")
    if preceded_by_running_text:
        score -= 3
        reasons.append("running prose before")
    if len(full_text) > 220 and not has_nearby_picture and not starts_at_boundary and not follows_figure_candidate:
        score -= 3
        reasons.append("long prose without nearby picture")
    if not caption_text or re.match(r"^[,;:&).]|\b(?:through|to|and)\b", caption_text, re.IGNORECASE):
        score -= 3
        reasons.append("weak caption text")

    if reference_like:
        return "figure_reference", float(score), ", ".join(reasons), True
    if score >= 4:
        return "figure_caption", float(score), ", ".join(reasons), False
    return "ignored", float(score), ", ".join(reasons), False


def _native_table_caption_indices(document, start_index: int) -> set[int]:
    indices: set[int] = set()
    paragraph_index = 0
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph_index += 1
            continue
        if child.tag != qn("w:tbl") or paragraph_index < start_index:
            continue
        for index, text in _previous_nonempty_paragraphs_with_indices(
            document.paragraphs, paragraph_index - 1, limit=5
        ):
            if should_skip_for_body_inventory(document.paragraphs[index], index, start_index):
                continue
            if not match_table_number(text):
                continue
            indices.add(index)
            caption_index = _next_nonempty_paragraph_index(document.paragraphs, index + 1)
            if caption_index is not None and caption_index < paragraph_index:
                indices.add(caption_index)
            break
    return indices


def _image_table_caption_indices(blocks: list[DocumentBlock], start_index: int) -> set[int]:
    picture_blocks = [
        block
        for block in blocks
        if block.has_picture
        and not (
            block.paragraph is not None
            and should_skip_for_body_inventory(block.paragraph, block.paragraph_index, start_index)
        )
    ]
    indices: set[int] = set()
    for picture_block in picture_blocks:
        for block in _nearby_blocks(blocks, picture_block.block_index, radius=5):
            if block.paragraph_index is None:
                continue
            if block.paragraph is not None and should_skip_for_body_inventory(block.paragraph, block.paragraph_index, start_index):
                continue
            if not match_table_number(block.text):
                continue
            indices.add(block.paragraph_index)
            caption_block = _next_text_block(blocks, block.block_index)
            if caption_block is not None and caption_block.paragraph_index is not None:
                indices.add(caption_block.paragraph_index)
    return indices


def _uncertain_image_table_candidates(
    blocks: list[DocumentBlock],
    existing: list[CaptionMatch],
    start_index: int,
) -> list[CaptionMatch]:
    existing_numbers = {item.number for item in existing}
    candidates: list[CaptionMatch] = []
    for picture_block in blocks:
        if not picture_block.has_picture:
            continue
        if picture_block.paragraph is not None and should_skip_for_body_inventory(
            picture_block.paragraph, picture_block.paragraph_index, start_index
        ):
            continue
        confidence = 0.65 if _has_gridlike_image(picture_block.paragraph) else 0.35
        for block in _nearby_blocks(blocks, picture_block.block_index, radius=5):
            if block.paragraph is not None and should_skip_for_body_inventory(block.paragraph, block.paragraph_index, start_index):
                continue
            if match_table_number(block.text):
                continue
            for match in TABLE_MENTION_RE.finditer(block.text):
                token = match.group(1) or match.group(2)
                number = int(token) if token.isdigit() else roman_to_int(token) or 0
                if not number or number in existing_numbers:
                    continue
                candidates.append(
                    CaptionMatch(
                        kind="table",
                        paragraph_index=block.paragraph_index,
                        number=number,
                        text="",
                        raw_text=normalize_space(block.text),
                        pattern=TABLE_MENTION_RE.pattern,
                        prefix="Table",
                        number_token=token,
                        inline_caption=False,
                        start_char=match.start(),
                        end_char=match.end(),
                        source="image_table_candidate",
                        confidence=confidence,
                        paragraph_style=block.paragraph_style,
                        block_index=block.block_index,
                        paragraph=block.paragraph,
                    )
                )
    return candidates


def _has_gridlike_image(paragraph) -> bool:
    return any(_image_blob_looks_gridlike(blob) for blob in _paragraph_image_blobs(paragraph))


def _paragraph_image_blobs(paragraph) -> list[bytes]:
    if paragraph is None:
        return []
    try:
        relationship_ids = paragraph._p.xpath(".//a:blip/@r:embed")
    except Exception:
        return []
    blobs: list[bytes] = []
    for relationship_id in relationship_ids:
        try:
            part = paragraph.part.related_parts.get(relationship_id)
        except Exception:
            part = None
        blob = getattr(part, "blob", None)
        if blob:
            blobs.append(blob)
    return blobs


def _image_blob_looks_gridlike(blob: bytes) -> bool:
    try:
        from PIL import Image
        import numpy as np

        image = Image.open(BytesIO(blob)).convert("L")
    except Exception:
        return False
    width, height = image.size
    if width < 30 or height < 30:
        return False
    arr = np.asarray(image)
    dark = arr < 120
    horizontal_lines = int((dark.mean(axis=1) > 0.45).sum())
    vertical_lines = int((dark.mean(axis=0) > 0.45).sum())
    return horizontal_lines >= 2 and vertical_lines >= 2


def _table_source(
    number_index: int,
    caption_index: int | None,
    native_caption_indices: set[int],
    image_caption_indices: set[int],
) -> str:
    candidate_indices = {number_index}
    if caption_index is not None:
        candidate_indices.add(caption_index)
    if candidate_indices & native_caption_indices:
        return "native_table"
    if candidate_indices & image_caption_indices:
        return "image_table_candidate"
    return "caption_only_table"


def _table_confidence(
    number_index: int,
    caption_index: int | None,
    native_caption_indices: set[int],
    image_caption_indices: set[int],
) -> float:
    source = _table_source(number_index, caption_index, native_caption_indices, image_caption_indices)
    return {"native_table": 0.95, "image_table_candidate": 0.75, "caption_only_table": 0.55}.get(source, 0.0)


def _block_index_for_paragraph(blocks: list[DocumentBlock], paragraph_index: int) -> int | None:
    for block in blocks:
        if block.paragraph_index == paragraph_index and block.source == "paragraph":
            return block.block_index
    return None


def _nearby_blocks(blocks: list[DocumentBlock], block_index: int, radius: int) -> list[DocumentBlock]:
    return [block for block in blocks if abs(block.block_index - block_index) <= radius]


def _next_text_block(blocks: list[DocumentBlock], block_index: int) -> DocumentBlock | None:
    for block in blocks:
        if block.block_index <= block_index:
            continue
        if block.text:
            return block
    return None


def _previous_nonempty_paragraphs_with_indices(paragraphs, start: int, limit: int) -> list[tuple[int, str]]:
    lines: list[tuple[int, str]] = []
    for index in range(start, -1, -1):
        text = normalize_space(paragraphs[index].text)
        if not text:
            continue
        lines.append((index, text))
        if len(lines) >= limit:
            break
    return lines


def _dedupe_table_inventory(matches: list[CaptionMatch]) -> list[CaptionMatch]:
    priority = {
        "native_table": 4,
        "image_table_confirmed": 3,
        "image_table_candidate": 3,
        "caption_only_table": 2,
    }
    selected: list[CaptionMatch] = []
    for match in matches:
        duplicate_index = _find_duplicate_table_index(selected, match)
        if duplicate_index is None:
            selected.append(match)
            continue
        existing = selected[duplicate_index]
        if priority.get(match.source, 0) > priority.get(existing.source, 0):
            selected[duplicate_index] = match
    return sorted(selected, key=lambda item: (item.block_index if item.block_index is not None else 10**9, item.number))


def dedupe_table_inventory(matches: list[CaptionMatch]) -> list[CaptionMatch]:
    return _dedupe_table_inventory(matches)


def _find_duplicate_table_index(selected: list[CaptionMatch], candidate: CaptionMatch) -> int | None:
    for index, existing in enumerate(selected):
        if existing.number != candidate.number:
            continue
        same_paragraph = existing.paragraph_index == candidate.paragraph_index
        nearby_blocks = (
            existing.block_index is not None
            and candidate.block_index is not None
            and abs(existing.block_index - candidate.block_index) <= 5
        )
        if same_paragraph or nearby_blocks:
            return index
    return None


def _starts_at_caption_boundary(text: str, start: int) -> bool:
    if start <= 0:
        return True
    before = text[:start]
    stripped = before.rstrip()
    if not stripped:
        return True
    return stripped[-1] in ".!?\n\r\t"


def _reference_verb_soon_after_number(after: str) -> bool:
    head = normalize_space(after)[:45]
    return bool(head.lower().startswith("in ") and REFERENCE_VERB_ANYWHERE_RE.search(head))


def _extract_caption_text_from_segment(segment: str) -> str:
    match = FIGURE_MENTION_RE.match(segment)
    if not match:
        return ""
    return normalize_space(segment[match.end() :].lstrip(" .:"))


def _reference_numbers_from_mention(mention: FigureMention) -> list[int]:
    numbers = [mention.number]
    text = mention.text
    lower = text.lower()

    range_match = re.search(r"\b(?:through|to|-)\s*(\d{1,3})\b", text, re.IGNORECASE)
    if range_match:
        end = int(range_match.group(1))
        if mention.number < end <= mention.number + 30:
            numbers.extend(range(mention.number + 1, end + 1))

    if lower.startswith("figures"):
        head = re.split(
            r"\b(?:shows?|demonstrates?|illustrates?|validates?|highlights?|depicts?)\b",
            text,
            maxsplit=1,
            flags=re.IGNORECASE,
        )[0]
        numbers.extend(int(value) for value in FIGURE_LIST_NUMBER_RE.findall(head))

    seen: set[int] = set()
    unique: list[int] = []
    for number in numbers:
        if number not in seen:
            seen.add(number)
            unique.append(number)
    return unique


def _reference_text_for_number(text: str, number: int) -> str:
    match = re.search(rf"\b(?:Fig\.?|Figures?|Figure)\s*{number}\b|Figure{number}\b", text, re.IGNORECASE)
    if match:
        return match.group(0)
    return f"Fig. {number}"


def _dedupe_figure_captions(matches: list[CaptionMatch]) -> list[CaptionMatch]:
    selected: list[CaptionMatch] = []
    for match in matches:
        duplicate_index = _find_duplicate_caption_index(selected, match)
        if duplicate_index is None:
            selected.append(match)
            continue
        if _caption_priority(match) > _caption_priority(selected[duplicate_index]):
            selected[duplicate_index] = match
    return sorted(selected, key=lambda item: (item.block_index if item.block_index is not None else 10**9, item.start_char or 0))


def _find_duplicate_caption_index(selected: list[CaptionMatch], candidate: CaptionMatch) -> int | None:
    for index, existing in enumerate(selected):
        if existing.number != candidate.number:
            continue
        same_area = (
            existing.paragraph_index == candidate.paragraph_index
            or (
                existing.block_index is not None
                and candidate.block_index is not None
                and abs(existing.block_index - candidate.block_index) <= 5
            )
        )
        if same_area and _caption_texts_overlap(existing.text, candidate.text):
            return index
    return None


def _caption_priority(match: CaptionMatch) -> tuple[float, int]:
    # Prefer concise, high-confidence copies when Word stores duplicate text inside shapes.
    return (match.confidence, -len(match.raw_text))


def _caption_texts_overlap(left: str, right: str) -> bool:
    left_norm = re.sub(r"[^a-z0-9]+", "", left.lower())
    right_norm = re.sub(r"[^a-z0-9]+", "", right.lower())
    if not left_norm or not right_norm:
        return True
    shortest, longest = sorted((left_norm, right_norm), key=len)
    return len(shortest) < 12 or shortest in longest or longest[:18] == shortest[:18]


def _dedupe_figure_references(references: list[FigureReference]) -> list[FigureReference]:
    selected: list[FigureReference] = []
    seen: set[tuple[int, int | None, int, str]] = set()
    for reference in references:
        key = (reference.number, reference.paragraph_index, reference.start_char, reference.source)
        if key in seen:
            continue
        seen.add(key)
        selected.append(reference)
    return selected


def _canonical_figure_prefix(text: str) -> str:
    return "Figure" if text.lower().startswith("figure") else "Fig."


def _style_mentions_caption(style_name: str | None) -> bool:
    value = (style_name or "").lower()
    return "caption" in value or "figure" in value


def _paragraph_style_name(paragraph) -> str | None:
    try:
        return paragraph.style.name
    except Exception:
        return None


def _textbox_paragraphs(paragraph) -> list[Paragraph]:
    try:
        nodes = paragraph._p.xpath(".//w:txbxContent//w:p")
    except Exception:
        return []
    return [Paragraph(node, paragraph._parent) for node in nodes]


def _is_table_caption_match(match) -> bool:
    if len(match.groups()) < 3:
        return True
    caption_text = normalize_space(match.group(3))
    if not caption_text:
        return True
    tail_after_number = match.string[match.end(2) :].lstrip()
    has_explicit_separator = tail_after_number.startswith((".", ":", "\uff1a"))
    if not has_explicit_separator and caption_text[:1].islower():
        return False
    return TABLE_PROSE_START_RE.match(caption_text) is None


def _is_confirmed_table_caption(match: CaptionMatch) -> bool:
    return match.source in {"native_table", "image_table_confirmed", "image_table_candidate"} and match.confidence >= 0.7


def _next_nonempty_paragraph_index(paragraphs, start: int) -> int | None:
    for index in range(start, len(paragraphs)):
        if normalize_space(paragraphs[index].text):
            return index
    return None


def _next_table_caption_text_index(paragraphs, start: int) -> int | None:
    for index in range(start, len(paragraphs)):
        paragraph = paragraphs[index]
        if has_picture(paragraph):
            return None
        text = normalize_space(paragraph.text)
        if not text:
            continue
        if match_table_number(text) or match_figure_caption(text):
            return None
        if TABLE_MENTION_RE.search(text):
            return None
        return index
    return None
