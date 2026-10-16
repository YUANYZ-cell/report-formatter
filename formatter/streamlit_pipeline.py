from __future__ import annotations

import hashlib
import io
import json
import re
import uuid
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document

from . import caption_rules, equation_detector, equation_rules, formatter_engine
from .formatter_engine import FormattingOptions, ProcessingMode, ProcessResult


ROOT = Path(__file__).resolve().parents[1]
STALE_ROOT_FILES = ("formatted_report.docx", "format_check_report.txt")
STALE_OUTPUT_PATTERNS = (
    "formatted_report*.docx",
    "*_formatted_*.docx",
    "format_check_report*.txt",
    "*_check_report_*.txt",
)
PREVIEW_PATTERNS = (
    "V = IR",
    "Equation (1) defines",
    "P = VI",
    "Eq. (",
    "Motor speed response after step input",
    "TABLE II",
    "MEASURED MOTOR SPEED RESULTS",
)


@dataclass
class StreamlitPipelineRun:
    result: ProcessResult
    docx_bytes: bytes
    report_bytes: bytes
    run_id: str
    docx_file_name: str
    report_file_name: str
    input_file_name: str
    input_sha256: str
    output_docx_sha256: str
    template_sha256: str | None
    formatter_version: str
    module_paths: dict[str, str]
    key_text_preview: list[str]

    def debug_info(self) -> dict[str, Any]:
        return {
            "run_id": self.run_id,
            "input_file_name": self.input_file_name,
            "input_file_sha256": self.input_sha256,
            "output_docx_sha256": self.output_docx_sha256,
            "template_file_sha256": self.template_sha256,
            "formatter_engine.__file__": self.module_paths["formatter_engine"],
            "caption_rules.__file__": self.module_paths["caption_rules"],
            "equation_rules.__file__": self.module_paths["equation_rules"],
            "equation_detector.__file__": self.module_paths["equation_detector"],
            "formatter_version": self.formatter_version,
            "exported_docx_preview": self.key_text_preview,
        }


def build_processing_signature(
    *,
    report_bytes: bytes,
    report_file_name: str,
    template_bytes: bytes | None,
    template_file_name: str | None,
    mode: str | ProcessingMode,
    style_source: str,
    options: FormattingOptions,
) -> str:
    options_payload = asdict(options)
    # UI labels are rendered from the current locale, so changing only the UI language must not clear cached results.
    options_payload.pop("ui_language", None)
    payload = {
        "report_file_name": report_file_name,
        "report_sha256": sha256_bytes(report_bytes),
        "template_file_name": template_file_name,
        "template_sha256": sha256_bytes(template_bytes) if template_bytes else None,
        "mode": ProcessingMode(mode).value,
        "style_source": style_source,
        "options": options_payload,
        "formatter_version": formatter_engine.FORMATTER_VERSION,
    }
    encoded = json.dumps(payload, sort_keys=True, ensure_ascii=True, default=str).encode("utf-8")
    return sha256_bytes(encoded)


def process_uploaded_docx(
    *,
    report_bytes: bytes,
    input_file_name: str,
    template_bytes: bytes | None = None,
    template_file_name: str | None = None,
    mode: str | ProcessingMode = ProcessingMode.ANALYZE_ONLY,
    style_source: str = "Auto-detect from uploaded report",
    options: FormattingOptions | None = None,
) -> StreamlitPipelineRun:
    cleanup_old_outputs()
    options = options or FormattingOptions()
    run_id = make_run_id()

    report_doc = Document(io.BytesIO(report_bytes))
    template_doc = Document(io.BytesIO(template_bytes)) if template_bytes else None
    result = formatter_engine.process_document(
        report_doc,
        template_document=template_doc,
        mode=mode,
        style_source=style_source,
        options=options,
    )
    docx_buffer = io.BytesIO()
    report_doc.save(docx_buffer)
    docx_bytes = docx_buffer.getvalue()
    report_text = result.report_text or ""
    report_bytes_out = report_text.encode("utf-8") if report_text else b""
    docx_file_name, report_file_name = build_output_file_names(input_file_name, run_id)

    return StreamlitPipelineRun(
        result=result,
        docx_bytes=docx_bytes,
        report_bytes=report_bytes_out,
        run_id=run_id,
        docx_file_name=docx_file_name,
        report_file_name=report_file_name,
        input_file_name=input_file_name,
        input_sha256=sha256_bytes(report_bytes),
        output_docx_sha256=sha256_bytes(docx_bytes),
        template_sha256=sha256_bytes(template_bytes) if template_bytes else None,
        formatter_version=formatter_engine.FORMATTER_VERSION,
        module_paths=module_debug_paths(),
        key_text_preview=extract_key_text_preview(docx_bytes),
    )


def cleanup_old_outputs(root: Path = ROOT) -> list[str]:
    removed: list[str] = []
    for file_name in STALE_ROOT_FILES:
        path = root / file_name
        if path.exists() and path.is_file():
            path.unlink()
            removed.append(str(path))

    for directory_name in ("output", "temp", "tmp"):
        directory = root / directory_name
        if not directory.exists() or not directory.is_dir():
            continue
        for pattern in STALE_OUTPUT_PATTERNS:
            for path in directory.glob(pattern):
                if path.is_file():
                    path.unlink()
                    removed.append(str(path))
    return removed


def extract_docx_text_lines(docx_bytes: bytes) -> list[str]:
    document = Document(io.BytesIO(docx_bytes))
    return [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]


def extract_key_text_preview(docx_bytes: bytes, limit: int = 16) -> list[str]:
    lines = extract_docx_text_lines(docx_bytes)
    matched: list[str] = []
    for line in lines:
        if any(pattern in line for pattern in PREVIEW_PATTERNS):
            matched.append(line)
    if matched:
        return matched[:limit]
    return lines[:limit]


def module_debug_paths() -> dict[str, str]:
    return {
        "formatter_engine": str(Path(formatter_engine.__file__).resolve()),
        "caption_rules": str(Path(caption_rules.__file__).resolve()),
        "equation_rules": str(Path(equation_rules.__file__).resolve()),
        "equation_detector": str(Path(equation_detector.__file__).resolve()),
    }


def sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def make_run_id() -> str:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"{timestamp}_{uuid.uuid4().hex[:8]}"


def build_output_file_names(input_file_name: str, run_id: str) -> tuple[str, str]:
    stem = safe_file_stem(input_file_name)
    return f"{stem}_formatted_{run_id}.docx", f"{stem}_check_report_{run_id}.txt"


def safe_file_stem(file_name: str) -> str:
    stem = Path(file_name or "report").stem.strip() or "report"
    stem = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "_", stem)
    stem = re.sub(r"\s+", "_", stem)
    return stem[:120] or "report"
