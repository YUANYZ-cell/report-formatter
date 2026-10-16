from __future__ import annotations

import re

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph

from .caption_rules import FormattingAction
from .equation_utils import replace_text_span_in_runs
from .utils import int_to_roman, normalize_space, set_paragraph_text


DECIMAL_HEADING_RE = re.compile(r"^(?P<number>\d+(?:\.\d+){0,3})\.?\s+(?P<title>.+)$")
TRAILING_PAGE_RE = re.compile(r"(?:\.{2,}|\t|\s{2,})\s*\d+\s*$")
NON_SECTION_TITLES = {"abstract", "references", "acknowledgment", "acknowledgements", "appendix"}


def renumber_ieee_headings(
    document,
    *,
    start_index: int,
    ui_language: str = "en",
    apply_formatting: bool = True,
) -> list[FormattingAction]:
    """Map semantic heading levels first, then render IEEE section labels."""
    title_levels = _toc_title_levels(document, start_index)
    candidates = _body_heading_candidates(document, start_index, title_levels)
    if not candidates:
        return []

    if apply_formatting:
        _configure_heading_styles(document)

    actions: list[FormattingAction] = []
    counters = [0, 0, 0, 0]
    for paragraph_index, paragraph, level, title, consumed_chars in candidates:
        for lower in range(level, 4):
            counters[lower] = 0
        counters[level - 1] += 1
        rendered = _render_heading(level, counters[level - 1], title)
        before = normalize_space(paragraph.text)
        object_id = f"heading-{level}-{counters[level - 1]}-{paragraph_index}"

        if before == rendered:
            if apply_formatting:
                paragraph.style = document.styles[f"Heading {level}"]
                _clear_direct_numbering(paragraph)
                paragraph.paragraph_format.keep_with_next = True
                paragraph.paragraph_format.keep_together = True
            continue

        if consumed_chars < len(before):
            inserted = _insert_paragraph_before(paragraph, rendered)
            replace_text_span_in_runs(paragraph, 0, consumed_chars, "")
            _trim_leading_space(paragraph)
            target = inserted
        else:
            set_paragraph_text(paragraph, rendered)
            target = paragraph

        if apply_formatting:
            target.style = document.styles[f"Heading {level}"]
            _clear_direct_numbering(target)
            target.paragraph_format.keep_with_next = True
            target.paragraph_format.keep_together = True

        actions.append(
            FormattingAction(
                "heading",
                paragraph_index,
                before,
                rendered,
                _heading_message(level, ui_language),
                object_id=object_id,
                object_type="heading",
                source_anchor=f"paragraph:{paragraph_index}",
                heading_level=level,
                old_number=_leading_number(before),
                new_number=rendered.split(maxsplit=1)[0],
                is_cross_column=level == 1,
            )
        )
    return actions


def _toc_title_levels(document, start_index: int) -> dict[str, tuple[int, str]]:
    levels: dict[str, tuple[int, str]] = {}
    for paragraph in document.paragraphs[:start_index]:
        text = normalize_space(paragraph.text)
        match = DECIMAL_HEADING_RE.match(text)
        if not match:
            continue
        number = match.group("number")
        title = TRAILING_PAGE_RE.sub("", match.group("title")).strip(" .\t")
        duplicated = DECIMAL_HEADING_RE.match(title)
        if duplicated:
            number = duplicated.group("number")
            title = duplicated.group("title").strip(" .\t")
        if not title:
            continue
        levels[_title_key(title)] = (min(4, number.count(".") + 1), title)
    return levels


def _body_heading_candidates(document, start_index: int, title_levels: dict[str, tuple[int, str]]):
    candidates = []
    for index in range(start_index, len(document.paragraphs)):
        paragraph = document.paragraphs[index]
        text = normalize_space(paragraph.text)
        if not text or text.lower().strip(" .") in NON_SECTION_TITLES:
            continue
        style_name = (getattr(getattr(paragraph, "style", None), "name", "") or "").lower()
        style_level_match = re.search(r"heading\s*([1-4])", style_name)
        style_level = int(style_level_match.group(1)) if style_level_match else None
        decimal = DECIMAL_HEADING_RE.match(text)
        prefix_len = 0
        title_part = text
        decimal_level = None
        existing_ieee = _strip_existing_ieee_prefix(text, style_level)
        if existing_ieee:
            title_part = existing_ieee
            prefix_len = len(text) - len(existing_ieee)
        if decimal:
            decimal_level = min(4, decimal.group("number").count(".") + 1)
            title_part = decimal.group("title")
            prefix_len = decimal.start("title")

        exact = title_levels.get(_title_key(title_part))
        if exact:
            level, canonical = exact
            consumed = len(text)
        else:
            matched = _longest_title_prefix(title_part, title_levels)
            if matched:
                level, canonical, title_chars = matched
                consumed = prefix_len + title_chars
            elif style_level:
                level = style_level
                canonical = title_part
                consumed = len(text)
            elif decimal_level and _looks_like_short_heading(title_part):
                level = decimal_level
                canonical = title_part
                consumed = len(text)
            else:
                continue
        candidates.append((index, paragraph, level, canonical.strip(" .:"), consumed))
    return candidates


def _longest_title_prefix(text: str, levels: dict[str, tuple[int, str]]):
    normalized = normalize_space(text)
    matches = []
    for level, canonical in levels.values():
        if normalized.casefold().startswith(canonical.casefold()) and (
            len(normalized) == len(canonical) or normalized[len(canonical)].isspace()
        ):
            matches.append((len(canonical), level, canonical))
    if not matches:
        return None
    length, level, canonical = max(matches)
    return level, canonical, length


def _looks_like_short_heading(text: str) -> bool:
    words = re.findall(r"[A-Za-z]+", text)
    return 1 <= len(words) <= 12 and not text.rstrip().endswith(('.', '?', '!'))


def _strip_existing_ieee_prefix(text: str, level: int | None) -> str | None:
    patterns = {
        1: r"^[IVXLCDM]+\.\s+",
        2: r"^[A-Z]+\.\s+",
        3: r"^\d+\)\s+",
        4: r"^[a-z]+\)\s+",
    }
    if level not in patterns:
        return None
    match = re.match(patterns[level], text)
    if not match:
        return None
    return text[match.end() :].rstrip(":")


def _render_heading(level: int, number: int, title: str) -> str:
    if level == 1:
        return f"{int_to_roman(number)}. {title.upper()}"
    if level == 2:
        return f"{_alpha(number)}. {title}"
    if level == 3:
        return f"{number}) {title}:"
    return f"{_alpha(number).lower()}) {title}:"


def _alpha(number: int) -> str:
    result = ""
    value = number
    while value:
        value, remainder = divmod(value - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _configure_heading_styles(document) -> None:
    for level in range(1, 5):
        style = document.styles[f"Heading {level}"]
        style.font.name = "Times New Roman"
        style.font.size = Pt(10)
        style.font.bold = level == 1
        style.font.italic = level in {2, 3, 4}
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.space_before = Pt(6)
        style.paragraph_format.space_after = Pt(0)
        style_p_pr = style._element.get_or_add_pPr()
        style_num_pr = style_p_pr.find(qn("w:numPr"))
        if style_num_pr is not None:
            style_p_pr.remove(style_num_pr)


def _insert_paragraph_before(paragraph, text: str) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    inserted = Paragraph(new_p, paragraph._parent)
    inserted.add_run(text)
    return inserted


def _trim_leading_space(paragraph) -> None:
    for run in paragraph.runs:
        if run.text:
            run.text = run.text.lstrip()
            break


def _clear_direct_numbering(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is not None:
        p_pr.remove(num_pr)


def _title_key(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", text.casefold())


def _leading_number(text: str) -> str | None:
    match = DECIMAL_HEADING_RE.match(text)
    return match.group("number") if match else None


def _heading_message(level: int, ui_language: str) -> str:
    if ui_language == "zh":
        return f"已按语义层级将 {level} 级标题转换为 IEEE 编号与样式。"
    return f"Converted level-{level} heading to IEEE numbering and style."
